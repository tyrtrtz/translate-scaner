"""Read-only discovery of contiguous merges and option-splitting patterns."""

import argparse
import json
import os
import re
import shutil
import sqlite3
import tempfile
import unicodedata
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ignored_sheets import load_ignored, write_json
from merge_reviews import apply_review_file, candidate_identity, load_reviews
from matcher import (TAIL_MATCH, TEXT_TAG, baseline_sheets, breath_key, match_key, plain_key, prepare_inputs,
                     trailing_note_key,
                     source_metadata_columns, source_row_metadata, write_detail_sheet)
from scanner import hidden_row_numbers, normalize, scan_folder


LINE_BREAK = re.compile(r"\r\n|\r|\n|\\r\\n|\\[nr]")
DISPLAY_TAG = re.compile(r"</?(?:color|ruby)(?:\s*=[^<>]*|\s+[^<>]*)?>", re.I)
MARKER = re.compile(
    r"(?:(?<![A-Za-zＡ-Ｚａ-ｚ0-9０-９_%{])|(?<=\\n)|(?<=\\r))"
    r"(?:(?P<plain>[A-Za-zＡ-Ｚａ-ｚ]|[0-9０-９]{1,3})[.．:：、,，;；)）]"
    r"|[(（](?P<bracket>[A-Za-zＡ-Ｚａ-ｚ]|[0-9０-９]{1,3})[)）])"
)
SCENE_HEADERS = {"场景", "场景名", "场景编号"}
# Optional evidence only: formula/ambiguous columns stay unknown, never guessed.
METADATA_HEADERS = {
    "角色": ("说话角色名", "说话角色", "角色名", "角色", "Speaker"),
    "剧情段": ("序号", "剧情段", "段落编号"),
    "语音类型": ("语音类型",),
    "语音名": ("语音名\n(ver_3.28)", "语音名"),
    "播放说明": ("语音说明 · 播放逻辑",),
    "业务Key": ("Key",),
    "条目来源": ("Comment",),
}
MAX_CHARACTERS = 16000
PLACEHOLDER = re.compile(r"\{[^{}\r\n]+\}|(?<!%)%(?:\d+\$)?[sd]")


def comparison_key(text):
    """Keep placeholders, punctuation, case and internal spaces significant."""
    return LINE_BREAK.sub("", TEXT_TAG.sub("", text)).strip()


def trim_segment(text):
    return re.sub(r"^(?:\s|\\[nr])+|(?:\s|\\[nr])+$", "", text)


def split_patterns(text):
    """Return hypotheses with raw boundaries; none is an approved split rule."""
    result = []
    markers = [
        m for m in MARKER.finditer(text)
        if not (m.group().endswith((".", "．", ",", "，")) and text[m.end():m.end() + 1].isdigit())
    ]
    if len(markers) >= 2:
        labels = [unicodedata.normalize("NFKC", m.group("plain") or m.group("bracket")).upper()
                  for m in markers]
        kinds = {"数字" if label.isdigit() else "字母" for label in labels}
        rule = next(iter(kinds)) + "编号" if len(kinds) == 1 else "混合编号"
        ends = [m.start() for m in markers[1:]] + [len(text)]
        parts = [text[m.end():end] for m, end in zip(markers, ends)]
        prefix = text[:markers[0].start()]
        issues = []
        if trim_segment(prefix):
            issues.append("首个编号前有未归属文字")
        if len(set(labels)) != len(labels):
            issues.append("编号重复，可能包含多组选项")
        if len(kinds) > 1:
            issues.append("数字与字母编号混用")
        else:
            values = [int(label) if label.isdigit() else ord(label) - ord("A") + 1 for label in labels]
            if values != list(range(1, len(values) + 1)):
                issues.append("编号未从1或A连续排列")
        if any(not comparison_key(part) for part in parts):
            issues.append("存在空选项")
        if any(LINE_BREAK.search(trim_segment(part)) for part in parts):
            issues.append("选项内部含换行，需要确认其含义")
        result.append(dict(rule=rule, labels=labels, parts=parts, prefix=prefix,
                           separators=[m.group() for m in markers],
                           spans=[[m.end(), end] for m, end in zip(markers, ends)], issues=issues))
    breaks = list(LINE_BREAK.finditer(text))
    if breaks:
        kinds = {"字面换行" if m.group().startswith("\\") else "实际换行" for m in breaks}
        parts = LINE_BREAK.split(text)
        start = 0
        spans = []
        for m in breaks:
            spans.append([start, m.start()])
            start = m.end()
        spans.append([start, len(text)])
        result.append(dict(rule=next(iter(kinds)) if len(kinds) == 1 else "混合换行",
                           labels=[], parts=parts, prefix="", spans=spans,
                           separators=[m.group() for m in breaks],
                           issues=["存在空分段"] if any(not comparison_key(p) for p in parts) else []))
    return result


def bilingual_evidence(pattern, korean):
    alternatives = split_patterns(korean)
    if not korean.strip():
        return "缺译", alternatives
    if not alternatives:
        return "韩文无可识别分隔符；不能按中文长度切分", alternatives
    if pattern["labels"]:
        aligned = [p for p in alternatives if p["labels"] == pattern["labels"]]
        if aligned and not pattern["issues"] and any(not p["issues"] for p in aligned):
            return "编号、顺序与段数一致；内容归属仍需人工复核", alternatives
        if any(p["labels"] for p in alternatives):
            return "编号或分段存在歧义；需复核中韩归属", alternatives
    if any(len(p["parts"]) == len(pattern["parts"]) for p in alternatives):
        return "仅段数相同，无法据此确认中韩边界", alternatives
    return "中韩分段数量不一致", alternatives


@dataclass
class StructureReport:
    destination: Path
    check_report: object
    source_rows: int
    baseline_rows: int
    merge_candidates: int
    split_candidates: int
    pattern_samples: int
    unresolved: int
    skipped: int
    trial_groups: int
    trial_segments: int
    trial_pending: int
    whole_matched_rows: int
    merge_trial_rows: int
    merge_pending_rows: int
    merge_rejected_rows: int


def create_index(db):
    db.executescript("""
        CREATE TABLE records (
            id INTEGER PRIMARY KEY, kind TEXT, path TEXT, sheet TEXT, row INTEGER,
            block INTEGER, original TEXT, korean TEXT, text TEXT, scene TEXT, metadata TEXT, note TEXT);
        CREATE INDEX record_text ON records(kind,text);
        CREATE UNIQUE INDEX record_position ON records(kind,path,sheet,row);
        CREATE TABLE variants (
            id INTEGER PRIMARY KEY, record_id INTEGER, rule TEXT, text TEXT, details TEXT);
        CREATE INDEX variant_text ON variants(text);
        CREATE INDEX variant_record ON variants(record_id);
        CREATE TABLE links (
            id INTEGER PRIMARY KEY, kind TEXT, one_id INTEGER, part_ids TEXT, variant_id INTEGER);
        CREATE INDEX link_one ON links(kind,one_id);
        CREATE TABLE covered (record_id INTEGER PRIMARY KEY);
        CREATE TABLE whole_matches (
            record_id INTEGER PRIMARY KEY, match_type TEXT, candidates INTEGER, variants INTEGER);
        CREATE TABLE whole_matched_sources (record_id INTEGER PRIMARY KEY);
        CREATE TABLE whole_match_pairs (baseline_id INTEGER,source_id INTEGER,PRIMARY KEY(baseline_id,source_id));
        CREATE TABLE trial_coverage (source_id INTEGER PRIMARY KEY,matched INTEGER,total INTEGER);
        CREATE TABLE merge_targets (record_id INTEGER PRIMARY KEY);
    """)


def index_sheet(db, kind, path, sheet, header_row, text_column, korean_column=None,
                include_hidden=True, include_deprecated=False, skipped=None):
    header = next(sheet.iter_rows(min_row=header_row, max_row=header_row))
    notes_column, speaker_column = source_metadata_columns(header) if kind == "source" else (None, None)
    scenes = [c.column for c in header if normalize(c.value) in SCENE_HEADERS]
    if len(scenes) > 1:
        raise ValueError("存在多个场景列，无法确定句段边界")
    scene_column = scenes[0] if scenes else None
    metadata_columns = {
        name: [c.column for c in header if normalize(c.value) in {normalize(a) for a in aliases}]
        for name, aliases in METADATA_HEADERS.items()
    }
    hidden = hidden_row_numbers(sheet.parent, sheet) if not include_hidden else set()
    columns = [c for c in (text_column, korean_column, notes_column, speaker_column, scene_column) if c]
    columns.extend(c for matches in metadata_columns.values() for c in matches)
    first, last = min(columns), max(columns)
    previous_row, block, scene = 0, 0, ""
    hidden_count = deprecated_count = invalid_count = empty_count = 0
    for number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, min_col=first, max_col=last), header_row + 1):
        cell = lambda column: cells[column - first]
        if number in hidden:
            hidden_count += 1
            continue
        note_text, _, excluded = source_row_metadata(
            cell(notes_column) if notes_column else None,
            cell(speaker_column) if speaker_column else None, number, include_deprecated)
        if excluded:
            deprecated_count += excluded == "废弃"
            invalid_count += excluded == "无效说明"
            continue
        if any(cell(c).data_type == "f" for c in (text_column, korean_column, scene_column) if c):
            raise ValueError(f"第 {number} 行中文、韩文或场景含公式，请先转为文本值")
        scene_value = cell(scene_column).value if scene_column else None
        next_scene = normalize(str(scene_value)) if scene_value is not None else ""
        if next_scene and next_scene != scene:
            scene = next_scene
            block += 1
        original = cell(text_column).value
        if not isinstance(original, str) or not original.strip():
            empty_count += 1
            continue
        if number != previous_row + 1 or not comparison_key(original):
            block += 1
        previous_row = number
        korean = cell(korean_column).value if korean_column else None
        korean = "" if korean is None else str(korean)
        metadata = {}
        for name, matches in metadata_columns.items():
            if len(matches) > 1:
                metadata[name] = "[未知：多个同义列]"
            elif matches:
                value = cell(matches[0])
                metadata[name] = ("[未知：公式]" if value.data_type == "f" else
                                  normalize(str(value.value)) if value.value is not None else "")
        cursor = db.execute("INSERT INTO records VALUES (NULL,?,?,?,?,?,?,?,?,?,?,?)",
                            (kind, str(path), sheet.title, number, block, original, korean,
                             comparison_key(original), scene, json.dumps(metadata, ensure_ascii=False), note_text))
        if kind == "source":
            record_id = cursor.lastrowid
            db.execute("INSERT INTO variants VALUES (NULL,?,?,?,?)",
                       (record_id, "原文整体", comparison_key(original), "{}"))
            for pattern in split_patterns(original):
                joined = comparison_key(pattern["prefix"]) + "".join(comparison_key(p) for p in pattern["parts"])
                db.execute("INSERT INTO variants VALUES (NULL,?,?,?,?)",
                           (record_id, pattern["rule"], joined, json.dumps(pattern, ensure_ascii=False)))
        if not comparison_key(original):
            block += 1
    if skipped is not None:
        for count, reason in ((hidden_count, "排除隐藏行"), (deprecated_count, "排除废弃行"),
                              (invalid_count, "排除特殊说明与说话角色名均为空的无效说明行"),
                              (empty_count, "空白或非文本中文行")):
            if count:
                skipped.append([path.name, str(path), sheet.title, f"{reason} {count} 行；连续窗口不跨这些行"])


def load_records(db, report, candidates, baselines, skipped, progress):
    for path in baselines:
        progress(f"读取基准：{path.name}")
        db.execute("SAVEPOINT file")
        book = None
        local_skips = []
        try:
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            for sheet, header, column, _ in baseline_sheets(book, path, report.header_rows, for_output=False):
                index_sheet(db, "baseline", path, sheet, header, column, skipped=local_skips)
        except sqlite3.Error:
            raise
        except Exception as error:
            db.execute("ROLLBACK TO file")
            skipped.append([path.name, str(path), "", f"基准文件跳过：{type(error).__name__}: {error}"])
        else:
            skipped.extend(local_skips)
        finally:
            db.execute("RELEASE file")
            if book:
                book.close()
    for path, results in candidates.items():
        progress(f"读取台本：{path.name}")
        book = None
        try:
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            for result in results:
                db.execute("SAVEPOINT sheet")
                local_skips = []
                try:
                    sheet = book[result.sheet]
                    sheet.reset_dimensions()
                    index_sheet(db, "source", path, sheet, result.header_row, result.chinese_column,
                                result.korean_column, report.include_hidden_rows,
                                report.include_deprecated, local_skips)
                except sqlite3.Error:
                    raise
                except Exception as error:
                    db.execute("ROLLBACK TO sheet")
                    result.issue = f"正文读取失败：{type(error).__name__}: {error}"
                    skipped.append([path.name, str(path), result.sheet, result.issue])
                else:
                    skipped.extend(local_skips)
                finally:
                    db.execute("RELEASE sheet")
        except sqlite3.Error:
            raise
        except Exception as error:
            for result in results:
                result.issue = f"源文件读取失败：{type(error).__name__}: {error}"
                skipped.append([path.name, str(path), result.sheet, result.issue])
        finally:
            if book:
                book.close()
    for kind, name in (("source", "台本"), ("baseline", "基准")):
        if not db.execute("SELECT 1 FROM records WHERE kind=? LIMIT 1", (kind,)).fetchone():
            reasons = "\n".join(str(s[-1]) for s in skipped[:8])
            raise ValueError(f"没有可扫描的{name}正文，未生成结果。\n{reasons}")
    db.commit()


def index_whole_matches(db, progress):
    """Share exact/plain/breath keys; retain both ends of whole-sentence matches."""
    progress("先检查全部基准的整句匹配；合并反查仅处理 char.xlsx 中未命中的行")
    db.execute("""CREATE TEMP TABLE whole_source_keys
        (record_id INTEGER PRIMARY KEY,exact TEXT,plain TEXT,breath TEXT,breath_plain TEXT,korean TEXT)""")
    for source in db.execute("SELECT id,original,korean,note FROM records WHERE kind='source'"):
        breath = breath_key(source["original"]) if source["note"] == "气息" else ""
        db.execute("INSERT INTO whole_source_keys VALUES (?,?,?,?,?,?)",
                   (source["id"], match_key(source["original"]), plain_key(source["original"]),
                    breath, plain_key(breath), source["korean"].strip()))
    db.executescript("""
        CREATE INDEX whole_exact ON whole_source_keys(exact);
        CREATE INDEX whole_plain ON whole_source_keys(plain);
        CREATE INDEX whole_breath ON whole_source_keys(breath);
        CREATE INDEX whole_breath_plain ON whole_source_keys(breath_plain);
    """)
    for baseline in db.execute("SELECT id,original FROM records WHERE kind='baseline'"):
        for plain in (False, True):
            key = plain_key(baseline["original"]) if plain else match_key(baseline["original"])
            if not key:
                continue
            column, breath_column = ("plain", "breath_plain") if plain else ("exact", "breath")
            condition = f"{column}=? OR ({breath_column}!='' AND {breath_column}=?)"
            arguments = (key, breath_key(key))
            count, variants, exact = db.execute(f"""SELECT COUNT(*),COUNT(DISTINCT NULLIF(korean,'')),
                MAX({column}=?) FROM whole_source_keys WHERE {condition}""", (key, *arguments)).fetchone()
            if count:
                mode = "纯文本补匹配" if plain else "全文精确匹配" if exact else "气息去括号匹配"
                db.execute("INSERT INTO whole_matches VALUES (?,?,?,?)",
                           (baseline["id"], mode, count, variants))
                db.execute(f"""INSERT OR IGNORE INTO whole_matched_sources
                    SELECT record_id FROM whole_source_keys WHERE {condition}""", arguments)
                db.execute(f"""INSERT INTO whole_match_pairs
                    SELECT ?,record_id FROM whole_source_keys WHERE {condition}""", (baseline["id"], *arguments))
                break
    db.execute("CREATE TEMP TABLE whole_tail_keys (source_id INTEGER PRIMARY KEY,text TEXT)")
    for source in db.execute("""SELECT record_id,plain FROM whole_source_keys s WHERE NOT EXISTS
        (SELECT 1 FROM whole_matched_sources w WHERE w.record_id=s.record_id)"""):
        key = trailing_note_key(source["plain"])
        if key and key != source["plain"]:
            db.execute("INSERT INTO whole_tail_keys VALUES (?,?)", (source["record_id"], key))
    db.execute("CREATE INDEX whole_tail_text ON whole_tail_keys(text)")
    for baseline in db.execute("SELECT id,original FROM records WHERE kind='baseline'"):
        matches = list(db.execute("SELECT source_id FROM whole_tail_keys WHERE text=?", (plain_key(baseline["original"]),)))
        if not matches:
            continue
        db.executemany("INSERT OR IGNORE INTO whole_match_pairs VALUES (?,?)", [(baseline["id"], r[0]) for r in matches])
        db.executemany("INSERT OR IGNORE INTO whole_matched_sources VALUES (?)", matches)
        count, variants = db.execute("""SELECT COUNT(*),COUNT(DISTINCT NULLIF(s.korean,'')) FROM whole_match_pairs p
            JOIN whole_source_keys s ON s.record_id=p.source_id WHERE p.baseline_id=?""", (baseline["id"],)).fetchone()
        db.execute("""INSERT INTO whole_matches VALUES (?,?,?,?) ON CONFLICT(record_id) DO UPDATE
            SET candidates=excluded.candidates,variants=excluded.variants""", (baseline["id"], TAIL_MATCH, count, variants))
    db.execute("DROP TABLE whole_tail_keys")
    db.execute("DROP TABLE whole_source_keys")
    db.commit()


def is_merge_baseline(path):
    return Path(path).name.casefold() == "char.xlsx"


def index_merge_targets(db):
    """Restrict before discovery, so saved reviews cannot re-enable story merges."""
    for (path,) in db.execute("SELECT DISTINCT path FROM records WHERE kind='baseline'"):
        if is_merge_baseline(path):
            db.execute("""INSERT INTO merge_targets
                SELECT r.id FROM records r WHERE r.kind='baseline' AND r.path=?
                AND NOT EXISTS (SELECT 1 FROM whole_matches w WHERE w.record_id=r.id)""", (path,))
    db.commit()


def find_links(db, max_rows, progress):
    # ponytail: O(rows * max_rows) indexed probes. Keep the window bounded; use
    # a prefix index only if measurements on customer files justify the extra state.
    for kind, link_kind in (("source", "合并"), ("baseline", "拆分")):
        window, previous = deque(maxlen=max_rows), None
        for count, row in enumerate(db.execute("SELECT * FROM records WHERE kind=? ORDER BY path,sheet,row", (kind,)), 1):
            group = (row["path"], row["sheet"], row["block"])
            if group != previous or not row["text"]:
                window.clear()
            previous = group
            if not row["text"]:
                continue
            window.append(row)
            joined = row["text"]
            parts = [row["id"]]
            for before in reversed(list(window)[:-1]):
                joined = before["text"] + joined
                parts.insert(0, before["id"])
                if len(joined) > MAX_CHARACTERS:
                    break
                if kind == "source":
                    matches = db.execute("""SELECT r.id,NULL FROM records r
                        JOIN merge_targets m ON m.record_id=r.id WHERE r.kind='baseline' AND r.text=?""", (joined,))
                else:
                    matches = db.execute("SELECT record_id,id FROM variants WHERE text=?", (joined,))
                for one_id, variant_id in matches:
                    db.execute("INSERT INTO links VALUES (NULL,?,?,?,?)",
                               (link_kind, one_id, json.dumps(parts), variant_id))
                    db.executemany("INSERT OR IGNORE INTO covered VALUES (?)", [(i,) for i in [one_id, *parts]])
            if count % 5000 == 0:
                progress(f"扫描{link_kind}候选：已处理 {count} 行")
    db.commit()


def get_record(db, record_id):
    return db.execute("SELECT * FROM records WHERE id=?", (record_id,)).fetchone()


def context_check(db, sources, baselines):
    messages = []
    anchored = False
    for label, edge, direction in (("前", 0, -1), ("后", -1, 1)):
        source, baseline = sources[edge], baselines[edge]
        for distance in range(1, 4):
            pair = []
            for anchor in (source, baseline):
                pair.append(db.execute("""SELECT * FROM records WHERE kind=? AND path=? AND sheet=?
                    AND block=? AND row=?""", (anchor["kind"], anchor["path"], anchor["sheet"],
                                              anchor["block"], anchor["row"] + distance * direction)).fetchone())
            if any(item is None for item in pair):
                break
            if not pair[0]["text"] or pair[0]["text"] != pair[1]["text"]:
                break
            unique = all(db.execute("SELECT COUNT(*) FROM records WHERE kind=? AND text=?",
                                    (item["kind"], item["text"])).fetchone()[0] == 1 for item in pair)
            anchored |= unique
            messages.append(f"{label}{distance}句一致（{'双方全文唯一' if unique else '重复文本，仅供参考'}）："
                            f"源行{pair[0]['row']} / 基准行{pair[1]['row']}")
    return "；".join(messages) or "没有连续相同的邻句证据，位置需复核", anchored


def context_evidence(db, sources, baselines):
    return context_check(db, sources, baselines)[0]


def preview(text):
    if len(text) <= 24000:
        return text
    return text[:23500] + "\n[展示截断；完整内容按记录ID回查“原始内容”，分段可按字符范围回查]"


def rendered_patterns(patterns):
    blocks = []
    for pattern in patterns:
        lines = [f"规则：{pattern['rule']}"]
        if pattern["labels"]:
            lines.append("编号：" + " / ".join(pattern["labels"]))
        lines.append("分隔符：" + " / ".join(s.replace("\r", "\\r").replace("\n", "\\n") for s in pattern["separators"]))
        if pattern["prefix"].strip():
            lines.append("编号前文字：" + pattern["prefix"])
        for index, (part, (start, end)) in enumerate(zip(pattern["parts"], pattern["spans"]), 1):
            lines.append(f"第{index}段（原文字符{start + 1}～{end}）：" + (part or "[空分段]"))
        if pattern["issues"]:
            lines.append("需复核：" + "；".join(pattern["issues"]))
        blocks.append("\n".join(lines))
    return preview("\n\n".join(blocks)) or "未识别到分隔符"


def row_text(rows, column):
    return preview("\n".join(f"[记录{r['id']} / 行{r['row']}] {r[column]}" for r in rows))


def pattern_samples(db, stats):
    for row in db.execute("SELECT * FROM records WHERE kind='source' ORDER BY id"):
        cn = split_patterns(row["original"])
        ko = split_patterns(row["korean"])
        for language, patterns in (("中文", cn), ("韩文", ko)):
            for index, pattern in enumerate(patterns, 1):
                sample_id = f"S{row['id']}-{language}-{index}"
                evidence, alternatives = bilingual_evidence(pattern, row["korean"]) if language == "中文" else (
                    "韩文格式记录；需结合中文与基准核对归属", cn)
                matched = 0
                if language == "中文":
                    matched = db.execute("""SELECT COUNT(*) FROM links l JOIN variants v ON l.variant_id=v.id
                        WHERE v.record_id=? AND v.rule=?""", (row["id"], pattern["rule"])).fetchone()[0]
                issues = pattern["issues"] + ["仅格式候选，普通数字/字母或自然换行也可能触发；尚未批准为拆分规则"]
                entry = stats[(language, pattern["rule"])]
                entry["count"] += 1
                entry["files"].add(row["path"])
                entry["matched"] += bool(matched)
                entry["issues"] += bool(pattern["issues"])
                if len(entry["examples"]) < 5:
                    entry["examples"].append(sample_id)
                yield [sample_id, row["id"], row["path"], row["sheet"], row["row"], language, pattern["rule"],
                       row["original"], row["korean"], rendered_patterns([pattern]),
                       len(pattern["parts"]), evidence, rendered_patterns(alternatives),
                       matched if language == "中文" else None, "；".join(issues)]


def candidate_rows(db, kind, assessments=None):
    @lru_cache(maxsize=1024)
    def merge_info(one_id):
        candidates = list(db.execute("SELECT part_ids FROM links WHERE kind='合并' AND one_id=?", (one_id,)))
        translations = {tuple(get_record(db, i)["korean"].strip() for i in json.loads(c[0])) for c in candidates}
        return len(candidates), len(translations) > 1

    for link in db.execute("SELECT * FROM links WHERE kind=? ORDER BY id", (kind,)):
        one = get_record(db, link["one_id"])
        parts = [get_record(db, i) for i in json.loads(link["part_ids"])]
        sources, baselines = (parts, [one]) if kind == "合并" else ([one], parts)
        variant = db.execute("SELECT * FROM variants WHERE id=?", (link["variant_id"],)).fetchone()
        rule = "连续句段拼接" if variant is None else variant["rule"]
        if kind == "合并":
            variants, conflict = merge_info(one["id"])
            assessment = f"完整中文一致；同一基准共 {variants} 个合并候选；译文连接方式待确认"
            assessment += "；包含缺译" if any(not r["korean"].strip() for r in sources) else ""
            assessment += "；候选之间的译文序列不同，需复核冲突" if conflict else ""
            if assessments is not None:
                decision = assessments[link["id"]]
                assessment = decision["state"] + "：" + decision["reason"] + "；" + assessment
            segments = "译文按来源逐行展示，未自动拼接"
        else:
            pattern = json.loads(variant["details"])
            if not pattern:
                pattern = dict(parts=[r["original"] for r in parts], labels=[], issues=[], prefix="")
                assessment = "基准连续多行完整覆盖源中文；源中文无已验证拆分边界"
            else:
                boundaries = (not pattern["prefix"].strip()
                              and [comparison_key(p) for p in pattern["parts"]] == [r["text"] for r in parts])
                assessment = "中文分段边界与基准逐条一致" if boundaries else "仅拼接后的中文一致，分段边界或前缀仍需复核"
                if pattern["issues"]:
                    assessment += "；" + "；".join(pattern["issues"])
            evidence, alternatives = bilingual_evidence(pattern, one["korean"])
            assessment += "；" + evidence
            segments = rendered_patterns(alternatives)
        yield [link["id"], rule, sources[0]["path"], sources[0]["sheet"],
               "、".join(str(r["row"]) for r in sources), row_text(sources, "original"), row_text(sources, "korean"),
               baselines[0]["path"], baselines[0]["sheet"], "、".join(str(r["row"]) for r in baselines),
               row_text(baselines, "original"), context_evidence(db, sources, baselines), assessment, segments,
               metadata_text(sources), metadata_text(baselines)]


def unresolved_rows(db, max_rows):
    for row in db.execute("""SELECT r.* FROM records r WHERE NOT EXISTS
        (SELECT 1 FROM covered c WHERE c.record_id=r.id) AND NOT EXISTS
        (SELECT 1 FROM whole_matches w WHERE w.record_id=r.id) AND NOT EXISTS
        (SELECT 1 FROM whole_matched_sources w WHERE w.record_id=r.id) AND NOT EXISTS
        (SELECT 1 FROM records other WHERE other.kind=CASE r.kind WHEN 'source' THEN 'baseline' ELSE 'source' END
         AND other.text=r.text AND r.text!='')
        ORDER BY kind,path,sheet,row"""):
        reason = f"在连续 {max_rows} 行、拼接 {MAX_CHARACTERS} 字符的范围内未解释；不代表单侧多余或缺失"
        if not row["text"]:
            reason = "清理展示标签后无有效中文；作为窗口边界保留"
        elif len(row["text"]) > MAX_CHARACTERS:
            reason = f"比较文本超过 {MAX_CHARACTERS} 字符，未参与连续窗口匹配"
        if row["kind"] == "source" and split_patterns(row["original"]):
            reason += "；已发现格式候选，详见规则样本"
            coverage = db.execute("SELECT matched,total FROM trial_coverage WHERE source_id=?", (row["id"],)).fetchone()
            if coverage and coverage[0]:
                reason += f"；试拆分已定位{coverage[0]}/{coverage[1]}段，剩余见分段未匹配"
        if row["kind"] == "baseline" and not is_merge_baseline(row["path"]):
            reason += "；本基准不在 char.xlsx 合并范围内，未执行合并反查"
        else:
            reason += "；合并反查仅限 char.xlsx"
        yield [row["id"], "台本" if row["kind"] == "source" else "基准", row["path"], row["sheet"],
               row["row"], row["original"], row["korean"], reason]


def balanced_display_tags(text):
    stack = []
    for match in DISPLAY_TAG.finditer(text):
        tag = match.group().lower()
        name = re.match(r"</?([a-z]+)", tag).group(1)
        if tag.startswith("</"):
            if not stack or stack.pop() != name:
                return False
        else:
            stack.append(name)
    return not stack


def trial_merge(sources, baseline):
    """Keep sentence boundaries visible in this preview; don't invent Korean spacing."""
    if "".join(comparison_key(s["original"]) for s in sources) != comparison_key(baseline["original"]):
        return None, "连续源中文未完整覆盖基准"
    content = PLACEHOLDER.sub("", comparison_key(baseline["original"]))
    if not any(unicodedata.category(char).startswith("L") for char in content):
        return None, "只有标点、数字或占位符，缺少可确认的对白内容"
    for source in sources:
        if not source["korean"].strip():
            return None, f"源行{source['row']}缺译"
        if Counter(PLACEHOLDER.findall(source["original"])) != Counter(PLACEHOLDER.findall(source["korean"])):
            return None, f"源行{source['row']}中韩占位符不一致"
        if not all(balanced_display_tags(source[key]) for key in ("original", "korean")):
            return None, f"源行{source['row']}展示标签跨行或未闭合"
    return "\n".join(s["korean"].strip() for s in sources), ""


def metadata_text(rows):
    return preview("\n".join(
        f"行{r['row']}：" + "；".join(f"{k}={v}" for k, v in
                                    {"场景": r["scene"], **json.loads(r["metadata"])}.items() if v)
        for r in rows))


def metadata_conflict(sources, baseline):
    metadata = [json.loads(row["metadata"]) for row in [*sources, baseline]]

    def known(items, key):
        return {m[key] for m in items if m.get(key) and not m[key].startswith("[未知：")}

    if len(known(metadata[:-1], "角色")) > 1:
        return "连续台词属于不同角色"
    if len(known(metadata, "角色")) > 1:
        return "源角色与基准角色不一致"
    if len(known(metadata[:-1], "剧情段")) > 1:
        return "连续台词跨剧情段"
    if len(known(metadata[:-1], "语音类型")) > 1:
        return "连续台词的语音类型不同"
    return ""


def merge_assessments(db, reviews, review_report):
    candidates, by_baseline, owners = {}, defaultdict(list), defaultdict(set)
    for link in db.execute("SELECT * FROM links WHERE kind='合并' ORDER BY id"):
        sources = [get_record(db, i) for i in json.loads(link["part_ids"])]
        baseline = get_record(db, link["one_id"])
        evidence, anchored = context_check(db, sources, [baseline])
        neighbors = []
        for group in (sources, [baseline]):
            first, last = group[0], group[-1]
            neighbors.extend(db.execute("""SELECT * FROM records WHERE kind=? AND path=? AND sheet=?
                AND (row BETWEEN ? AND ? OR row BETWEEN ? AND ?) ORDER BY row""",
                (first["kind"], first["path"], first["sheet"], first["row"] - 3, first["row"] - 1,
                 last["row"] + 1, last["row"] + 3)))
        identity, fingerprint = candidate_identity(sources, baseline, neighbors, evidence)
        translation, error = trial_merge(sources, baseline)
        candidates[link["id"]] = dict(
            link=link, sources=sources, baseline=baseline, identity=identity, fingerprint=fingerprint,
            evidence=evidence, anchored=anchored, translation=translation, error=error,
            conflict=metadata_conflict(sources, baseline))
        by_baseline[baseline["id"]].append(candidates[link["id"]])
        for source in sources:
            owners[source["id"]].add(baseline["id"])
    reviews = apply_review_file(review_report, reviews, candidates)
    for group in by_baseline.values():
        for candidate in group:
            review = reviews.get(candidate["identity"])
            stale = review is not None and review["fingerprint"] != candidate["fingerprint"]
            decision = "" if review is None or stale else review["decision"]
            candidate.update(review=review, stale=stale, decision=decision)
            state, reason = "待确认", "缺少唯一邻句定位证据；同角色、同译文不能确认位置"
            if decision == "不适用":
                state, reason = "不适用", "人工复核不适用：" + review["comment"]
            elif candidate["conflict"]:
                state, reason = "不适用", candidate["conflict"]
            elif candidate["error"]:
                reason = candidate["error"]
            elif stale:
                reason = "复核已过期：内容、角色或相关上下文已变化，需要重新确认"
            elif decision == "待确认":
                reason = "人工保留待确认：" + review["comment"]
            elif any(len(owners[s["id"]]) > 1 for s in candidate["sources"]):
                reason = "源行同时参与多个基准的合并候选，位置归属不唯一"
            elif decision == "可适用":
                state, reason = "人工确认可适用", review["comment"] or "复核确认本组对应关系"
            elif len(group) > 1:
                reason = "存在多个来源，需确认具体位置；即使译文相同也不自动采用"
            elif candidate["anchored"]:
                state, reason = "可试合并", "唯一邻句定位通过；语义仍待人工复核"
            candidate.update(state=state, reason=reason)
        selected = [c for c in group if c["state"] in ("人工确认可适用", "可试合并")]
        if len(selected) > 1:
            for candidate in selected:
                candidate.update(state="待确认", reason="同一基准有多组通过复核的来源，请只确认实际对应的一组")
    return candidates, reviews


def write_merge_trial_report(db, destination, max_rows, progress, assessments):
    by_baseline = defaultdict(list)
    for link in db.execute("SELECT * FROM links WHERE kind='合并' ORDER BY id"):
        sources = [get_record(db, i) for i in json.loads(link["part_ids"])]
        by_baseline[link["one_id"]].append((link, sources))
    accepted, candidates, review_rows, pending, rejected = [], [], [], 0, 0
    for baseline_id, links in by_baseline.items():
        baseline = get_record(db, baseline_id)
        selected = [assessments[link["id"]] for link, _ in links
                    if assessments[link["id"]]["state"] in ("可试合并", "人工确认可适用")]
        if selected:
            chosen = selected[0]
            accepted.append([f"M{baseline_id}", baseline["original"], chosen["translation"], 1, chosen["state"],
                             baseline["path"], baseline["sheet"], baseline["row"],
                             chosen["link"]["id"]])
        elif all(assessments[link["id"]]["state"] == "不适用" for link, _ in links):
            rejected += 1
        else:
            pending += 1
        for link, sources in links:
            item = assessments[link["id"]]
            judgment = item["state"] + "：" + item["reason"]
            candidates.append([link["id"], f"M{baseline_id}", baseline["original"], item["translation"], judgment,
                               sources[0]["path"], sources[0]["sheet"], "、".join(str(s["row"]) for s in sources),
                               row_text(sources, "original"), row_text(sources, "korean"),
                               baseline["path"], baseline["sheet"], baseline["row"],
                               item["evidence"], metadata_text(sources), metadata_text([baseline])])
            review_rows.append([link["id"], item["decision"],
                                (item["review"] or {}).get("comment", ""), judgment,
                                sources[0]["path"], sources[0]["sheet"], "、".join(str(s["row"]) for s in sources),
                                row_text(sources, "original"), row_text(sources, "korean"),
                                baseline["path"], baseline["sheet"], baseline["row"], baseline["original"],
                                metadata_text(sources), metadata_text([baseline]), item["evidence"],
                                item["identity"], item["fingerprint"]])
    baseline_rows = db.execute("SELECT COUNT(*) FROM records WHERE kind='baseline'").fetchone()[0]
    whole_counts = dict(db.execute("SELECT match_type,COUNT(*) FROM whole_matches GROUP BY match_type"))
    whole = sum(whole_counts.values())
    merge_targets = db.execute("SELECT COUNT(*) FROM merge_targets").fetchone()[0]
    book = Workbook(write_only=True)
    write_detail_sheet(book, "合并建议",
        ["组ID", "基准中文", "试合并韩文", "来源组数", "状态", "基准完整路径", "基准Sheet", "基准行号", "候选ID"],
        accepted, [16, 70, 75, 14, 55, 55, 22, 14, 22])
    write_detail_sheet(book, "全部反查候选",
        ["候选ID", "组ID", "基准中文", "候选合并韩文", "采用判断", "源完整路径", "源Sheet", "源实际行号",
         "源中文（含记录ID）", "源韩文（含记录ID）", "基准完整路径", "基准Sheet", "基准行号", "邻句证据", "源角色与剧情信息", "基准角色与条目信息"],
        candidates, [14, 16, 70, 75, 80, 55, 22, 22, 70, 75, 55, 22, 14, 70, 70, 70])
    write_detail_sheet(book, "复核填写",
        ["候选ID", "复核结论", "复核意见", "当前判断", "源完整路径", "源Sheet", "源实际行号",
         "源中文", "源韩文", "基准完整路径", "基准Sheet", "基准行号", "基准中文",
         "源角色与剧情信息", "基准角色与条目信息", "邻句证据", "对应标识", "内容指纹"],
        review_rows, [12, 18, 70, 75, 55, 22, 22, 70, 75, 55, 22, 14, 70, 70, 70, 70, 68, 68])
    write_detail_sheet(book, "整句命中检查",
        ["基准完整路径", "基准Sheet", "基准行号", "基准中文", "整句命中方式", "源候选数", "非空译文种类数", "处理"],
        ([r["path"], r["sheet"], r["row"], r["original"], r["match_type"], r["candidates"], r["variants"],
          "已有中文候选，全部缺译；不进入合并反查" if r["variants"] == 0 else
          "已有中文候选；推荐或冲突仍由常规匹配处理，不进入合并反查"]
         for r in db.execute("""SELECT r.*,w.match_type,w.candidates,w.variants FROM whole_matches w
                                JOIN records r ON r.id=w.record_id ORDER BY r.id""")),
        [55, 22, 14, 70, 24, 14, 22, 85], progress)
    write_detail_sheet(book, "合并范围",
        ["基准完整路径", "参与基准行数", "整句命中行数", "进入合并反查行数", "范围说明"],
        ([r["path"], r["total"], r["whole"], r["targets"],
          "仅未命中整句的行参与合并反查，仍需角色及位置核验" if is_merge_baseline(r["path"]) else
          "不参与合并反查；保留整句匹配及拆分扫描，不能据此认定不存在合并"]
         for r in db.execute("""SELECT r.path,COUNT(*) AS total,COUNT(w.record_id) AS whole,
             COUNT(m.record_id) AS targets FROM records r
             LEFT JOIN whole_matches w ON w.record_id=r.id
             LEFT JOIN merge_targets m ON m.record_id=r.id
             WHERE r.kind='baseline' GROUP BY r.path ORDER BY r.path""")),
        [65, 18, 18, 22, 85])
    write_detail_sheet(book, "流程统计与说明", ["项目", "数量或说明"], [
        ["参与基准行数", baseline_rows],
        ["整句精确命中", whole_counts.get("全文精确匹配", 0)],
        ["气息去括号命中", whole_counts.get("气息去括号匹配", 0)],
        ["仅纯文本命中", whole_counts.get("纯文本补匹配", 0)],
        ["句尾去括号命中", whole_counts.get(TAIL_MATCH, 0)],
        ["进入合并反查的基准行", merge_targets],
        ["未命中整句但不在合并范围的基准行", baseline_rows - whole - merge_targets],
        ["反查找到候选的基准行", len(by_baseline)],
        ["反查候选组数", len(candidates)],
        ["试合并建议基准行", len(accepted)],
        ["反查有候选但暂不采用的基准行", pending],
        ["全部候选不适用的基准行", rejected],
        ["反查仍无候选的基准行", merge_targets - len(by_baseline)],
        ["流程", "先对全部基准检查整句匹配；仅基准文件名为char.xlsx（不区分大小写）且所有整句阶段均无中文候选的行反查连续源行"],
        ["合并范围", "story.xlsx及其他基准不参与合并反查；未读取到char.xlsx正文时合并候选为零，不扩大到其他文件。分文件数量见合并范围表"],
        ["已知场景", "目前已确认样本来自角色必杀技分段台词；语音名及播放说明供核对，char.xlsx中的其他候选仍需定位或人工确认"],
        ["命中与推荐", "整句命中检查不代替正式匹配的上下文推荐；有中文候选但缺译或冲突的行不算未匹配，不用合并覆盖"],
        ["合并方式", "按源实际行号顺序，以实际换行连接去除首尾空白后的韩文；内部换行、标签、标点、占位符保留；保留整组原文供回查"],
        ["采用条件", "人工明确确认，或单一来源且有双方全文唯一的连续邻句证据；角色/剧情段冲突、缺译、占位符或标签异常、跨基准争用仍阻止采用"],
        ["多个候选", "同角色或同译文不代表同一位置；多个来源须明确复核具体一组；理论可用归为待确认"],
        ["复核方法", "在复核填写中填写可适用、不适用或待确认及意见；界面选择该文件后重跑。空结论不覆盖旧记录，颜色不作为批准"],
        ["复核保存", "结论保存在输出根目录合并复核记录.json，下次扫描自动读取。路径、行号变化后需重新对应；内容、角色或相关上下文变化后旧复核失效"],
        ["范围外复核", "历史story复核记录保留但不参与本次合并；含范围外候选的旧复核表请勿重导，应使用最新复核填写表"],
        ["范围", f"沿用结构差异扫描.xlsx中的范围和跳过清单，最多连续{max_rows}行；反查未命中不代表不存在对应；拆分预览仍独立生成"],
        ["试用阶段", "仅生成独立合并建议，语义及位置仍需复核；未改动原文件、未回填TextTrans，也未改写正式匹配结果"],
    ], [38, 115])
    book.save(destination)
    book.close()
    return whole, len(accepted), pending, rejected


def split_translation(chinese, korean):
    """Use numbered options first, otherwise keep newline-delimited Korean verbatim."""
    patterns = []
    cn_patterns = split_patterns(chinese)
    numbered = any(p["labels"] for p in cn_patterns)
    for language, text in (("中文", chinese), ("韩文", korean)):
        pattern = next((p for p in split_patterns(text)
                        if (p["rule"] == "数字编号" if numbered else "换行" in p["rule"])), None)
        if pattern is None:
            return None, f"{language}没有可用的{'数字编号' if numbered else '数字编号或换行分段'}；不能可靠分配译文"
        if not numbered:
            # Ignore only blank cell edges; internal blank lines must not shift translations.
            for edge in (0, -1):
                while pattern["parts"] and not pattern["parts"][edge].strip():
                    pattern["parts"].pop(edge)
                    pattern["spans"].pop(edge)
            pattern["issues"] = ["存在空分段"] if any(not comparison_key(p) for p in pattern["parts"]) else []
            if len(pattern["parts"]) < 2:
                return None, f"{language}不足两条有效换行分段"
        if pattern["issues"]:
            return None, f"{language}：" + "；".join(pattern["issues"])
        patterns.append(pattern)
    cn, ko = patterns
    if cn["labels"] != ko["labels"] or len(cn["parts"]) != len(ko["parts"]):
        return None, "中韩编号、顺序或段数不一致"
    if not numbered:
        ko_numbered = next((p for p in split_patterns(korean) if p["labels"]), None)
        if ko_numbered and (ko_numbered["issues"] or len(ko_numbered["parts"]) != len(cn["parts"])):
            return None, "韩文编号或行边界有歧义；保留整组待复核"
    for index, (chinese_part, korean_part) in enumerate(zip(cn["parts"], ko["parts"]), 1):
        if Counter(PLACEHOLDER.findall(chinese_part)) != Counter(PLACEHOLDER.findall(korean_part)):
            return None, f"第{index}段中韩占位符不一致"
        if not all(balanced_display_tags(p) for p in (chinese_part, korean_part)):
            return None, f"第{index}段展示标签跨越边界或未闭合"
    return patterns, ""


def trial_numbered_split(chinese, korean, baseline_texts):
    patterns, reason = split_translation(chinese, korean)
    if patterns and [comparison_key(p) for p in patterns[0]["parts"]] != [comparison_key(p) for p in baseline_texts]:
        return None, "中文分段边界与基准未逐条一致"
    return patterns, reason


def write_trial_report(db, destination, progress, max_rows=8):
    # One physical source can match several windows; do not pick the first.
    mappings, owners = defaultdict(set), defaultdict(set)
    for link in db.execute("SELECT * FROM links WHERE kind='拆分' ORDER BY id"):
        parts = tuple(json.loads(link["part_ids"]))
        mappings[link["one_id"]].add(parts)
    prepared, pending = [], []
    for source in db.execute("SELECT * FROM records WHERE kind='source' ORDER BY id"):
        targets = mappings[source["id"]]
        if not targets and not (split_patterns(source["original"]) or split_patterns(source["korean"])):
            continue
        patterns, reason = split_translation(source["original"], source["korean"])
        if patterns and (len(patterns[0]["parts"]) > max_rows or len(comparison_key(source["original"])) > MAX_CHARACTERS):
            patterns, reason = None, f"分段数超过{max_rows}或正文超过{MAX_CHARACTERS}字符；请调整范围或人工复核"
        if patterns is None:
            pending.append([source["id"], source["original"], source["korean"], reason,
                            source["path"], source["sheet"], source["row"], len(targets)])
        else:
            keys = [comparison_key(p) for p in patterns[0]["parts"]]
            full = [ids for ids in targets if [get_record(db, i)["text"] for i in ids] == keys]
            choices = []
            for index, key in enumerate(keys):
                ids = [full[0][index]] if len(full) == 1 else [r[0] for r in db.execute(
                    "SELECT id FROM records WHERE kind='baseline' AND text=? ORDER BY id", (key,))]
                choices.append(ids)
                for record_id in ids:
                    owners[record_id].add((source["id"], index))
            prepared.append((source, patterns, choices))

    groups, unmatched = [], []
    for source, patterns, choices in prepared:
        accepted, reasons = [], []
        for index, ids in enumerate(choices):
            reason = ""
            if not ids:
                reason = "本段在基准中没有对应中文"
            elif len(ids) > 1:
                reason = f"本段对应{len(ids)}个基准位置，归属不唯一"
            elif len(owners[ids[0]]) > 1:
                reason = "同一基准行有多个源分段候选，来源冲突"
            if reason:
                reasons.append(f"第{index + 1}段：{reason}")
                unmatched.append([f"T{source['id']}", index + 1, patterns[0]["parts"][index], patterns[1]["parts"][index],
                                  reason, source["path"], source["sheet"], source["row"], len(ids)])
            else:
                accepted.append((index, get_record(db, ids[0])))
                db.execute("INSERT OR IGNORE INTO covered VALUES (?)", (ids[0],))
        db.execute("INSERT INTO trial_coverage VALUES (?,?,?)", (source["id"], len(accepted), len(choices)))
        if accepted:
            groups.append((source, accepted, patterns))
            if len(accepted) == len(choices):
                db.execute("INSERT OR IGNORE INTO covered VALUES (?)", (source["id"],))
        if reasons:
            pending.append([source["id"], source["original"], source["korean"],
                            f"已匹配{len(accepted)}/{len(choices)}段；" + "；".join(reasons),
                            source["path"], source["sheet"], source["row"], len(mappings[source["id"]])])

    def details():
        for source, accepted, (cn, ko) in groups:
            evidence = context_evidence(db, [source], [b for _, b in accepted])
            for index, baseline in accepted:
                yield [f"T{source['id']}", cn["labels"][index] if cn["labels"] else str(index + 1),
                       trim_segment(cn["parts"][index]), trim_segment(ko["parts"][index]) if cn["labels"] else ko["parts"][index],
                       baseline["original"], "试拆分，待人工复核语义",
                       baseline["path"], baseline["sheet"], baseline["row"],
                       source["path"], source["sheet"], source["row"],
                       cn["parts"][index], ko["parts"][index],
                       f"{cn['spans'][index][0] + 1}～{cn['spans'][index][1]}",
                       f"{ko['spans'][index][0] + 1}～{ko['spans'][index][1]}", evidence,
                       cn["rule"], f"{len(accepted)}/{len(cn['parts'])}段"]

    book = Workbook(write_only=True)
    segments = write_detail_sheet(book, "拆分明细",
        ["组ID", "选项编号", "拆分中文", "拆分韩文", "基准中文", "状态",
         "基准完整路径", "基准Sheet", "基准行号", "源完整路径", "源Sheet", "源行号",
         "中文分段原值", "韩文分段原值", "中文字符位置", "韩文字符位置", "邻句证据", "拆分规则", "整组匹配情况"],
        details(), [16, 12, 45, 55, 45, 30, 50, 20, 12, 50, 20, 12, 50, 60, 20, 20, 65, 24, 18], progress)
    write_detail_sheet(book, "整组原文",
        ["组ID", "源中文原值", "源韩文原值", "段数", "源完整路径", "源Sheet", "源行号"],
        ([f"T{s['id']}", s["original"], s["korean"], len(p[0]["parts"]), s["path"], s["sheet"], s["row"]]
         for s, _, p in groups), [16, 65, 75, 10, 50, 20, 12])
    write_detail_sheet(book, "分段未匹配",
        ["组ID", "段序号", "分段中文", "分段韩文", "原因", "源完整路径", "源Sheet", "源行号", "基准候选数"],
        unmatched, [16, 12, 60, 60, 80, 55, 22, 14, 18])
    write_detail_sheet(book, "分段候选",
        ["组ID", "段序号", "分段中文", "基准完整路径", "基准Sheet", "基准行号", "基准中文"],
        ([f"T{s['id']}", i + 1, p[0]["parts"][i], b["path"], b["sheet"], b["row"], b["original"]]
         for s, p, choices in prepared for i, ids in enumerate(choices) for record_id in ids
         for b in [get_record(db, record_id)]), [16, 12, 60, 55, 22, 14, 60])
    write_detail_sheet(book, "暂不拆分",
        ["源记录ID", "源中文", "源韩文", "原因", "源完整路径", "源Sheet", "源行号", "基准位置组数"],
        pending, [16, 65, 75, 75, 50, 20, 12, 18])
    write_detail_sheet(book, "拆分说明", ["项目", "值"], [
        ["结果", f"试拆分 {len(groups)} 组 / {segments} 条；暂不拆分 {len(pending)} 个源行"],
        ["本版规则", "优先数字编号；中文无编号时按实际或字面换行拆分，中韩段数须一致。中文逐段查找，可部分匹配，保留原段序号"],
        ["位置要求", "优先采用唯一的整组连续对应；其他段只有唯一基准位置且无来源冲突才出预览；所有对应仅供复核"],
        ["文本处理", "编号模式移除编号及分段首尾空白；换行模式保留每段韩语原值和编号。两侧仅去除空白边缘行，内部空行不跳过，不按中文长度切韩语"],
        ["拒绝条件", "前缀、重复或跳号、空段、选项内部换行、中韩占位符差异、跨段或未闭合展示标签均暂不拆分"],
        ["待复核", "本表仅为试拆分预览，编号一致不证明译文语义正确；不回填基准TextTrans，不改动输入文件"],
        ["部分匹配", "拆分明细只列可定位的段；其余见分段未匹配，全部对应见分段候选；暂不拆分包含部分匹配组的剩余问题"],
        ["覆盖范围", "沿用同目录结构差异扫描.xlsx的扫描设置及跳过清单；暂不拆分只列格式或拆分候选，不代表所有这些行都需要拆分"],
    ], [24, 110])
    book.save(destination)
    book.close()
    return len(groups), segments, len(pending)


def write_report(db, destination, skipped, report, baseline_root, max_rows, progress, assessments):
    book = Workbook(write_only=True)
    stats = defaultdict(lambda: dict(count=0, files=set(), matched=0, issues=0, examples=[]))
    # Candidate sheets are evidence only: no translated output cells are changed.
    headers = ["候选ID", "规则", "源完整路径", "源Sheet", "源实际行号", "源中文（含记录ID）", "源韩文（含记录ID）",
               "基准完整路径", "基准Sheet", "基准实际行号", "基准中文（含记录ID）", "邻句证据", "判断与复核原因", "韩文分段候选",
               "源角色与剧情信息", "基准角色与条目信息"]
    widths = [12, 22, 50, 22, 22, 65, 65, 50, 22, 22, 65, 65, 75, 75, 70, 70]
    counts = [write_detail_sheet(book, f"{kind}候选", headers, candidate_rows(db, kind, assessments), widths, progress) for kind in ("合并", "拆分")]
    samples = write_detail_sheet(book, "规则样本",
        ["样本ID", "源记录ID", "源完整路径", "源Sheet", "源行号", "语言", "候选规则", "源中文", "源韩文",
         "分段与原文字符位置", "段数", "中韩验证", "另一语言分段", "基准候选数", "问题"],
        pattern_samples(db, stats), [24, 14, 50, 22, 12, 10, 22, 65, 65, 80, 10, 65, 80, 16, 75], progress)
    write_detail_sheet(book, "规则汇总", ["语言", "候选规则", "样本数", "涉及文件数", "有基准候选的样本数", "格式异常样本数", "示例样本ID", "说明"],
        ([language, rule, s["count"], len(s["files"]), s["matched"] if language == "中文" else None,
          s["issues"], "；".join(s["examples"]),
          "全部样本见规则样本；基准候选数仅统计中文；规则尚待确认"] for (language, rule), s in sorted(stats.items())),
        [12, 24, 14, 16, 24, 22, 65, 65])
    unresolved = write_detail_sheet(book, "未解释", ["记录ID", "文件类型", "完整路径", "Sheet", "行号", "原中文", "原韩文", "原因"],
                                    unresolved_rows(db, max_rows), [12, 12, 50, 22, 12, 65, 65, 80], progress)
    write_detail_sheet(book, "原始内容", ["记录ID", "文件类型", "完整路径", "Sheet", "行号", "场景", "原中文", "原韩文", "角色与剧情信息"],
        ([r["id"], "台本" if r["kind"] == "source" else "基准", r["path"], r["sheet"], r["row"], r["scene"], r["original"], r["korean"], metadata_text([r])]
         for r in db.execute("SELECT * FROM records ORDER BY id")), [12, 12, 50, 22, 12, 24, 80, 80, 70], progress)
    write_detail_sheet(book, "跳过清单", ["文件名", "完整路径", "Sheet", "原因或范围"], skipped, [30, 60, 24, 90])
    write_detail_sheet(book, "扫描说明", ["项目", "值"], [
        ["源文件夹", str(report.root)], ["基准文件夹", str(baseline_root)],
        ["扫描时间", datetime.now().isoformat(timespec="seconds")], ["最大连续行数", max_rows],
        ["最大拼接字符数", MAX_CHARACTERS],
        ["比较规则", "去首尾空白，忽略<color>等英文尖括号标签、实际与字面换行；保留占位符、内部空格及标点；原文与译文标签保留供复核"],
        ["候选含义", "仅发现当前数据中的可能对应关系；不自动分配、合并或回填译文；不执行单侧增删扫描"],
        ["合并反查顺序", "先对全部基准检查整句命中；只有char.xlsx（文件名不区分大小写）中无中文候选的行参与合并反查，已有候选但缺译或冲突的行保留常规处理"],
        ["合并范围", "story.xlsx及其他基准不参与合并反查；继续整句匹配及拆分扫描。范围外行不能标记为反查未命中，分文件统计见试合并结果.xlsx的合并范围表"],
        ["试合并预览", "同目录试合并结果.xlsx提供换行连接的韩文预览、全部候选、整句命中检查和流程统计"],
        ["试拆分预览", "同目录试拆分结果.xlsx按数字编号或换行生成独立中韩分段预览，支持部分定位，具体条件见拆分说明；候选仍需人工复核"],
        ["窗口边界", "同文件同Sheet连续物理行；不跨空白、非文本、排除行或显式场景变化。没有场景列时无法保证不跨语义场景"],
        ["场景识别", "可识别场景、场景名、场景编号中的唯一列，非空场景值沿用到下次变化；多个场景列需先澄清"],
        ["搜索范围", "整组候选限双方连续2至最大行数；试拆分还逐段查找基准并报告未定位原因；不证明未解释内容不存在对应"],
        ["中韩验证", "编号、顺序和数量仅提供结构证据，无法证明语义对应。没有韩文边界时不按中文长度切割"],
        ["原文回查", "候选展示过长会明确提示截断；完整文本见原始内容。字符位置从1起、含首尾，按未清理原文计数"],
        ["候选计数", "不同来源、不同窗口或不同拆法分别计数，不代表可以自动同步的唯一行数；所有候选均待复核"],
        ["包含子文件夹", str(report.recursive)], ["包含隐藏源Sheet", str(report.include_hidden_sheets)],
        ["包含隐藏源行", str(report.include_hidden_rows)], ["包含废弃源行", str(report.include_deprecated)],
        ["基准范围", "递归扫描全部xlsx/xlsm工作表Text列，包含隐藏行和工作表；读取失败列入跳过清单"],
    ], [28, 110])
    book.save(destination)
    book.close()
    return *counts, samples, unresolved


def scan_structure(report, baseline_root, output_root, ignored_sheets=(), progress=None, max_rows=8, review_report=None):
    if type(max_rows) is not int or not 2 <= max_rows <= 50:
        raise ValueError("最大连续行数必须是 2～50 的整数")
    progress = progress or (lambda text: None)
    fresh, candidates, skipped, baselines, baseline_root, output_root = prepare_inputs(
        report, baseline_root, output_root, ignored_sheets, progress)
    output_root.mkdir(parents=True, exist_ok=True)
    review_path = output_root / "合并复核记录.json"
    reviews = load_reviews(review_path)
    with tempfile.TemporaryDirectory(prefix="结构扫描临时_", dir=output_root) as temporary:
        staging = Path(temporary)
        db = sqlite3.connect(staging / "index.sqlite3")
        db.row_factory = sqlite3.Row
        try:
            create_index(db)
            load_records(db, fresh, candidates, baselines, skipped, progress)
            index_whole_matches(db, progress)
            index_merge_targets(db)
            find_links(db, max_rows, progress)
            assessments, reviews = merge_assessments(db, reviews, review_report)
            progress("生成编号及换行试拆分预览")
            trial_counts = write_trial_report(db, staging / "试拆分结果.xlsx", progress, max_rows)
            counts = write_report(db, staging / "结构差异扫描.xlsx", skipped, fresh, baseline_root, max_rows, progress, assessments)
            progress("生成 char.xlsx 未匹配行的反查试合并预览")
            merge_counts = write_merge_trial_report(db, staging / "试合并结果.xlsx", max_rows, progress, assessments)
            source_rows, baseline_rows = [db.execute("SELECT COUNT(*) FROM records WHERE kind=?", (k,)).fetchone()[0]
                                          for k in ("source", "baseline")]
        finally:
            db.close()
        (staging / "index.sqlite3").unlink()
        destination = output_root / ("结构扫描_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + staging.name[-6:])
        os.rename(staging, destination)
    try:
        if review_report:
            write_json({"version": 1, "reviews": reviews}, review_path, prefix="reviews_")
    except Exception:
        shutil.rmtree(destination)
        raise
    return StructureReport(destination, fresh, source_rows, baseline_rows, *counts, len(skipped), *trial_counts, *merge_counts)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="扫描连续合并候选及中韩拆分格式，不修改输入文件")
    parser.add_argument("source")
    parser.add_argument("baseline")
    parser.add_argument("output")
    parser.add_argument("--max-rows", type=int, default=8)
    parser.add_argument("--header-rows", type=int, default=10)
    parser.add_argument("--review-report", help="含复核填写工作表的试合并结果.xlsx；结论自动保存到输出根目录")
    args = parser.parse_args()
    ignored = load_ignored()
    report = scan_folder(args.source, header_rows=args.header_rows, ignored_sheets=ignored)
    result = scan_structure(report, args.baseline, args.output, ignored, print, args.max_rows, args.review_report)
    print(f"扫描完成：合并 {result.merge_candidates}，拆分 {result.split_candidates}，格式样本 {result.pattern_samples}")
    print(result.destination / "结构差异扫描.xlsx")
    print(f"试拆分：{result.trial_groups} 组 / {result.trial_segments} 条；暂不拆分 {result.trial_pending} 个源行")
    print(result.destination / "试拆分结果.xlsx")
    print(f"整句命中 {result.whole_matched_rows} 行；试合并 {result.merge_trial_rows} 行；待确认 {result.merge_pending_rows} 行；不适用 {result.merge_rejected_rows} 行")
    print(result.destination / "试合并结果.xlsx")
