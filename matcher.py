"""Match approved source sheets using a disk index; save baseline copies only."""

import os
import sqlite3
import tempfile
import shutil
import re
import json
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from io import TextIOWrapper
from zipfile import ZipFile
from xml.sax.expatreader import create_parser
from xml.sax.handler import feature_external_ges
from xml.sax.saxutils import XMLGenerator, escape
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ignored_sheets import sheet_key
from scanner import SUPPORTED, hidden_row_numbers, normalize, scan_folder
from updater import VERSION


def match_key(value):
    return value.strip() if isinstance(value, str) else ""


PLAIN_REVIEW_NOTE = "仅纯文本匹配，仍需复核标签"
PUNCT_REVIEW_NOTE = "比较时忽略指定标点差异（…/...、，/,、；/;、：/:）；保留旧候选共同判断"
PLAIN_PUNCTUATION = str.maketrans({"…": "...", "，": ",", "；": ";", "：": ":"})
PUNCT_MATCH = "标点等价补匹配"
TAIL_MATCH = "句尾去括号补匹配"
FUZZY_MATCH = "模糊匹配"
FUZZY_MIN_PERCENT = 80
CONTEXT_SKIP_NOTES = {"演出", "气息"}
TEXT_TAG = re.compile(r"</?[A-Za-z][^<>\r\n]*>")
CODE_TAGS = re.compile(r"\{[^{}\r\n]*\}|\\[nr]|[\r\n]")
RECOMMENDED_PAIRS = """
    SELECT DISTINCT chosen.text,chosen.korean_key
    FROM recommendations r JOIN source_rows chosen ON chosen.rowid=r.source_id
    UNION SELECT chosen.text,chosen.korean_key FROM source_rows chosen JOIN (
        SELECT source_id FROM structure_sources GROUP BY source_id
        HAVING MAX(complete)=1 OR COUNT(DISTINCT part)=MAX(total)
    ) covered ON covered.source_id=chosen.rowid
"""


def plain_key(value, *, punctuation=True):
    text = CODE_TAGS.sub("", TEXT_TAG.sub("", match_key(value))).strip()
    return text.translate(PLAIN_PUNCTUATION) if punctuation else text


def source_metadata_columns(header):
    columns = []
    for name in ("特殊说明", "说话角色名"):
        found = [c.column for c in header if normalize(c.value) == name]
        if len(found) > 1:
            raise ValueError(f"存在多个{name}列，无法判断上下文清理规则")
        columns.append(found[0] if found else None)
    remarks = tuple(c.column for c in header if normalize(c.value) in {"备注", "비고"})
    return *columns, remarks


def source_row_metadata(note, speaker, include_deprecated, remarks=()):
    """Share source eligibility between matching and structural discovery."""
    note_text = normalize(str(note.value)) if note is not None and note.value is not None else ""
    speaker_text = normalize(str(speaker.value)) if speaker is not None and speaker.value is not None else ""
    markers = ([note] if note is not None else []) + list(remarks)
    if not include_deprecated:
        for cell in markers:
            if cell.data_type != "f" and normalize(cell.value) in {"废弃", "폐기"}:
                return note_text, speaker_text, f"废弃：{cell.coordinate}={normalize(cell.value)}（未勾选包含废弃行）"
    # Unknown note formulas affect this row only. Speaker metadata is optional;
    # missing names/notes do not make otherwise valid Chinese/Korean invalid.
    for cell in markers:
        if cell.data_type == "f" and (cell is note or not include_deprecated):
            return note_text, speaker_text, f"说明字段 {cell.coordinate} 含公式，无法判断标记；仅跳过本行，请先转成文本值"
    return note_text, speaker_text, ""


def breath_key(value):
    """Remove balanced Chinese/ASCII parenthetical directions only for breath lines."""
    text = match_key(value)
    stack, result = [], []
    pairs = {"(": ")", "（": "）"}
    for char in text:
        if char in pairs:
            stack.append(pairs[char])
        elif char in (")", "）"):
            if not stack or stack.pop() != char:
                return text  # Incomplete or mixed-up brackets need manual review.
        elif not stack:
            result.append(char)
    return text if stack else "".join(result).strip()


def trailing_note_key(value):
    """Remove only balanced parenthetical groups at the end; keep malformed text."""
    text = match_key(value)
    end = len(text)
    pairs = {")": "(", "）": "（"}
    stack = []
    for char in text:
        if char in ("(", "（"):
            stack.append(char)
        elif char in pairs and (not stack or stack.pop() != pairs[char]):
            return text
    if stack:
        return text
    while end and text[end - 1] in pairs:
        stack, start = [], None
        for index in range(end - 1, -1, -1):
            char = text[index]
            if char in pairs:
                stack.append(pairs[char])
            elif char in ("(", "（"):
                if not stack or stack.pop() != char:
                    return text
                if not stack:
                    start = index
                    break
        if start is None:
            return text
        end = len(text[:start].rstrip())
    return text[:end]


def boundary_newlines(value):
    """Detect CR/LF in the whitespace before and after the baseline text."""
    return tuple(any(c in edge for c in "\r\n") for edge in (
        re.match(r"^\s*", value).group(), re.search(r"\s*$", value).group()))


def align_boundary_newlines(value, keep_start, keep_end):
    """Align boundary newline presence; retain spaces and internal newlines."""
    if not isinstance(value, str):
        return value
    without_newlines = lambda match: match.group().replace("\r", "").replace("\n", "")
    if not keep_start:
        value = re.sub(r"^\s+", without_newlines, value, count=1)
    elif value.strip() and not boundary_newlines(value)[0]:
        value = "\n" + value
    if not keep_end:
        value = re.sub(r"\s+$", without_newlines, value, count=1)
    elif value.strip() and not boundary_newlines(value)[1]:
        value += "\n"
    return value


def text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # A source string beginning with '=' is text, never a formula.
    return cell


@dataclass
class MatchReport:
    destination: Path
    check_report: object
    baseline_rows: int
    matched_rows: int
    source_rows: int
    source_matched: int
    skipped: int
    outputs: list
    missing_translations: int = 0
    conflict_texts: int = 0
    match_pairs: int = 0
    recommended_rows: int = 0
    baseline_files: int = 0


def baseline_sheets(workbook, path, header_rows, *, for_output=True):
    for sheet in workbook:
        if hasattr(sheet, "reset_dimensions"):
            sheet.reset_dimensions()
        candidates = []
        for number, row in enumerate(sheet.iter_rows(max_row=header_rows), 1):
            headers = {cell.column: normalize(cell.value) for cell in row if cell.value is not None}
            for column, title in headers.items():
                if title == "Text":
                    candidates.append((number, column, headers))
        location = f"{path.name} / {sheet.title}"
        if not candidates:
            raise ValueError(f"基准表头缺失：{location} 的第 1～{header_rows} 行未找到 Text（完整名称匹配，忽略首尾空白）")
        if len(candidates) > 1:
            positions = "、".join(f"{get_column_letter(column)}{row}" for row, column, _ in candidates)
            raise ValueError(f"基准表头不明确：{location} 找到多个 Text：{positions}，请保留唯一的 Text 表头")
        header_row, text_column, headers = candidates[0]
        if for_output and ("匹配中文1" in headers.values() or "推荐状态" in headers.values()):
            raise ValueError(f"{location} 已有检索结果，请提供未回填的基准文件")
        output_column = max(18, max(headers) + 1)
        if for_output and output_column + 15 > 16384:
            raise ValueError(f"{location} 的表头右侧不足 16 列，无法写入检索结果")
        yield sheet, header_row, text_column, output_column


class BaselineWriter(XMLGenerator):
    """Stream original worksheet XML unchanged in meaning, adding result cells."""

    def __init__(self, stream, connection, path, sheet, header_row, output_column):
        # Avoid the platform's automatic LF -> CRLF conversion in worksheet XML.
        stream = TextIOWrapper(stream, encoding="utf-8", newline="\n", write_through=True)
        super().__init__(stream, encoding="utf-8", short_empty_elements=True)
        self.connection, self.path, self.sheet = connection, str(path), sheet
        self.header_row = header_row
        self.column_shift = output_column - 18
        self.output_letters = {get_column_letter(c) for c in range(output_column, output_column + 16)}
        self.row = None
        self.protected_cell = False
        self.protected_value = False
        self.prefix = ""
        self.omit_dimension = False
        self.has_columns = False
        self.matched_rows = 0

    def startElement(self, name, attrs):
        local = name.rsplit(":", 1)[-1]
        if local == "worksheet":
            self.prefix = name[:-len(local)]
        if local == "dimension":
            # Optional metadata can be wrong in the customer's baseline files.
            # Without it readers determine dimensions from the actual cells.
            self.omit_dimension = True
            return
        if local == "cols":
            self.has_columns = True
        if local == "sheetData" and not self.has_columns:
            super().startElement(self.prefix + "cols", {})
            self.output_columns()
            super().endElement(self.prefix + "cols")
        if local == "row":
            self.row = int(attrs["r"])
        if local == "c":
            column = re.match(r"[A-Z]+", attrs.get("r", "")).group()
            self.protected_cell = column in self.output_letters
        if self.protected_cell and local == "f":
            self.existing_content()
        if self.protected_cell and local in ("v", "t"):
            self.protected_value = True
        if self.protected_cell:
            return
        super().startElement(name, attrs)

    def characters(self, content):
        if self.protected_cell and self.protected_value and content:
            self.existing_content()
        if not self.omit_dimension and not self.protected_cell:
            self.xml_characters(content)

    def xml_characters(self, content):
        # XML parsers normalize literal CR/CRLF; character references preserve
        # retained source newlines and the original baseline cell values.
        if "\r" in content:
            self._finish_pending_start_element()
            self._write(escape(content, {"\r": "&#13;"}))
        else:
            super().characters(content)

    def existing_content(self):
        first, last = (get_column_letter(c + self.column_shift) for c in (18, 33))
        raise ValueError(f"{Path(self.path).name} / {self.sheet} 的 {first}～{last} 列已有内容，请提供未回填的基准文件")

    def output_columns(self):
        for first, last, width in ((18, 26, 45), (27, 29, 18), (30, 32, 55), (33, 33, 24)):
            super().startElement(self.prefix + "col", {"min": str(first + self.column_shift), "max": str(last + self.column_shift), "width": str(width), "customWidth": "1"})
            super().endElement(self.prefix + "col")

    def output_cell(self, column, value):
        if value is None:
            return
        attrs = {"r": f"{get_column_letter(column + self.column_shift)}{self.row}"}
        if isinstance(value, str):
            attrs["t"] = "inlineStr"
        super().startElement(self.prefix + "c", attrs)
        if isinstance(value, str):
            super().startElement(self.prefix + "is", {})
            super().startElement(self.prefix + "t", {"xml:space": "preserve"})
            self.xml_characters(value)
            super().endElement(self.prefix + "t")
            super().endElement(self.prefix + "is")
        else:
            super().startElement(self.prefix + "v", {})
            super().characters(str(value))
            super().endElement(self.prefix + "v")
        super().endElement(self.prefix + "c")

    def endElement(self, name):
        local = name.rsplit(":", 1)[-1]
        if local == "dimension":
            self.omit_dimension = False
            return
        if local in ("v", "t"):
            self.protected_value = False
        if local == "c":
            omitted = self.protected_cell
            self.protected_cell = False
            if omitted:
                return
        if self.protected_cell:
            return
        if local == "cols":
            self.output_columns()
        if local == "row":
            if self.row == self.header_row:
                for column in range(18, 27):
                    title = ("匹配中文", "匹配韩语", "来源路径 / Sheet / 行号")[(column - 18) % 3]
                    self.output_cell(column, f"{title}{(column - 18) // 3 + 1}")
                self.output_cell(27, "匹配条数")
                self.output_cell(28, "缺译条数")
                self.output_cell(29, "译文冲突")
                for column, title in enumerate(("推荐韩语", "推荐来源", "推荐依据", "推荐状态"), 30):
                    self.output_cell(column, title)
            else:
                record = self.connection.execute("SELECT text,leading_newline,trailing_newline FROM baseline_rows WHERE path=? AND sheet=? AND row=?",
                                              (self.path, self.sheet, self.row)).fetchone()
                if record:
                    key = (self.path, self.sheet, self.row)
                    keep_start, keep_end = record[1:]
                    summary = self.connection.execute("SELECT n,missing FROM baseline_counts WHERE path=? AND sheet=? AND row=?", key).fetchone()
                    count, missing = summary if summary else (0, 0)
                    self.matched_rows += bool(count)
                    if 0 < count <= 3:
                        hits = self.connection.execute("""SELECT s.chinese,s.korean,
                            s.path || '；Sheet：' || s.sheet || '；行号：' || s.row AS source,s.rowid AS ordering
                            FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
                            WHERE e.path=? AND e.sheet=? AND e.row=?
                            UNION ALL SELECT chinese,korean,source,9223372036854775807 FROM structure_recommendations
                            WHERE path=? AND sheet=? AND row=? ORDER BY ordering LIMIT 3""", (*key, *key))
                        for index, (chinese, korean, source, _) in enumerate(hits):
                            chinese = align_boundary_newlines(chinese, keep_start, keep_end)
                            korean = align_boundary_newlines(korean, keep_start, keep_end)
                            for offset, value in enumerate((chinese, korean, source)):
                                self.output_cell(18 + index * 3 + offset, value)
                    self.output_cell(27, count)
                    self.output_cell(28, missing)
                    conflict = self.connection.execute("SELECT 1 FROM baseline_counts WHERE path=? AND sheet=? AND row=? AND variants>1", key).fetchone()
                    self.output_cell(29, "是" if conflict else "否" if count else "未匹配")
                    recommendation = self.connection.execute("""
                        SELECT s.korean,s.path,s.sheet,s.row,r.basis,r.status
                        FROM recommendations r LEFT JOIN source_rows s ON s.rowid=r.source_id
                        WHERE r.path=? AND r.sheet=? AND r.row=?
                    """, (self.path, self.sheet, self.row)).fetchone()
                    korean, filename, title, number, basis, status = recommendation
                    source = f"{filename}；Sheet：{title}；行号：{number}" if filename is not None else None
                    structural = self.connection.execute("""SELECT korean,source,basis,status
                        FROM structure_recommendations WHERE path=? AND sheet=? AND row=?""", key).fetchone()
                    if structural:
                        korean, source, basis, status = structural
                    self.output_cell(30, align_boundary_newlines(korean, keep_start, keep_end))
                    self.output_cell(31, source)
                    self.output_cell(32, basis)
                    self.output_cell(33, status)
            self.row = None
        super().endElement(name)


def write_baseline_copy(path, destination, connection, sheets):
    matched = 0
    with ZipFile(path) as source, ZipFile(destination, "w") as output:
        for item in source.infolist():
            with source.open(item) as original, output.open(item, "w") as copied:
                if item.filename in sheets:
                    writer = BaselineWriter(copied, connection, path, *sheets[item.filename])
                    parser = create_parser()
                    parser.setFeature(feature_external_ges, False)
                    parser.setContentHandler(writer)
                    parser.parse(original)
                    matched += writer.matched_rows
                else:
                    shutil.copyfileobj(original, copied, length=1024 * 1024)
    return matched


def write_detail_sheet(book, title, headers, rows, widths, progress=None, row_limit=1048576):
    """Stream data rows; split automatically at Excel's worksheet row limit."""
    sheet = None
    page = count = total = 0
    for values in rows:
        if sheet is None or count == row_limit:
            if sheet is not None:
                sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
            page += 1
            sheet = book.create_sheet(title if page == 1 else f"{title}_{page}")
            sheet.freeze_panes = "A2"
            for column, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(column)].width = width
            header_cells = []
            for value in headers:
                cell = WriteOnlyCell(sheet, value=value)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245B78")
                header_cells.append(cell)
            sheet.append(header_cells)
            count = 1
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[len(cells)].endswith("率"):
                cell.number_format = "0.00%"
            elif headers[len(cells)] == "相似度":
                cell.number_format = "0.0%"
            cells.append(cell)
        sheet.append(cells)
        count += 1
        total += 1
        if progress and total % 5000 == 0:
            progress(f"写出{title}：{total} 条")
    if sheet is None:
        sheet = book.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.append(headers)
        count = 1
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    return total


def read_context(connection, table, path, sheet, sequence, plain=False, limit=3):
    """Source windows extend until three usable lines, retaining skipped evidence."""
    column = "plain_text" if plain else "text"
    source = table == "source_context_rows"
    expression = (f"CASE WHEN context_skip='气息' THEN {'breath_plain' if plain else 'breath_text'} "
                  f"ELSE {column} END") if source else column
    nonempty = f"({column}!='' OR context_skip!='')" if source else f"{column}!=''"
    result = []
    for operator, order in (("<", "DESC"), (">", "ASC")):
        cursor = connection.execute(f"""
            SELECT row,{expression}{',context_skip' if source else ''} FROM {table}
            WHERE path=? AND sheet=? AND seq {operator} ?
            AND {nonempty}
            ORDER BY seq {order} {'' if source else 'LIMIT ?'}
        """, (path, sheet, sequence) if source else (path, sheet, sequence, limit))
        records, usable = [], 0
        for record in cursor:
            records.append(list(record))
            usable += not source or not record[2]
            if usable >= limit:
                break
        result.append(records)
    return tuple(result)


def align_context(connection, path, sheet, sequence, baseline, source, plain=False):
    """Consume an ignored line on the baseline side only when its text also matches."""
    limit = max(3, *(len(side) for side in source))
    if limit > 3:
        baseline = read_context(connection, "baseline_rows", path, sheet, sequence, plain, limit)
    used_baseline, used_source, skipped = [], [], []
    for direction, base, raw_source in zip(("前", "后"), baseline, source):
        base_index, bp, sp = 0, [], []
        for row, text, note in raw_source:
            if note:
                base_text = base[base_index][1] if base_index < len(base) else ""
                if note == "气息":
                    base_text = breath_key(base_text)
                if base_index < len(base) and text == base_text:
                    position = f"同时越过基准行{base[base_index][0]}"
                    base_index += 1
                else:
                    position = "基准邻句无同文，仅越过台本行"
                skipped.append(f"{direction}：源行{row}（{note}）；{position}")
                continue
            sp.append([row, text])
            if base_index < len(base):
                bp.append(base[base_index])
                base_index += 1
        bp.extend(base[base_index:base_index + 3 - len(bp)])
        used_baseline.append(bp)
        used_source.append(sp)
    return *used_baseline, *used_source, "\n".join(skipped)


def context_text(records, direction):
    return "\n".join(f"{direction}{index}句 [行{row}] {text}" for index, (row, text) in enumerate(records, 1))


def context_statistics(source, previous_equal, following_equal):
    counts = [(len(side), sum(r[2] == "演出" for r in side), sum(r[2] == "气息" for r in side), equal)
              for side, equal in zip(source, (previous_equal, following_equal))]
    totals = tuple(sum(values) for values in zip(*counts))
    return "\n".join(f"{label}：查询{n}行，演出{stage}行，气息{breath}行，匹配{equal}句"
                     for label, (n, stage, breath, equal) in
                     zip(("向上", "向下", "合计（不含当前句）"), [*counts, totals]))


def initialize_recommendations(connection):
    connection.executescript("""
        CREATE TABLE source_context (source_id INTEGER PRIMARY KEY, previous TEXT, following TEXT,plain_previous TEXT,plain_following TEXT);
        CREATE TABLE comparisons (path TEXT,sheet TEXT,row INTEGER,source_id INTEGER,
            previous_equal INTEGER,following_equal INTEGER,full_sides INTEGER,match_type TEXT,
            baseline_previous TEXT,baseline_following TEXT,source_previous TEXT,source_following TEXT,skipped TEXT,query_statistics TEXT,
            similarity REAL,
            PRIMARY KEY(path,sheet,row,source_id));
        CREATE INDEX comparison_source ON comparisons(source_id);
        CREATE TABLE recommendations (path TEXT,sheet TEXT,row INTEGER,source_id INTEGER,basis TEXT,status TEXT,match_type TEXT,
            PRIMARY KEY(path,sheet,row));
        CREATE INDEX recommendation_source ON recommendations(source_id);
        CREATE TABLE structure_recommendations (path TEXT,sheet TEXT,row INTEGER,chinese TEXT,korean TEXT,
            source TEXT,basis TEXT,status TEXT,kind TEXT,group_id TEXT,PRIMARY KEY(path,sheet,row));
        CREATE TABLE structure_sources (path TEXT,sheet TEXT,row INTEGER,source_id INTEGER,
            part INTEGER,total INTEGER,complete INTEGER,PRIMARY KEY(path,sheet,row,source_id));
        CREATE INDEX structure_source_id ON structure_sources(source_id);
        CREATE TABLE structure_pending (path TEXT,sheet TEXT,row INTEGER,chinese TEXT,korean TEXT,
            source TEXT,basis TEXT,status TEXT,kind TEXT,group_id TEXT,reason TEXT);
        CREATE TABLE structure_pending_sources (source_id INTEGER,reason TEXT);
        CREATE INDEX structure_pending_source ON structure_pending_sources(source_id);
    """)


def build_recommendations(connection, progress=None, *, supplement=False):
    """Run legacy rounds first; later rounds may only fill unresolved positions."""
    if not supplement:
        initialize_recommendations(connection)
    before = connection.execute("SELECT COUNT(*) FROM recommendations WHERE source_id IS NOT NULL").fetchone()[0]
    rounds = (PUNCT_MATCH, TAIL_MATCH) if supplement else ("全文精确匹配", "纯文本补匹配", TAIL_MATCH)
    for round_type in rounds:
        tail = round_type == TAIL_MATCH
        plain = round_type != "全文精确匹配"
        if tail:
            connection.execute("DROP TABLE IF EXISTS tail_sources")
            connection.execute("CREATE TEMP TABLE tail_sources (source_id INTEGER PRIMARY KEY,text TEXT)")
            for source_id, text in connection.execute("""SELECT rowid,plain_text FROM source_rows s
                WHERE NOT EXISTS (SELECT 1 FROM comparisons c WHERE c.source_id=s.rowid)"""):
                key = trailing_note_key(text)
                if key and key != text:
                    connection.execute("INSERT INTO tail_sources VALUES (?,?)", (source_id, key))
            connection.execute("CREATE INDEX tail_text ON tail_sources(text)")
        column = "plain_text" if plain else "text"
        condition = """WHERE b.plain_text!='' AND NOT EXISTS
            (SELECT 1 FROM recommendations r WHERE r.path=b.path AND r.sheet=b.sheet AND r.row=b.row AND r.source_id IS NOT NULL)""" if plain else ""
        if tail:
            condition = "WHERE b.plain_text!='' AND EXISTS (SELECT 1 FROM tail_sources t WHERE t.text=b.plain_text)"
        if supplement:
            condition += """ AND NOT EXISTS (SELECT 1 FROM recommendations r
                WHERE r.path=b.path AND r.sheet=b.sheet AND r.row=b.row AND r.source_id IS NOT NULL)
                AND NOT EXISTS (SELECT 1 FROM structure_recommendations t
                    WHERE t.path=b.path AND t.sheet=b.sheet AND t.row=b.row)
                AND NOT EXISTS (SELECT 1 FROM structure_pending t
                    WHERE t.path=b.path AND t.sheet=b.sheet AND t.row=b.row)"""
        if progress:
            progress(("v1.1.12 补匹配：" if supplement else "v1.1.9 优先匹配：") + round_type + "……")
        records = connection.execute(f"SELECT b.path,b.sheet,b.row,b.{column},b.seq FROM baseline_rows b {condition} ORDER BY b.path,b.sheet,b.row")
        for index, (path, sheet, row, text, sequence) in enumerate(records, 1):
            match_type = round_type
            previous_recommendation = connection.execute(
                "SELECT source_id,basis,status,match_type FROM recommendations WHERE path=? AND sheet=? AND row=?",
                (path, sheet, row)).fetchone() if tail else None
            previous, following = read_context(connection, "baseline_rows", path, sheet, sequence, plain)
            breath_text = breath_key(text)
            if plain:
                count = connection.execute("""SELECT COUNT(*),SUM(korean_key=''),
                    COUNT(DISTINCT CASE WHEN korean_key!='' THEN korean_key END)
                    FROM source_rows WHERE plain_text=? OR (breath_plain!='' AND breath_plain=?)""", (text, breath_text)).fetchone()
            else:
                count = connection.execute("SELECT n,missing,variants FROM counts WHERE text=?", (text,)).fetchone()
            best_score, best_id, best_key, tied = (-1, -1), None, None, False
            best_previous = best_following = 0
            column = "plain_text" if plain else "text"
            breath_column = "breath_plain" if plain else "breath_text"
            candidate_condition = f"{column}=? OR ({breath_column}!='' AND {breath_column}=?)"
            arguments = (text, breath_text)
            if tail:
                candidate_condition = "rowid IN (SELECT source_id FROM tail_sources WHERE text=?)"
                arguments = (text,)
            if supplement:
                # Re-score all prior candidates with the broader context rules;
                # never hide an existing conflict by considering new sources alone.
                candidate_condition = f"""({candidate_condition}) OR rowid IN
                    (SELECT source_id FROM comparisons WHERE path=? AND sheet=? AND row=?)"""
                arguments = (*arguments, path, sheet, row)
            candidates = connection.execute(f"""
                SELECT rowid,korean_key,path,sheet,seq,{column} FROM source_rows
                WHERE {candidate_condition} ORDER BY rowid
            """, arguments)
            for source_id, korean_key, source_path, source_sheet, source_sequence, source_text in candidates:
                candidate_type = "气息去括号匹配" if not plain and source_text != text else match_type
                if supplement:
                    prior = connection.execute("SELECT match_type FROM comparisons WHERE path=? AND sheet=? AND row=? AND source_id=?",
                                               (path, sheet, row, source_id)).fetchone()
                    if prior:
                        candidate_type = prior[0]
                context = connection.execute("SELECT previous,following FROM source_context WHERE source_id=?", (source_id,)).fetchone()
                if context is None:
                    sp, sf = read_context(connection, "source_context_rows", source_path, source_sheet, source_sequence)
                    connection.execute("INSERT INTO source_context (source_id,previous,following) VALUES (?,?,?)", (source_id, json.dumps(sp, ensure_ascii=False), json.dumps(sf, ensure_ascii=False)))
                else:
                    sp, sf = map(json.loads, context)
                if plain:
                    context = connection.execute("SELECT plain_previous,plain_following FROM source_context WHERE source_id=?", (source_id,)).fetchone()
                    if context[0] is None:
                        sp, sf = read_context(connection, "source_context_rows", source_path, source_sheet, source_sequence, plain=True)
                        connection.execute("UPDATE source_context SET plain_previous=?,plain_following=? WHERE source_id=?",
                                           (json.dumps(sp, ensure_ascii=False), json.dumps(sf, ensure_ascii=False), source_id))
                    else:
                        sp, sf = map(json.loads, context)
                bp, bf, used_sp, used_sf, skipped = align_context(
                    connection, path, sheet, sequence, (previous, following), (sp, sf), plain)
                pe = sum(a[1] == b[1] for a, b in zip(bp, used_sp))
                fe = sum(a[1] == b[1] for a, b in zip(bf, used_sf))
                full = int(len(bp) == len(used_sp) == pe == 3) + int(len(bf) == len(used_sf) == fe == 3)
                connection.execute("INSERT OR REPLACE INTO comparisons VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (path, sheet, row, source_id, pe, fe, full, candidate_type,
                     *(json.dumps(items, ensure_ascii=False) for items in (bp, bf, used_sp, used_sf)), skipped,
                     ("v1.1.12标点等价重评上下文；" if supplement else "") + context_statistics((sp, sf), pe, fe), None))
                if not korean_key:
                    continue
                score = (full, pe + fe)
                if score > best_score:
                    best_score, best_id, best_key, tied = score, source_id, korean_key, False
                    best_previous, best_following = pe, fe
                elif score == best_score and korean_key != best_key:
                    tied = True
            if tail and previous_recommendation and previous_recommendation[0] is not None:
                connection.execute("UPDATE recommendations SET basis=? WHERE path=? AND sheet=? AND row=?",
                    ("保留先前匹配阶段推荐：" + previous_recommendation[1] + "；新增句尾去括号来源仅供复核", path, sheet, row))
                continue
            if tail or supplement:
                # New candidates must compete with every already-found source, not hide conflicts.
                count = connection.execute("""SELECT COUNT(*),SUM(s.korean_key=''),
                    COUNT(DISTINCT NULLIF(s.korean_key,'')) FROM comparisons c
                    JOIN source_rows s ON s.rowid=c.source_id WHERE c.path=? AND c.sheet=? AND c.row=?""",
                    (path, sheet, row)).fetchone()
                best_score, best_id, best_key, tied = (-1, -1), None, None, False
                for source_id, korean_key, pe, fe, full, candidate_type in connection.execute("""
                    SELECT s.rowid,s.korean_key,c.previous_equal,c.following_equal,c.full_sides,c.match_type
                    FROM comparisons c JOIN source_rows s ON s.rowid=c.source_id
                    WHERE c.path=? AND c.sheet=? AND c.row=? ORDER BY s.rowid""", (path, sheet, row)):
                    if not korean_key:
                        continue
                    score = (full, pe + fe)
                    if score > best_score:
                        best_score, best_id, best_key, tied = score, source_id, korean_key, False
                        best_previous, best_following, match_type = pe, fe, candidate_type
                    elif score == best_score and korean_key != best_key:
                        tied = True
            if count is None or count[0] == 0:
                status, basis, best_id = "未匹配", "没有忽略标签后正文一致的源行" if plain else "没有中文全文匹配的源行", None
            elif count[2] == 0:
                status, basis, best_id = "全部缺译", "中文已匹配，但所有来源韩语均为空或全空白", None
            elif count[2] == 1:
                status, basis = "唯一译文", "全部非空韩语去除首尾空白后相同；选取上下文得分最高的来源，并列取首条"
            elif best_score[0] == 0:
                status, basis, best_id = "待确认", "存在多种译文，没有任一来源的完整前3句或完整后3句一致", None
            elif tied:
                status, basis, best_id = "待确认", "存在多种译文，最高上下文得分对应不同韩语，无法唯一推荐", None
            else:
                status = "上下文推荐"
                basis = f"前{best_previous}/3句一致，后{best_following}/3句一致；完整一致侧数{best_score[0]}，一致位置数{best_score[1]}"
            if count and count[0]:
                basis += f"；参与本轮匹配的候选{count[0]}条，非空韩语{count[2]}种（仅统计参与扫描且通过过滤的来源）"
            if best_id is not None:
                basis += "；完美匹配：上下各3句有效台词一致" if best_score[0] == 2 else "；非完美匹配：上下未同时各满足3句"
                breath = connection.execute(f"SELECT {column},{breath_column} FROM source_rows WHERE rowid=?", (best_id,)).fetchone()
                if breath[1] and breath[1] == breath_text and breath[0] != text:
                    basis += "；气息台词忽略括号说明（原文保留）"
                    if not plain:
                        match_type = "气息去括号匹配"
            if tail:
                basis = "句尾去括号补查后综合候选；" + basis
                if best_id is not None and match_type == TAIL_MATCH:
                    basis = "忽略源中文句尾括号说明后匹配，仍需复核；" + PLAIN_REVIEW_NOTE + "；" + basis
            elif plain:
                basis = (PLAIN_REVIEW_NOTE + "；" if best_id is not None else "纯文本补匹配未能推荐；") + basis
            basis = (("v1.1.12规则补匹配；" + PUNCT_REVIEW_NOTE + "；") if supplement else
                     "v1.1.9规则优先；") + basis
            connection.execute("INSERT OR REPLACE INTO recommendations VALUES (?,?,?,?,?,?,?)", (path, sheet, row, best_id, basis, status, match_type))
            if progress and index % 5000 == 0:
                progress(f"比较上下文并推荐译文：{index} 个基准行")
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS baseline_counts
            (path TEXT,sheet TEXT,row INTEGER,n INTEGER,missing INTEGER,variants INTEGER);
        CREATE UNIQUE INDEX IF NOT EXISTS baseline_count_position ON baseline_counts(path,sheet,row);
        INSERT OR REPLACE INTO baseline_counts
            SELECT e.path,e.sheet,e.row,COUNT(*),SUM(s.korean_key=''),
                COUNT(DISTINCT CASE WHEN s.korean_key!='' THEN s.korean_key END) AS variants
            FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
            WHERE NOT EXISTS (SELECT 1 FROM structure_recommendations t
                WHERE t.path=e.path AND t.sheet=e.sheet AND t.row=e.row)
            GROUP BY e.path,e.sheet,e.row;
        CREATE TABLE IF NOT EXISTS matched_sources (source_id INTEGER PRIMARY KEY);
        INSERT OR IGNORE INTO matched_sources SELECT DISTINCT source_id FROM comparisons;
    """)
    recommended = connection.execute("SELECT COUNT(*) FROM recommendations WHERE source_id IS NOT NULL").fetchone()[0]
    connection.commit()
    return recommended - before


def apply_punctuation_recommendations(connection, progress=None):
    """Broaden disk keys only after legacy whole/structure results are locked."""
    connection.create_function("punctuation_key", 1, lambda text: text.translate(PLAIN_PUNCTUATION), deterministic=True)
    connection.execute("UPDATE baseline_rows SET plain_text=punctuation_key(plain_text)")
    for table in ("source_rows", "source_context_rows"):
        connection.execute(f"UPDATE {table} SET plain_text=punctuation_key(plain_text),breath_plain=punctuation_key(breath_plain)")
    connection.execute("UPDATE source_context SET plain_previous=NULL,plain_following=NULL")
    return build_recommendations(connection, progress, supplement=True)


def apply_structure_recommendations(connection, candidates):
    """Supplement unresolved targets and retain physical sources for every segment."""
    groups = {}
    for item in candidates:
        key = (item["path"], item["sheet"], item["row"])
        groups.setdefault(key, []).append(item)
    for key, items in groups.items():
        baseline = connection.execute("SELECT text FROM baseline_rows WHERE path=? AND sheet=? AND row=?", key).fetchone()
        chosen = connection.execute("SELECT source_id FROM recommendations WHERE path=? AND sheet=? AND row=?", key).fetchone()
        for item in items:
            if baseline is None or baseline[0] != match_key(item["chinese"]):
                raise ValueError("结构扫描后的基准内容已变化，请重新扫描")
            source_ids = []
            for source in item["sources"]:
                record = connection.execute("""SELECT rowid,chinese,korean FROM source_rows
                    WHERE path=? AND sheet=? AND row=?""", (source["path"], source["sheet"], source["row"])).fetchone()
                if record is None or record[1] != source["original"] or (record[2] or "") != source["korean"]:
                    raise ValueError("结构扫描后的源内容或参与范围已变化，请重新扫描")
                source_ids.append(record[0])
            if not source_ids or not match_key(item["korean"]) or item["kind"] not in ("合并", "拆分"):
                raise ValueError("结构推荐缺少来源、译文或有效类型")
            values = (*key, item["chinese"], item["korean"], item["source"], "v1.1.9规则优先；" + item["basis"],
                      item["kind"] + "推荐·待复核", item["kind"], item["group_id"])
            reason = "保留已有整句推荐" if chosen[0] is not None else "同一基准有多个结构推荐，需确认采用哪一组" if len(items) > 1 else ""
            if reason:
                connection.execute("INSERT INTO structure_pending VALUES (?,?,?,?,?,?,?,?,?,?,?)", (*values, reason))
                connection.executemany("INSERT INTO structure_pending_sources VALUES (?,?)", [(sid, reason) for sid in source_ids])
                continue
            connection.execute("INSERT INTO structure_recommendations VALUES (?,?,?,?,?,?,?,?,?,?)", values)
            connection.executemany("INSERT INTO structure_sources VALUES (?,?,?,?,?,?,?)",
                [(*key, sid, item.get("part", 0), item.get("total", 1), int(item["kind"] == "合并")) for sid in source_ids])
            connection.executemany("INSERT OR IGNORE INTO matched_sources VALUES (?)", [(sid,) for sid in source_ids])
            variants = connection.execute("""SELECT COUNT(*) FROM (
                SELECT s.korean_key FROM comparisons c JOIN source_rows s ON s.rowid=c.source_id
                WHERE c.path=? AND c.sheet=? AND c.row=? AND s.korean_key!='' UNION SELECT ?
            )""", (*key, match_key(item["korean"]))).fetchone()[0]
            connection.execute("""INSERT INTO baseline_counts VALUES (?,?,?,1,0,?)
                ON CONFLICT(path,sheet,row) DO UPDATE SET n=n+1,variants=excluded.variants""", (*key, variants))
    connection.commit()
    return connection.execute("SELECT COUNT(*) FROM structure_recommendations").fetchone()[0]


def apply_fuzzy_recommendations(connection, progress=None):
    """Last-pass suggestions for unmatched sources and targets; index text on disk."""
    connection.executescript(f"""
        CREATE TEMP TABLE fuzzy_sources (source_id INTEGER PRIMARY KEY,text TEXT);
        INSERT INTO fuzzy_sources
            SELECT s.rowid,s.plain_text FROM source_rows s
            WHERE s.plain_text!=''
              AND NOT EXISTS (SELECT 1 FROM matched_sources m WHERE m.source_id=s.rowid)
              AND NOT EXISTS (SELECT 1 FROM structure_pending_sources p WHERE p.source_id=s.rowid)
              AND NOT EXISTS (SELECT 1 FROM ({RECOMMENDED_PAIRS}) r
                              WHERE r.text=s.text AND r.korean_key=s.korean_key);
        CREATE INDEX fuzzy_source_text ON fuzzy_sources(text);
        CREATE INDEX baseline_plain ON baseline_rows(plain_text);
        CREATE INDEX structure_pending_position ON structure_pending(path,sheet,row);
        CREATE TEMP TABLE fuzzy_targets (id INTEGER PRIMARY KEY,text TEXT UNIQUE);
        INSERT INTO fuzzy_targets(text)
            SELECT DISTINCT b.plain_text FROM baseline_rows b WHERE b.plain_text!=''
              AND EXISTS (SELECT 1 FROM fuzzy_sources)
              AND EXISTS (SELECT 1 FROM recommendations r
                          WHERE r.path=b.path AND r.sheet=b.sheet AND r.row=b.row AND r.source_id IS NULL)
              AND NOT EXISTS (SELECT 1 FROM structure_recommendations t
                              WHERE t.path=b.path AND t.sheet=b.sheet AND t.row=b.row)
              AND NOT EXISTS (SELECT 1 FROM structure_pending p
                              WHERE p.path=b.path AND p.sheet=b.sheet AND p.row=b.row);
        CREATE TEMP TABLE fuzzy_characters (
            character TEXT,size INTEGER,target_id INTEGER,n INTEGER,
            PRIMARY KEY(character,size,target_id)) WITHOUT ROWID;
        CREATE TEMP TABLE fuzzy_query (character TEXT PRIMARY KEY,n INTEGER) WITHOUT ROWID;
        CREATE TEMP TABLE fuzzy_scores (target_id INTEGER,text TEXT,similarity REAL,
            PRIMARY KEY(target_id,text)) WITHOUT ROWID;
    """)
    if progress:
        progress(f"对仍未匹配的句子查找相似文本（至少 {FUZZY_MIN_PERCENT}%）……")
    has_text = lambda value: any(unicodedata.category(c).startswith("L") for c in value)
    for target_id, text in connection.execute("SELECT id,text FROM fuzzy_targets"):
        if has_text(text):
            connection.executemany("INSERT INTO fuzzy_characters VALUES (?,?,?,?)",
                [(c, len(text), target_id, n) for c, n in Counter(text).items()])
    for index, (text,) in enumerate(connection.execute("SELECT DISTINCT text FROM fuzzy_sources ORDER BY text"), 1):
        if not has_text(text):
            continue
        connection.execute("DELETE FROM fuzzy_query")
        connection.executemany("INSERT INTO fuzzy_query VALUES (?,?)", Counter(text).items())
        size = len(text)
        lower = (size * FUZZY_MIN_PERCENT + 199 - FUZZY_MIN_PERCENT) // (200 - FUZZY_MIN_PERCENT)
        upper = size * (200 - FUZZY_MIN_PERCENT) // FUZZY_MIN_PERCENT
        # Length and shared character counts are upper bounds on SequenceMatcher's
        # ratio, so this pruning cannot discard an above-threshold comparison.
        candidates = connection.execute("""SELECT t.id,t.text FROM fuzzy_targets t JOIN (
            SELECT f.target_id FROM fuzzy_query q JOIN fuzzy_characters f
                ON f.character=q.character AND f.size BETWEEN ? AND ?
            GROUP BY f.target_id HAVING 200*SUM(MIN(f.n,q.n)) >= ?*(f.size+?)
        ) c ON c.target_id=t.id""", (lower, upper, FUZZY_MIN_PERCENT, size))
        matcher = SequenceMatcher(None, "", text, autojunk=False)
        for target_id, baseline_text in candidates:
            matcher.set_seq1(baseline_text)
            similarity = matcher.ratio()
            if similarity * 100 >= FUZZY_MIN_PERCENT:
                connection.execute("INSERT INTO fuzzy_scores VALUES (?,?,?)", (target_id, text, similarity))
        if progress and index % 100 == 0:
            progress(f"模糊匹配：已比较 {index} 种未匹配源句")
    # Expand the fixed input snapshot only after scoring; one source may serve
    # several baseline occurrences, independent of traversal order.
    connection.execute("""INSERT INTO comparisons
        SELECT b.path,b.sheet,b.row,s.source_id,0,0,0,?, '[]','[]','[]','[]','',
               '本轮仅比较中文相似度，未使用上下文得分',f.similarity
        FROM fuzzy_scores f JOIN fuzzy_targets t ON t.id=f.target_id
        JOIN fuzzy_sources s ON s.text=f.text JOIN baseline_rows b ON b.plain_text=t.text
        WHERE EXISTS (SELECT 1 FROM recommendations r
                      WHERE r.path=b.path AND r.sheet=b.sheet AND r.row=b.row AND r.source_id IS NULL)
          AND NOT EXISTS (SELECT 1 FROM structure_recommendations t
                          WHERE t.path=b.path AND t.sheet=b.sheet AND t.row=b.row)
          AND NOT EXISTS (SELECT 1 FROM structure_pending p
                          WHERE p.path=b.path AND p.sheet=b.sheet AND p.row=b.row)""", (FUZZY_MATCH,))
    recommended = 0
    for path, sheet, row, n, missing, variants, highest in connection.execute("""
        SELECT e.path,e.sheet,e.row,COUNT(*),SUM(s.korean_key=''),
               COUNT(DISTINCT NULLIF(s.korean_key,'')),MAX(e.similarity)
        FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
        WHERE (e.path,e.sheet,e.row) IN (SELECT path,sheet,row FROM comparisons WHERE match_type=?)
        GROUP BY e.path,e.sheet,e.row""", (FUZZY_MATCH,)):
        key = (path, sheet, row)
        fuzzy_highest = highest
        best = connection.execute("""SELECT e.source_id,s.korean_key,COALESCE(e.similarity,1.0)
            FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
            WHERE e.path=? AND e.sheet=? AND e.row=? AND s.korean_key!=''
            ORDER BY COALESCE(e.similarity,1.0) DESC,e.source_id LIMIT 1""", key).fetchone()
        source_id, status = None, "全部缺译"
        reason = "达到阈值的来源韩语均为空，无法推荐"
        if best:
            highest = best[2]
            tied = connection.execute("""SELECT 1 FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
                WHERE e.path=? AND e.sheet=? AND e.row=? AND COALESCE(e.similarity,1.0)=?
                  AND s.korean_key!='' AND s.korean_key!=? LIMIT 1""", (*key, highest, best[1])).fetchone()
            if tied:
                status, reason = "待确认", "最高相似度对应不同韩语，无法唯一推荐"
            else:
                source_id, status = best[0], "模糊推荐·待复核"
                reason = "采用最高相似度的非空译文；同分同译取读取顺序首条"
                recommended += 1
        has_prior = connection.execute("SELECT 1 FROM comparisons WHERE path=? AND sheet=? AND row=? AND similarity IS NULL", key).fetchone()
        similarity_note = (f"新增模糊候选最高相似度 {fuzzy_highest:.1%}" if has_prior and source_id is None else
                           f"相似度 {highest:.1%}")
        basis = (f"v1.1.12规则补匹配；模糊匹配，{similarity_note}（最低 {FUZZY_MIN_PERCENT}%）；{reason}；"
                 "按去标签及换行、统一指定标点后的中文字符顺序比较；仍需复核文字差异、语义及标签")
        if has_prior:
            basis += "；保留旧规则候选共同判断，已等价命中的旧候选优先于近似候选，仍有不同译文并列时不推荐"
        connection.execute("UPDATE recommendations SET source_id=?,basis=?,status=?,match_type=? WHERE path=? AND sheet=? AND row=?",
                           (source_id, basis, status, FUZZY_MATCH, *key))
        connection.execute("INSERT OR REPLACE INTO baseline_counts VALUES (?,?,?,?,?,?)", (*key, n, missing, variants))
    connection.execute("INSERT OR IGNORE INTO matched_sources SELECT source_id FROM comparisons WHERE match_type=?", (FUZZY_MATCH,))
    for table in ("fuzzy_sources", "fuzzy_targets", "fuzzy_characters", "fuzzy_query", "fuzzy_scores"):
        connection.execute(f"DROP TABLE {table}")
    connection.commit()
    return recommended


def write_baseline_statistics(destination, connection, baseline_parts):
    headers = ["非空中文行数", "中文匹配行数", "有译文行数", "推荐行数", "未匹配行数", "命中但全缺译行数",
               "冲突行数", "冲突已推荐行数", "冲突待确认行数", "中文匹配率", "译文覆盖率", "推荐率"]

    def counts(path, sheet=None):
        clause = "b.path=?" + (" AND b.sheet=?" if sheet is not None else "")
        arguments = (str(path), sheet) if sheet is not None else (str(path),)
        n, matched, available, recommended, missing, conflicts, resolved = connection.execute(f"""
            SELECT COUNT(*),COALESCE(SUM(c.n>0),0),COALESCE(SUM(c.variants>0),0),
                   COALESCE(SUM(r.source_id IS NOT NULL OR t.row IS NOT NULL),0),COALESCE(SUM(c.variants=0),0),
                   COUNT(f.row),COALESCE(SUM(f.row IS NOT NULL AND (r.source_id IS NOT NULL OR t.row IS NOT NULL)),0)
            FROM baseline_rows b LEFT JOIN baseline_counts c ON c.path=b.path AND c.sheet=b.sheet AND c.row=b.row
            LEFT JOIN baseline_counts f ON f.path=b.path AND f.sheet=b.sheet AND f.row=b.row AND f.variants>1
            JOIN recommendations r ON r.path=b.path AND r.sheet=b.sheet AND r.row=b.row
            LEFT JOIN structure_recommendations t ON t.path=b.path AND t.sheet=b.sheet AND t.row=b.row WHERE {clause}
        """, arguments).fetchone()
        return [n, matched, available, recommended, n - matched, missing, conflicts, resolved, conflicts - resolved]

    def with_rates(values):
        n = values[0]
        return [*values, *(values[i] / n if n else None for i in (1, 2, 3))]

    book = Workbook(write_only=True)
    file_rows, sheet_rows, total = [], [], [0] * 9
    for path, parts in baseline_parts.items():
        values = counts(path)
        total = [a + b for a, b in zip(total, values)]
        file_rows.append([path.name, str(path), *with_rates(values)])
        for sheet, _, _ in parts.values():
            sheet_rows.append([path.name, str(path), sheet, *with_rates(counts(path, sheet))])
    write_detail_sheet(book, "文件统计", ["文件名", "完整路径", *headers], file_rows, [35, 65, *([20] * len(headers))])
    write_detail_sheet(book, "Sheet统计", ["文件名", "完整路径", "Sheet", *headers], sheet_rows, [35, 65, 25, *([20] * len(headers))])
    write_detail_sheet(book, "总体统计", ["基准文件数", "基准 Sheet 数", *headers],
                       [[len(baseline_parts), len(sheet_rows), *with_rates(total)]], [20] * (len(headers) + 2))
    book.save(destination)
    book.close()


def write_matching_details(destination, connection, progress=None):
    book = Workbook(write_only=True)
    matches = connection.execute("""
        SELECT b.path,b.sheet,b.row,s.path,s.sheet,s.row,s.chinese,s.korean,
               CASE WHEN s.korean_key='' THEN '是' ELSE '否' END,
               CASE WHEN c.variants>1 THEN '是' ELSE '否' END,
               e.baseline_previous,e.baseline_following,e.source_previous,e.source_following,
               e.previous_equal,e.following_equal,e.full_sides,
               CASE WHEN r.source_id=s.rowid THEN '是' ELSE '否' END,
               COALESCE(t.status,r.status),COALESCE(t.basis,r.basis),e.match_type,e.skipped,e.query_statistics,b.text,e.similarity
        FROM comparisons e JOIN source_rows s ON s.rowid=e.source_id
        JOIN baseline_rows b ON b.path=e.path AND b.sheet=e.sheet AND b.row=e.row
        JOIN baseline_counts c ON c.path=b.path AND c.sheet=b.sheet AND c.row=b.row
        JOIN recommendations r ON r.path=b.path AND r.sheet=b.sheet AND r.row=b.row
        LEFT JOIN structure_recommendations t ON t.path=b.path AND t.sheet=b.sheet AND t.row=b.row
        ORDER BY s.rowid,b.path,b.sheet,b.row
    """)
    rows = ([Path(b).name,b,bs,br,Path(p).name,p,ss,sr,cn,ko,missing,conflict,
             context_text(json.loads(bp), "前"),context_text(json.loads(bf), "后"),
             context_text(json.loads(sp), "前"),context_text(json.loads(sf), "后"),
             pe,fe,full,pe+fe,selected,status,basis,match_type,
             "是" if full == 2 else "否",
             skipped if len(skipped) <= 32000 else skipped[:32000] + "\n（展示截断，完整特殊行见上下文排除明细）", statistics,
             baseline_text, similarity]
            for b,bs,br,p,ss,sr,cn,ko,missing,conflict,bp,bf,sp,sf,pe,fe,full,selected,status,basis,match_type,skipped,statistics,baseline_text,similarity in matches)
    pairs = write_detail_sheet(book, "匹配明细",
                               ["基准文件名","基准完整路径","基准 Sheet","基准行号","源文件名","源完整路径","源 Sheet","源行号","源中文","韩语","缺译","译文冲突",
                                "基准前3句","基准后3句","来源前3句","来源后3句","前句一致数","后句一致数","完整一致侧数","一致位置总数","推荐来源","推荐状态","推荐依据","匹配方式",
                                "完美匹配","上下文跳过位置","上下文查询统计","基准中文","相似度"],
                               rows, [30,65,25,16,35,65,25,16,55,55,12,16,65,65,65,65,18,18,20,20,16,24,70,24,16,75,75,65,16], progress)
    structure_headers = ["基准完整路径", "基准Sheet", "基准行号", "基准中文", "推荐韩语", "推荐来源",
                         "推荐依据", "推荐状态", "类型", "组ID"]
    pairs += write_detail_sheet(book, "合并拆分推荐", structure_headers,
        connection.execute("SELECT * FROM structure_recommendations ORDER BY path,sheet,row"),
        [60, 22, 14, 70, 75, 80, 90, 26, 14, 18], progress)
    write_detail_sheet(book, "结构推荐未写入", [*structure_headers, "未写入原因"],
        connection.execute("SELECT * FROM structure_pending ORDER BY path,sheet,row"),
        [60, 22, 14, 70, 75, 80, 90, 26, 14, 18, 75], progress)
    write_detail_sheet(book, "上下文排除明细",
        ["源完整路径", "源Sheet", "源行号", "源中文", "源韩文", "特殊说明", "说话角色名", "处理"],
        connection.execute("SELECT * FROM context_exclusions ORDER BY path,sheet,row"),
        [65, 25, 16, 70, 70, 20, 25, 75], progress)
    for title, condition in (("未匹配明细", "t.source_id IS NULL"), ("缺译明细", "s.korean_key=''"), ("译文冲突", "c.text IS NOT NULL")):
        # Conditions are fixed program constants, never user input.
        records = connection.execute(f"""
            SELECT s.path,s.sheet,s.row,s.chinese,s.korean,
                   CASE WHEN t.source_id IS NOT NULL THEN '是' ELSE '否' END,
                   CASE WHEN s.korean_key='' THEN '是' ELSE '否' END,
                   (SELECT COUNT(DISTINCT v.korean_key) FROM source_rows v WHERE v.text=s.text AND v.korean_key!='')
            FROM source_rows s LEFT JOIN matched_sources t ON t.source_id=s.rowid
            LEFT JOIN conflicts c ON c.text=s.text WHERE {condition} ORDER BY s.text,s.rowid
        """)
        rows = ([Path(p).name,p,sheet,row,cn,ko,matched,missing,variants]
                for p,sheet,row,cn,ko,matched,missing,variants in records)
        write_detail_sheet(book, title, ["源文件名","源完整路径","源 Sheet","源行号","源中文","韩语","命中基准","缺译","非空韩语种类数"],
                           rows, [35,65,25,16,55,55,16,12,24], progress)
    records = connection.execute(f"""
        WITH recommended_pairs AS ({RECOMMENDED_PAIRS})
        SELECT s.path,s.sheet,s.row,s.chinese,s.korean,
               CASE WHEN t.source_id IS NOT NULL THEN '是' ELSE '否' END,
               CASE WHEN s.korean_key='' THEN '是' ELSE '否' END,
               CASE WHEN EXISTS (SELECT 1 FROM structure_sources u WHERE u.source_id=s.rowid)
                    THEN '拆分后仅部分分段获推荐，剩余分段待复核'
                    WHEN EXISTS (SELECT 1 FROM structure_pending_sources p WHERE p.source_id=s.rowid)
                    THEN '结构候选未写入：' || (SELECT GROUP_CONCAT(DISTINCT p.reason)
                                              FROM structure_pending_sources p WHERE p.source_id=s.rowid)
                    WHEN t.source_id IS NULL THEN '基准文件中找不到'
                    ELSE '基准文件中找到了但不符合推荐规则' END
        FROM source_rows s LEFT JOIN matched_sources t ON t.source_id=s.rowid
        WHERE NOT EXISTS (SELECT 1 FROM recommended_pairs r
                          WHERE r.text=s.text AND r.korean_key=s.korean_key)
        ORDER BY s.rowid
    """)
    rows = ([Path(p).name,p,sheet,row,cn,ko,matched,missing,reason]
            for p,sheet,row,cn,ko,matched,missing,reason in records)
    write_detail_sheet(book, "未推荐明细",
                       ["源文件名","源完整路径","源 Sheet","源行号","源中文","韩语","命中基准","缺译","未推荐原因"],
                       rows, [35,65,25,16,55,55,16,12,55], progress)
    book.save(destination)
    book.close()
    return pairs


def issue_counts(connection, path, sheet=None):
    clause = "s.path=?" + (" AND s.sheet=?" if sheet is not None else "")
    arguments = (str(path), sheet) if sheet is not None else (str(path),)
    return connection.execute(f"""
        SELECT COALESCE(SUM(s.korean_key=''),0),COUNT(DISTINCT c.text)
        FROM source_rows s LEFT JOIN conflicts c ON c.text=s.text WHERE {clause}
    """, arguments).fetchone()


def write_statistics(destination, file_rows, sheet_rows, skipped, report, baseline_root):
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows, widths in (
        ("文件统计", ["文件名", "完整路径", "参与 Sheet 数", "跳过 Sheet / 文件项数", "非空中文行数", "匹配行数", "未匹配行数", "匹配率", "处理状态", "缺译行数", "冲突中文数"], file_rows, [38, 70, 18, 24, 18, 18, 18, 16, 24,18,18]),
        ("Sheet统计", ["文件名", "完整路径", "Sheet", "非空中文行数", "匹配行数", "未匹配行数", "匹配率", "缺译行数", "冲突中文数"], sheet_rows, [38, 70, 30, 18, 18, 18, 16,18,18]),
        ("跳过清单", ["文件名", "完整路径", "Sheet", "跳过原因"], skipped, [38, 70, 30, 85]),
        ("扫描说明", ["项目", "本次设置"], [
            ["程序版本", VERSION], ["源文件夹", str(report.root)], ["基准文件夹", str(baseline_root)],
            ["表头检查行数", report.header_rows], ["中文表头", "、".join(report.chinese_headers)],
            ["韩语表头", "、".join(report.korean_headers)],
            *[[label, "是" if value else "否"] for label, value in (
                ("包含子文件夹", report.recursive), ("包含隐藏工作表", report.include_hidden_sheets),
                ("包含隐藏行", report.include_hidden_rows), ("包含废弃行", report.include_deprecated))],
            ["本地忽略 Sheet 数", report.ignored_sheets],
            ["匹配优先级", "先按v1.1.9整句及启用的合并/拆分推荐；已推荐结果锁定。其余按v1.1.12标点等价、80%模糊补查，保留旧候选。"],
            ["候选范围", "唯一译文仅指实际参与候选的非空韩语只有一种；未扫描或被排除的来源不参与判断。"],
            ["排除记录", "整文件/Sheet 见跳过清单；废弃、说明公式、隐藏行及上下文跳过见匹配明细.xlsx 的上下文排除明细。"],
        ], [28, 100]),
    ):
        sheet = book.create_sheet(title)
        for row_number, values in enumerate([headers, *rows], 1):
            for column, value in enumerate(values, 1):
                cell = text_cell(sheet, row_number, column, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="245B78")
                elif headers[column - 1] == "匹配率":
                    cell.number_format = "0.00%"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
    book.save(destination)
    book.close()


def prepare_inputs(report, baseline_root, output_root, ignored_sheets=(), progress=None):
    """Validate roots and recheck sheets for both matching and structure scans."""
    baseline_root = Path(baseline_root).expanduser().resolve()
    output_root = Path(output_root).expanduser().resolve()
    source_root = report.root.resolve()
    if not baseline_root.is_dir():
        raise ValueError("请选择存在的基准文件夹")
    if source_root == baseline_root or source_root in baseline_root.parents or baseline_root in source_root.parents:
        raise ValueError("基准文件夹和源文件夹必须分开，不能互相包含")
    if any(output_root == root or root in output_root.parents for root in (source_root, baseline_root)):
        raise ValueError("输出文件夹不能位于源文件夹或基准文件夹内，以免再次扫描结果")
    baselines = sorted(p for p in baseline_root.rglob("*") if p.is_file()
                       and p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
    if not baselines:
        raise ValueError("基准文件夹中没有 .xlsx 或 .xlsm 文件")
    ignored = set(ignored_sheets)
    previous_exceptions = {sheet_key(r.path, r.sheet): r for r in report.exceptions}
    # Recheck at execution time so changed/new sheets cannot bypass validation.
    fresh = scan_folder(source_root, report.header_rows, report.chinese_headers,
                        report.korean_headers, ignored_sheets=ignored,
                        recursive=report.recursive, include_hidden_sheets=report.include_hidden_sheets,
                        include_hidden_rows=report.include_hidden_rows, include_deprecated=report.include_deprecated,
                        progress=lambda i, n, p: progress(f"复查表头 {i}/{n}：{p.name}") if progress else None)
    candidates, skipped = {}, []
    for result in fresh.skipped_sheets:
        skipped.append([result.path.name, str(result.path), result.sheet, result.issue])
    for index, result in enumerate(fresh.results):
        key = sheet_key(result.path, result.sheet)
        if key in previous_exceptions:
            if result.passed:
                result = previous_exceptions[key]
                fresh.results[index] = result
        if result.passed and sheet_key(result.path, "") in previous_exceptions:
            result.issue = previous_exceptions[sheet_key(result.path, "")].issue
        if result.passed and sheet_key(result.path, "") not in previous_exceptions:
            candidates.setdefault(result.path, []).append(result)
        else:
            reason = result.issue or previous_exceptions[sheet_key(result.path, "")].issue
            skipped.append([result.path.name, str(result.path), result.sheet, "异常：" + reason])
    # Include only ignore records that actually exist in this source folder.
    for path in fresh.file_paths:
        records = [sheet for filename, sheet in ignored if filename == sheet_key(path, "")[0]]
        if records:
            try:
                book = load_workbook(path, read_only=True, keep_links=False)
                try:
                    for sheet in book.sheetnames:
                        if sheet_key(path, sheet) in ignored:
                            skipped.append([path.name, str(path), sheet, "本地已忽略"])
                finally:
                    book.close()
            except Exception:
                pass  # File-level read errors already appear in the header report.
    if not candidates:
        raise ValueError("没有可参与匹配的 Sheet；请先检查表头并处理异常或调整忽略记录")
    return fresh, candidates, skipped, baselines, baseline_root, output_root


def match_folder(report, baseline_root, output_root, ignored_sheets=(), progress=None, structure_candidates=()):
    fresh, candidates, skipped, baselines, baseline_root, output_root = prepare_inputs(
        report, baseline_root, output_root, ignored_sheets, progress)
    output_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="匹配临时_", dir=output_root) as temporary:
        staging = Path(temporary)
        connection = sqlite3.connect(staging / "index.sqlite3")
        baseline_rows = matched_rows = 0
        sheet_rows = []
        try:
            connection.executescript("""
                CREATE TABLE targets (text TEXT PRIMARY KEY,breath_text TEXT);
                CREATE INDEX target_breath ON targets(breath_text);
                CREATE TABLE baseline_rows (path TEXT, sheet TEXT, row INTEGER, text TEXT, seq INTEGER,
                    leading_newline INTEGER,trailing_newline INTEGER,plain_text TEXT,PRIMARY KEY (path, sheet, row));
                CREATE INDEX baseline_text ON baseline_rows(text);
                CREATE UNIQUE INDEX baseline_sequence ON baseline_rows(path,sheet,seq);
                CREATE TABLE source_rows (text TEXT, chinese TEXT, korean TEXT, path TEXT, sheet TEXT, row INTEGER,
                    korean_key TEXT, seq INTEGER,plain_text TEXT,breath_text TEXT,breath_plain TEXT);
                CREATE INDEX source_text ON source_rows(text);
                CREATE INDEX source_position ON source_rows(path,sheet);
                CREATE UNIQUE INDEX source_sequence ON source_rows(path,sheet,seq);
                CREATE INDEX source_plain ON source_rows(plain_text);
                CREATE INDEX source_breath ON source_rows(breath_text);
                CREATE INDEX source_breath_plain ON source_rows(breath_plain);
                CREATE TABLE source_context_rows (
                    path TEXT,sheet TEXT,row INTEGER,seq INTEGER,text TEXT,plain_text TEXT,
                    context_skip TEXT,breath_text TEXT,breath_plain TEXT);
                CREATE UNIQUE INDEX context_sequence ON source_context_rows(path,sheet,seq);
                CREATE INDEX baseline_plain_sequence ON baseline_rows(path,sheet,seq) WHERE plain_text!='';
                CREATE VIEW hits AS SELECT s.rowid AS rowid,s.*,t.text AS target_text FROM source_rows s
                    JOIN targets t ON t.text=s.text OR (s.breath_text!='' AND t.breath_text=s.breath_text);
                CREATE TABLE context_exclusions (
                    path TEXT,sheet TEXT,row INTEGER,chinese TEXT,korean TEXT,note TEXT,speaker TEXT,reason TEXT);
            """)
            baseline_parts = {}
            for path in baselines:
                if progress:
                    progress(f"读取基准中文：{path.name}")
                book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet, header_row, text_column, output_column in baseline_sheets(book, path, report.header_rows):
                        baseline_parts.setdefault(path, {})[sheet._worksheet_path.lstrip("/")] = (sheet.title, header_row, output_column)
                        sequence = 0
                        for row in sheet.iter_rows(min_row=header_row + 1, min_col=text_column, max_col=text_column):
                            cell = row[0]
                            if cell.data_type == "f":
                                raise ValueError(f"基准中文不能是公式：{path.name} / {sheet.title} / {cell.coordinate}")
                            key = match_key(cell.value)
                            if key:
                                baseline_rows += 1
                                sequence += 1
                                connection.execute("INSERT OR IGNORE INTO targets VALUES (?,?)", (key, breath_key(key)))
                                connection.execute("INSERT INTO baseline_rows VALUES (?,?,?,?,?,?,?,?)",
                                                   (str(path), sheet.title, cell.row, key, sequence, *boundary_newlines(cell.value), plain_key(cell.value, punctuation=False)))
                finally:
                    book.close()
            if not baseline_rows:
                raise ValueError("基准文件的 Text 列没有非空中文文本")
            connection.commit()
            for index, (path, results) in enumerate(candidates.items(), 1):
                if progress:
                    progress(f"检索源文件 {index}/{len(candidates)}：{path.name}")
                book = None
                try:
                    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                    for result in results:
                        connection.execute("SAVEPOINT sheet")
                        try:
                            sheet = book[result.sheet]
                            sheet.reset_dimensions()
                            hidden_rows = hidden_row_numbers(book, sheet) if not report.include_hidden_rows else set()
                            total = matched = sequence = 0
                            columns = [result.chinese_column, result.korean_column]
                            header = next(sheet.iter_rows(min_row=result.header_row, max_row=result.header_row))
                            notes_column, speaker_column, remarks_columns = source_metadata_columns(header)
                            columns.extend(c for c in (notes_column, speaker_column, *remarks_columns) if c)
                            first, last = min(columns), max(columns)
                            for number, row in enumerate(sheet.iter_rows(min_row=result.header_row + 1,
                                                                        min_col=first, max_col=last), result.header_row + 1):
                                chinese = row[result.chinese_column - first]
                                korean = row[result.korean_column - first]
                                note = row[notes_column - first] if notes_column else None
                                speaker = row[speaker_column - first] if speaker_column else None
                                if number in hidden_rows:
                                    note_text = str(note.value) if note is not None and note.value is not None else ""
                                    speaker_text = str(speaker.value) if speaker is not None and speaker.value is not None else ""
                                    excluded = "隐藏行（未勾选包含隐藏行）"
                                else:
                                    note_text, speaker_text, excluded = source_row_metadata(
                                        note, speaker, report.include_deprecated,
                                        [row[c - first] for c in remarks_columns])
                                if (excluded and (match_key(chinese.value) or note_text)) or note_text in CONTEXT_SKIP_NOTES:
                                    connection.execute("INSERT INTO context_exclusions VALUES (?,?,?,?,?,?,?,?)",
                                        (str(path), sheet.title, number, chinese.value,
                                         None if korean.value is None else str(korean.value), note_text, speaker_text,
                                         "排除本行：" + excluded if excluded else
                                         "上下文不计有效句，继续向外补查；有正文时仍参与整句匹配" +
                                         ("；比较时忽略成对括号说明" if note_text == "气息" else "")))
                                if excluded:
                                    continue
                                if chinese.data_type == "f" or korean.data_type == "f":
                                    raise ValueError(f"第 {number} 行中韩文本含公式，请先转成文本值")
                                key = match_key(chinese.value)
                                context_skip = note_text if note_text in CONTEXT_SKIP_NOTES else ""
                                breath = breath_key(chinese.value) if note_text == "气息" else ""
                                if key or context_skip:
                                    sequence += 1
                                    connection.execute("INSERT INTO source_context_rows VALUES (?,?,?,?,?,?,?,?,?)",
                                        (str(path), sheet.title, number, sequence, key, plain_key(chinese.value, punctuation=False),
                                         context_skip, breath, plain_key(breath, punctuation=False)))
                                if not key:
                                    continue
                                total += 1
                                translation = None if korean.value is None else str(korean.value)
                                connection.execute("INSERT INTO source_rows VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                                                   (key, chinese.value, translation, str(path), sheet.title, number,
                                                    match_key(translation), sequence, plain_key(chinese.value, punctuation=False),
                                                    breath, plain_key(breath, punctuation=False)))
                                if connection.execute("SELECT 1 FROM targets WHERE text=? OR (?!='' AND breath_text=?)",
                                                      (key, breath, breath)).fetchone():
                                    matched += 1
                            connection.execute("RELEASE sheet")
                            sheet_rows.append([path.name, str(path), result.sheet, total, matched, total - matched,
                                               matched / total if total else None])
                        except sqlite3.Error:
                            raise
                        except Exception as error:
                            connection.execute("ROLLBACK TO sheet")
                            connection.execute("RELEASE sheet")
                            result.issue = f"正文读取失败：{type(error).__name__}: {error}"
                            skipped.append([path.name, str(path), result.sheet, result.issue])
                    connection.commit()
                except sqlite3.Error:
                    raise
                except Exception as error:
                    for result in results:
                        result.issue = f"文件读取失败：{type(error).__name__}: {error}"
                        skipped.append([path.name, str(path), result.sheet, result.issue])
                finally:
                    if book is not None:
                        book.close()
            if not sheet_rows:
                failures = [f"{r.path.name} / {r.sheet}：{r.issue}" for results in candidates.values() for r in results if r.issue]
                raise ValueError("所有候选 Sheet 正文读取失败，未生成匹配结果：\n" + "\n".join(failures[:10]))
            connection.executescript("""
                CREATE TABLE counts AS SELECT target_text AS text, COUNT(*) AS n,SUM(korean_key='') AS missing,
                    COUNT(DISTINCT CASE WHEN korean_key!='' THEN korean_key END) AS variants FROM hits GROUP BY target_text;
                CREATE UNIQUE INDEX count_text ON counts(text);
                CREATE TABLE conflicts AS SELECT text,COUNT(DISTINCT korean_key) AS variants FROM source_rows
                    WHERE korean_key!='' GROUP BY text HAVING COUNT(DISTINCT korean_key)>1;
                CREATE UNIQUE INDEX conflict_text ON conflicts(text);
            """)
            missing_translations = connection.execute("SELECT COUNT(*) FROM source_rows WHERE korean_key='' ").fetchone()[0]
            conflict_texts = connection.execute("SELECT COUNT(*) FROM conflicts").fetchone()[0]
            if progress:
                progress("比较前后三句并生成推荐……")
            recommended_rows = build_recommendations(connection, progress)
            recommended_rows += apply_structure_recommendations(connection, structure_candidates)
            recommended_rows += apply_punctuation_recommendations(connection, progress)
            recommended_rows += apply_fuzzy_recommendations(connection, progress)
            for row in sheet_rows:
                row[4] = connection.execute("""SELECT COUNT(*) FROM source_rows s JOIN matched_sources m ON m.source_id=s.rowid
                    WHERE s.path=? AND s.sheet=?""", (row[1], row[2])).fetchone()[0]
                row[5], row[6] = row[3] - row[4], row[4] / row[3] if row[3] else None
            outputs = []
            for path in baselines:
                if progress:
                    progress(f"写入基准副本：{path.name}")
                relative = path.relative_to(baseline_root)
                destination = staging / "基准匹配结果" / relative
                destination.parent.mkdir(parents=True, exist_ok=True)
                matched_rows += write_baseline_copy(path, destination, connection, baseline_parts[path])
                outputs.append(Path("基准匹配结果") / relative)
            if progress:
                progress("写出匹配、未匹配、缺译、冲突及未推荐明细……")
            match_pairs = write_matching_details(staging / "匹配明细.xlsx", connection, progress)
            outputs.append(Path("匹配明细.xlsx"))
            write_baseline_statistics(staging / "基准文件匹配统计.xlsx", connection, baseline_parts)
            outputs.append(Path("基准文件匹配统计.xlsx"))
            for row in sheet_rows:
                row.extend(issue_counts(connection, row[1], row[2]))
            file_rows = []
            for filename in sorted({r[1] for r in sheet_rows} | {r[1] for r in skipped}):
                rows = [r for r in sheet_rows if r[1] == filename]
                skip_count = sum(r[1] == filename for r in skipped)
                total = sum(r[3] for r in rows)
                matched = sum(r[4] for r in rows)
                file_rows.append([Path(filename).name, filename, len(rows), skip_count, total, matched, total - matched,
                                  matched / total if total else None, "部分跳过" if rows and skip_count else "已扫描" if rows else "全部跳过",
                                  *issue_counts(connection, filename)])
            write_statistics(staging / "源文件匹配统计.xlsx", file_rows, sheet_rows, skipped, fresh, baseline_root)
            outputs.append(Path("源文件匹配统计.xlsx"))
        finally:
            connection.close()
        (staging / "index.sqlite3").unlink()
        destination = output_root / ("匹配结果_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + staging.name[-6:])
        os.rename(staging, destination)
    return MatchReport(destination, fresh, baseline_rows, matched_rows, sum(r[3] for r in sheet_rows),
                       sum(r[4] for r in sheet_rows), len(skipped), [destination / p for p in outputs],
                       missing_translations, conflict_texts, match_pairs, recommended_rows, len(baselines))
