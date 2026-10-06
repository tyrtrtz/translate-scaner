"""One desktop smoke test: layout and the check → ignore → match workflow."""

import tempfile
import time
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from app import Application, AutoScrollbar
from settings import load_settings


class DesktopWorkflow(unittest.TestCase):
    def test_layout_and_workflow(self):
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"Tk display unavailable: {str(error).splitlines()[0]}")
        self.addCleanup(root.destroy)
        with tempfile.TemporaryDirectory() as directory, patch("app.messagebox.showerror") as errors:
            base = Path(directory)
            source, baseline = base / "源文件", base / "基准"
            source.mkdir()
            baseline.mkdir()
            book = Workbook()
            book.active.title = "翻译"
            book.active.append(["中文原文", "韩文"])
            book.active.append(["你好", "안녕하세요"])
            book.create_sheet("参考").append(["修改记录"])
            book.save(source / "sample.xlsx")
            book.close()
            book = Workbook()
            book.active["G1"] = "Text"
            book.active["G2"] = "你好"
            book.save(baseline / "baseline.xlsx")
            book.close()

            app = Application(root, ignore_file=base / "ignored.json", settings_file=base / "settings.json")
            root.geometry("1060x760")
            root.update()
            with patch("app.check_for_update", return_value=None), patch("app.messagebox.showinfo") as notices:
                app.update_button.invoke()
                self.wait_for_task(root, app)
                self.assertIn("已是最新版", app.update_message.get())
                notices.assert_called_once()
            with patch("app.check_for_update", side_effect=OSError("offline")):
                app.update_button.invoke()
                self.wait_for_task(root, app)
                self.assertIn("无法连接", app.update_message.get())
                self.assertTrue(app.start_button.instate(["!disabled"]))
                errors.assert_called_once()
                errors.reset_mock()
            with patch("app.check_for_update", return_value={"version": "1.2.0"}), \
                    patch("app.can_install_updates", return_value=True), \
                    patch("app.messagebox.askyesno", return_value=False) as confirm:
                app.set_running(True)
                app.check_updates(manual=False)
                deadline = time.monotonic() + 10
                while app.update_checking and time.monotonic() < deadline:
                    root.update()
                    time.sleep(0.01)
                self.assertFalse(app.update_checking)
                self.assertTrue(app.update_button.instate(["disabled"]))
                confirm.assert_not_called()
                app.set_running(False)
                app.update_button.invoke()
                confirm.assert_called_once()
                app.available_update = None
            self.assertTrue(app.status_label.winfo_ismapped())
            self.assertLessEqual(app.status_label.winfo_rooty() + app.status_label.winfo_height(), root.winfo_rooty() + root.winfo_height())
            self.assertTrue(app.match_button.instate(["disabled"]))
            self.assertTrue(app.exception_empty.winfo_ismapped())
            bars = [widget for widget in app.table.master.winfo_children() if isinstance(widget, AutoScrollbar)]
            horizontal = next(widget for widget in bars if str(widget.cget("orient")) == "horizontal")
            vertical = next(widget for widget in bars if str(widget.cget("orient")) == "vertical")
            self.assertFalse(horizontal.winfo_ismapped(), "Empty tables should not have scrollbars")
            self.assertFalse(vertical.winfo_ismapped())
            style = app.tabs.tk.call("ttk::style", "layout", "TNotebook")
            self.assertNotIn("Notebook.client", str(style))
            self.assertEqual(app.folder_entry.master.master.cget("background"), "#ffffff")
            for widget in (app.start_button, app.save_settings_button, app.ignore_button):
                self.assertGreaterEqual(widget.winfo_width(), widget.winfo_reqwidth())
                self.assertLessEqual(widget.winfo_y() + widget.winfo_height(), widget.master.winfo_height())
            app.tabs.select(app.match_tab)
            root.update()
            for widget in (app.match_button, app.match_open_button, app.match_details_button):
                self.assertGreaterEqual(widget.winfo_width(), widget.winfo_reqwidth())
                self.assertLessEqual(widget.winfo_x() + widget.winfo_width(), widget.master.winfo_width())

            app.folder.set(str(source))
            app.baseline_folder.set(str(baseline))
            app.output_folder.set(str(base / "结果"))
            app.start_button.invoke()
            self.assertTrue(app.start_button.instate(["disabled"]))
            self.wait_for_task(root, app)
            self.assertEqual(app.metric_values["files"].get(), "1")
            self.assertEqual(app.metric_values["passed"].get(), "1")
            self.assertEqual(app.metric_values["exceptions"].get(), "1")
            self.assertFalse(app.exception_empty.winfo_ismapped())
            self.assertTrue(horizontal.winfo_ismapped())
            # Custom thumb images must retain Tk's actual drag bindings.
            y = horizontal.winfo_height() // 2
            self.assertTrue(horizontal.identify(10, y).endswith("thumb"))
            horizontal.event_generate("<ButtonPress-1>", x=10, y=y)
            horizontal.event_generate("<B1-Motion>", x=80, y=y)
            horizontal.event_generate("<ButtonRelease-1>", x=80, y=y)
            root.update()
            self.assertGreater(app.table.xview()[0], 0, "Scrollbar dragging must scroll the table")
            app.table.xview_moveto(0)
            self.assertTrue(app.open_button.instate(["disabled"]))
            app.table.selection_set("0")
            app.table.check_all()
            root.update()
            self.assertEqual(app.selection_count.get(), "已勾选 1 项")
            self.assertTrue(app.open_button.instate(["!disabled"]))
            app.ignore_button.invoke()
            root.update()
            self.assertFalse(horizontal.winfo_ismapped())
            self.assertEqual(app.metric_values["exceptions"].get(), "0")
            self.assertEqual(app.metric_values["ignored"].get(), "1")
            self.assertTrue(app.exception_empty.winfo_ismapped())
            self.assertEqual(app.selection_count.get(), "已勾选 0 项")
            self.assertTrue(app.match_button.instate(["!disabled"]))

            app.tabs.select(app.match_tab)
            app.match_button.invoke()
            self.wait_for_task(root, app)
            self.assertIsNotNone(app.match_report)
            self.assertEqual(app.match_report.matched_rows, 1)
            self.assertTrue((app.match_report.destination / "匹配明细.xlsx").exists())
            self.assertTrue(app.match_open_button.instate(["!disabled"]))
            self.assertIn("检索完成", app.match_summary.get())
            self.assertEqual(app.match_summary_text.cget("state"), "disabled")
            self.assertIn("检索完成", app.match_summary_text.get("1.0", "end"))
            self.assertEqual(load_settings(base / "settings.json")["source_folder"], str(source))
            app.tabs.select(app.ignored_tab)
            app.ignored_table.check_all()
            root.update()
            app.restore_button.invoke()
            root.update()
            self.assertIsNone(app.report)
            self.assertEqual(app.metric_values["files"].get(), "—")
            self.assertEqual(app.metric_values["ignored"].get(), "0")
            self.assertTrue(app.ignored_empty.winfo_ismapped())
            self.assertTrue(app.match_button.instate(["disabled"]))
            errors.assert_not_called()

    def wait_for_task(self, root, app):
        deadline = time.monotonic() + 10
        while (app.running or app.update_checking) and time.monotonic() < deadline:
            root.update()
            time.sleep(0.01)
        root.update()
        self.assertFalse(app.running, "Background task did not finish")
        self.assertFalse(app.update_checking, "Update check did not finish")


if __name__ == "__main__":
    unittest.main()
