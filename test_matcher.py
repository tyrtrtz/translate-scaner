import hashlib
import tempfile
import unittest
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ignored_sheets import sheet_key
from matcher import TAIL_MATCH, align_boundary_newlines, breath_key, match_folder, plain_key, trailing_note_key, write_detail_sheet
from scanner import scan_folder


class TranslationMatches(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "源한국어"
        self.baseline = self.root / "基准"
        self.output = self.root / "结果"
        self.source.mkdir()
        self.baseline.mkdir()

    def write_source(self, rows, name="客户.xlsx", title="正文"):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        sheet.append(["台词", "한국어"])
        for row in rows:
            sheet.append(row)
        path = self.source / name
        book.save(path)
        book.close()
        return path

    def write_baseline(self, texts, name="story.xlsx"):
        book = Workbook()
        sheet = book.active
        sheet.cell(1, 7, "Text")
        sheet.cell(1, 1, "Key")
        sheet.cell(2, 1, "=1+1")
        sheet.cell(2, 1).font = Font(bold=True)
        sheet.column_dimensions["G"].width = 40
        for number, text in enumerate(texts, 2):
            sheet.cell(number, 7, text)
        path = self.baseline / name
        book.save(path)
        book.close()
        return path

    def test_baseline_headers_move_independently_per_sheet(self):
        self.write_source([["前一", "앞1"], ["前二", "앞2"], ["前三", "앞3"], ["目标", "추천"]])
        self.write_source([["其他一", "다른1"], ["其他二", "다른2"], ["其他三", "다른3"], ["目标", "다른 번역"]], name="另一个.xlsx")
        book = Workbook()
        book.remove(book.active)
        layouts = [("左侧", 3, 4, 2, 18), ("右侧", 1, 10, 8, 18), ("宽表", 5, 19, 52, 53)]
        for name, row, text, translation, _ in layouts:
            sheet = book.create_sheet(name)
            if row > 1:
                sheet['A1'] = '说明页眉'
                sheet['G2'] = '不是正文'
            sheet.cell(row, text, " Text \n")
            sheet.cell(row, translation, "TextTrans")
            for offset, value in enumerate(("前一", "前二", "前三", "\n目标\n"), 1):
                sheet.cell(row + offset, text, value)
                sheet.cell(row + offset, translation, '客户已有译文')
        path = self.baseline / '不同布局.xlsx'
        book.save(path)
        book.close()
        before = hashlib.sha256(path.read_bytes()).digest()
        result = match_folder(scan_folder(self.source), self.baseline, self.output)
        self.assertEqual((result.baseline_rows, result.matched_rows, result.recommended_rows, result.match_pairs), (12, 12, 12, 15))
        book = load_workbook(result.destination / '基准匹配结果/不同布局.xlsx')
        for name, row, text, translation, first in layouts:
            sheet = book[name]
            self.assertEqual(sheet.cell(row, text).value, ' Text \n')
            self.assertEqual(sheet.cell(row, translation).value, 'TextTrans')
            self.assertEqual(sheet.cell(row + 4, translation).value, '客户已有译文')
            self.assertEqual(sheet.cell(row, first).value, '匹配中文1')
            self.assertEqual(sheet.cell(row + 4, first + 12).value, '\n추천\n')
            self.assertEqual(sheet.cell(row + 4, first + 15).value, '上下文推荐')
            self.assertEqual(sheet.cell(row + 4, first + 9).value, 2)
            if row > 1:
                self.assertEqual(sheet['G2'].value, '不是正文')
                self.assertIsNone(sheet.cell(1, first).value)
        book.close()
        book = load_workbook(result.destination / '基准文件匹配统计.xlsx')
        self.assertEqual({r[2]: r[3:7] for r in list(book['Sheet统计'].values)[1:]},
                         {name: (4, 4, 4, 4) for name, *_ in layouts})
        book.close()
        book = load_workbook(result.destination / '匹配明细.xlsx')
        targets = [r for r in list(book['匹配明细'].values)[1:] if r[8] == '目标']
        self.assertEqual({(r[2], r[3]) for r in targets}, {(name, row + 4) for name, row, *_ in layouts})
        book.close()
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_baseline_header_errors_include_positions_and_configured_range(self):
        self.write_source([["中文", "번역"]])
        path = self.write_baseline(["中文"])
        for cells, limit, message in (
            ({'G1': None, 'B3': 'Text'}, 2, '第 1～2 行未找到 Text'),
            ({'G1': 'Text', 'B1': 'Text'}, 10, '多个 Text：B1、G1'),
            ({'G1': 'Text', 'B3': 'Text'}, 10, '多个 Text：G1、B3'),
            ({'G1': 'Text', 'XFD1': 'TextTrans'}, 10, '不足 16 列'),
        ):
            with self.subTest(cells=cells):
                book = Workbook()
                for coordinate, value in cells.items():
                    book.active[coordinate] = value
                book.active['G4'] = '中文'
                book.save(path)
                book.close()
                with self.assertRaisesRegex(ValueError, message):
                    match_folder(scan_folder(self.source, limit), self.baseline, self.output)
                self.assertEqual(list(self.output.iterdir()), [])
        # A header on the last permitted row is accepted, even without TextTrans.
        book = Workbook()
        book.active['B3'] = 'Text'
        book.active['B4'] = '中文'
        book.save(path)
        book.close()
        result = match_folder(scan_folder(self.source, 3), self.baseline, self.output)
        book = load_workbook(result.destination / '基准匹配结果/story.xlsx')
        self.assertEqual(book.active['R3'].value, '匹配中文1')
        self.assertEqual(book.active['AD4'].value, '번역')
        book.close()

    def test_shifted_results_are_protected_and_cannot_be_used_as_input(self):
        self.write_source([["中文", "번역"]])
        path = self.write_baseline(["中文"])
        book = load_workbook(path)
        book.active['AZ1'] = 'TextTrans'
        book.active['BA2'] = '原内容'
        book.save(path)
        book.close()
        report = scan_folder(self.source)
        with self.assertRaisesRegex(ValueError, 'BA～BP 列已有内容'):
            match_folder(report, self.baseline, self.output)
        book = load_workbook(path)
        book.active['BA2'] = None
        book.save(path)
        book.close()
        result = match_folder(report, self.baseline, self.output)
        path.write_bytes((result.destination / '基准匹配结果/story.xlsx').read_bytes())
        with self.assertRaisesRegex(ValueError, '已有检索结果'):
            match_folder(report, self.baseline, self.output)
        # Formula rejection follows the identified Text column, too.
        book = Workbook()
        book.active['C1'] = 'Text'
        book.active['C2'] = '=1+1'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, '基准中文不能是公式.*C2'):
            match_folder(report, self.baseline, self.output)

    def test_output_boundary_newlines_follow_each_baseline_row(self):
        def preserve_cell_text(path, cells):
            # Without lxml, openpyxl's fixture writer emits literal CR. XML
            # readers normalize it, so encode CR as Excel's character reference.
            with ZipFile(path) as archive:
                parts = [(item, archive.read(item)) for item in archive.infolist()]
            with ZipFile(path, "w") as archive:
                for item, data in parts:
                    if item.filename == "xl/worksheets/sheet1.xml":
                        root = ET.fromstring(data)
                        for cell in root.findall(".//{*}c"):
                            if cell.get("r") in cells:
                                cell.find("{*}is/{*}t").text = cells[cell.get("r")]
                        data = ET.tostring(root, encoding="utf-8").replace(b"\r", b"&#13;")
                    archive.writestr(item, data)

        chinese = " \n\n\t甲\n乙 \t\n\n "
        korean = " \n\n\t번역\n둘 \t\n\n "
        source = self.write_source([[chinese, korean], ["无译文", None],
                                    ["\n三组\n", "\n단일\n"], ["\r三组\r", "\r단일\r"], ["\r\n三组\r\n", "\r\n단일\r\n"]])
        baselines = ["甲\n乙", "\r\n甲\n乙", "甲\n乙\r\n", " \r\n甲\n乙\n ", "无译文", "三组"]
        baseline = self.write_baseline(baselines)
        preserve_cell_text(source, {"A5": "\r三组\r", "B5": "\r단일\r",
                                    "A6": "\r\n三组\r\n", "B6": "\r\n단일\r\n"})
        preserve_cell_text(baseline, {f"G{row}": text for row, text in enumerate(baselines, 2)})
        # A fourth source occurrence suppresses R:Z, but AD must still be cleaned.
        many = self.write_source([["\r\n多条\r\n", "\r\n여러\r\n"]] * 4, name="多条.xlsx")
        preserve_cell_text(many, {f"{column}{row}": text for row in range(2, 6)
                                 for column, text in (("A", "\r\n多条\r\n"), ("B", "\r\n여러\r\n"))})
        other = self.write_baseline(["多条", "\r\n多条\r\n"], name="另一个.xlsx")
        preserve_cell_text(other, {"G3": "\r\n多条\r\n"})
        original = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline, other)}
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_rows, result.matched_rows, result.recommended_rows, result.source_rows,
                          result.source_matched, result.match_pairs, result.conflict_texts), (8, 8, 7, 9, 9, 16, 0))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        clean_chinese, clean_korean = " \t甲\n乙 \t ", " \t번역\n둘 \t "
        expected = [(clean_chinese, clean_korean),
                    (" \n\n\t甲\n乙 \t ", " \n\n\t번역\n둘 \t "),
                    (" \t甲\n乙 \t\n\n ", " \t번역\n둘 \t\n\n "),
                    (chinese, korean)]
        for row, (cn, ko) in enumerate(expected, 2):
            self.assertEqual(book.active.cell(row, 7).value, baselines[row - 2])
            self.assertEqual(book.active.cell(row, 18).value, cn)
            self.assertEqual(book.active.cell(row, 19).value, ko)
            self.assertEqual(book.active.cell(row, 30).value, ko)
            self.assertEqual(book.active.cell(row, 33).value, "唯一译文")
        self.assertIsNone(book.active["S6"].value)
        self.assertIsNone(book.active["AD6"].value)
        self.assertEqual([book.active.cell(7,c).value for c in (18,21,24)], ["三组"] * 3)
        self.assertEqual([book.active.cell(7,c).value for c in (19,22,25,30)], ["단일"] * 4)
        book.close()
        book = load_workbook(result.destination / "基准匹配结果/另一个.xlsx")
        self.assertIsNone(book.active["R2"].value)
        self.assertEqual(book.active["AD2"].value, "여러")
        self.assertEqual(book.active["AA2"].value, 4)
        self.assertEqual(book.active["G3"].value, "\r\n多条\r\n")
        self.assertEqual(book.active["AD3"].value, "\r\n여러\r\n")
        book.close()
        # Evidence remains raw, and none of the original workbooks is saved.
        book = load_workbook(result.destination / "匹配明细.xlsx")
        self.assertTrue(all(r[8:10] == (chinese, korean) for r in list(book["匹配明细"].values)[1:] if r[4] == source.name and r[7] == 2))
        book.close()
        self.assertEqual(original, {p: hashlib.sha256(p.read_bytes()).digest() for p in original})
        self.assertEqual(align_boundary_newlines("\n=文字\n", False, False), "=文字")
        self.assertEqual(align_boundary_newlines(" \r\n\t ", False, False), " \t ")
        self.assertEqual(align_boundary_newlines("原文无换行", True, True), "\n原文无换行\n")
        self.assertEqual(align_boundary_newlines(" \t正文\t ", True, False), "\n \t正文\t ")
        self.assertEqual(align_boundary_newlines(" \r\n正文\n\n ", True, True), " \r\n正文\n\n ")
        self.assertEqual(align_boundary_newlines("正文\n\n", True, False), "\n正文")
        self.assertEqual(align_boundary_newlines("\n正文", False, True), "正文\n")
        self.assertEqual(align_boundary_newlines(None, True, True), None)
        self.assertEqual(align_boundary_newlines(" \t ", True, True), " \t ")

    def test_missing_boundary_newlines_are_added_per_baseline(self):
        self.write_source([["中文\n正文", "번역\n본문"]] * 3 + [["缺译", None], ["未推荐", "甲译文"], ["未推荐", "乙译文"]])
        self.write_source([["多条", "여러"]] * 4, name="多条.xlsx")
        baselines = ["中文\n正文", "\n中文\n正文", "中文\n正文\n", "\n\n中文\n正文\n\n", "\n缺译\n", "\n未命中\n", "\n未推荐\n", "\n多条\n"]
        self.write_baseline(baselines)
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_rows, result.matched_rows, result.recommended_rows), (8, 7, 5))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        expected = [("中文\n正文", "번역\n본문"), ("\n中文\n正文", "\n번역\n본문"),
                    ("中文\n正文\n", "번역\n본문\n"), ("\n中文\n正文\n", "\n번역\n본문\n")]
        for row, (chinese, korean) in enumerate(expected, 2):
            self.assertEqual(book.active.cell(row, 7).value, baselines[row - 2])
            self.assertEqual([book.active.cell(row, c).value for c in (18, 21, 24)], [chinese] * 3)
            self.assertEqual([book.active.cell(row, c).value for c in (19, 22, 25, 30)], [korean] * 4)
        for row in (6, 7, 8):
            self.assertIsNone(book.active.cell(row, 30).value)
        self.assertIsNone(book.active['S6'].value)
        self.assertEqual(book.active['AG8'].value, '待确认')
        self.assertIsNone(book.active['R9'].value)
        self.assertEqual(book.active['AD9'].value, '\n여러\n')
        self.assertEqual(book.active['AA9'].value, 4)
        book.close()
        book = load_workbook(result.destination / "匹配明细.xlsx")
        self.assertEqual([(row[8], row[9]) for row in list(book['匹配明细'].values)[1:] if row[8] == '中文\n正文'],
                         [('中文\n正文', '번역\n본문')] * 12)
        book.close()

    def test_counts_details_stats_unicode_and_originals_unchanged(self):
        rows = [["一条", "번역1"], ["  两条  ", "번역2"], ["两条", "번역3"],
                *[["三条", f"번역{i}"] for i in range(3)],
                *[["四条", f"번역{i}"] for i in range(4)],
                ["无基准", "외부"], [None, "空中文"], ["   ", "空白"], [" 空译文 ", None], ["内部 空格", "구별"]]
        source = self.write_source(rows)
        baseline = self.write_baseline([" 一条 ", "两条", "三条", "四条", "未找到", "空译文", "内部空格"])
        book = load_workbook(source)
        book.active.row_dimensions[2].hidden = True
        book.active.auto_filter.ref = "A1:B15"
        hidden = book.create_sheet("隐藏")
        hidden.append(["台词", "韩文"])
        hidden.append(["独有隐藏", "숨김"])
        hidden.sheet_state = "hidden"
        book.save(source)
        book.close()
        self.write_baseline(["独有隐藏", "一条"], name="char.xlsx")
        original = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, *self.baseline.iterdir())}
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_rows, result.matched_rows), (9, 7))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        sheet = book.active
        self.assertEqual([sheet.cell(i, 27).value for i in range(2, 9)], [1, 2, 3, 4, 0, 1, 0])
        self.assertEqual(sheet["R2"].value, "一条")
        self.assertEqual(sheet["S2"].value, "번역1")
        self.assertIn("行号：2", sheet["T2"].value)
        self.assertIn("Sheet：正文", sheet["T2"].value)
        self.assertEqual(sheet["R3"].value, "  两条  ")
        self.assertEqual(sheet["U3"].value, "两条")
        self.assertEqual(sheet["X4"].value, "三条")
        self.assertTrue(all(sheet.cell(5, c).value is None for c in range(18, 27)))
        self.assertIsNone(sheet["S7"].value)
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertTrue(sheet["A2"].font.bold)
        self.assertEqual(sheet.column_dimensions["G"].width, 40)
        book.close()
        stats = load_workbook(result.destination / "源文件匹配统计.xlsx")
        row = list(stats["文件统计"].values)[1]
        self.assertEqual(row[2:7], (2, 0, 14, 12, 2))
        self.assertAlmostEqual(row[7], 12 / 14)
        self.assertEqual(stats["文件统计"]["H2"].number_format, "0.00%")
        stats.close()
        self.assertEqual(original, {p: hashlib.sha256(p.read_bytes()).digest() for p in original})
        self.assertFalse(any(self.output.glob("匹配临时_*")))
        self.assertFalse(any(result.destination.glob("*.sqlite3")))

    def test_ignored_and_preexisting_abnormal_sheets_never_count(self):
        path = self.write_source([["命中", "有效"]])
        book = load_workbook(path)
        abnormal = book.create_sheet("异常")
        abnormal.append(["台词"])
        abnormal.append(["命中", "不应输出"])
        ignored = book.create_sheet("忽略")
        ignored.append(["台词", "韩文"])
        ignored.append(["命中", "不应输出"])
        book.save(path)
        book.close()
        record = {sheet_key(path, "忽略")}
        report = scan_folder(self.source, 1, ignored_sheets=record)
        # A sheet fixed after checking remains excluded until the user rechecks.
        book = load_workbook(path)
        book["异常"]["B1"] = "韩文"
        book.save(path)
        book.close()
        self.write_baseline(["命中"])
        result = match_folder(report, self.baseline, self.output, ignored_sheets=record)
        self.assertEqual((result.source_rows, result.source_matched, result.skipped), (1, 1, 2))
        self.assertEqual([r.sheet for r in result.check_report.exceptions], ["异常"])
        stats = load_workbook(result.destination / "源文件匹配统计.xlsx")
        self.assertEqual({r[2] for r in list(stats["跳过清单"].values)[1:]}, {"异常", "忽略"})
        stats.close()

    def test_changed_header_excluded_and_body_failure_rolls_back(self):
        path = self.write_source([["命中", "错误"]], title="变更")
        book = load_workbook(path)
        for title, rows in (("失败", [["命中", "半途数据"], ["命中", "=1+1"]]), ("合法", [["命中", "正确"]])):
            sheet = book.create_sheet(title)
            sheet.append(["台词", "韩文"])
            for row in rows:
                sheet.append(row)
        book.save(path)
        book.close()
        report = scan_folder(self.source, 1)
        book = load_workbook(path)
        book["变更"]["B1"] = "新未知表头"
        book.save(path)
        book.close()
        self.write_baseline(["命中"])
        result = match_folder(report, self.baseline, self.output)
        self.assertEqual((result.source_rows, result.source_matched, result.skipped), (1, 1, 2))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertEqual(book.active["AA2"].value, 1)
        self.assertEqual(book.active["S2"].value, "正确")
        book.close()

    def test_literal_formula_like_text_not_executed(self):
        path = self.write_source([["=文字", "=번역"]])
        book = load_workbook(path)
        book.active["A2"].data_type = book.active["B2"].data_type = "s"
        book.save(path)
        book.close()
        baseline = self.write_baseline(["=文字"])
        book = load_workbook(baseline)
        book.active["G2"].data_type = "s"
        book.save(baseline)
        book.close()
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertEqual(book.active["S2"].value, "=번역")
        self.assertEqual(book.active["S2"].data_type, "s")
        book.close()

    def test_scan_switches_control_matching_scope_and_denominator(self):
        path = self.write_source([["命中", "可见"], ["命中", "隐藏行"], ["未匹配", "可见"]])
        book = load_workbook(path)
        book.active.row_dimensions[3].hidden = True
        hidden = book.create_sheet("隐藏表")
        hidden.append(["台词", "韩文"])
        hidden.append(["命中", "隐藏表译文"])
        hidden.sheet_state = "veryHidden"
        book.save(path)
        book.close()
        nested = self.source / "子目录"
        nested.mkdir()
        self.write_source([["命中", "子目录译文"]], name="子目录/嵌套.xlsx")
        self.write_baseline(["命中"])
        report = scan_folder(self.source, 1, recursive=False, include_hidden_sheets=False, include_hidden_rows=False)
        self.assertEqual(report.files, 1)
        self.assertEqual(report.passed_sheets, 1)
        self.assertEqual(len(report.skipped_sheets), 1)
        result = match_folder(report, self.baseline, self.output)
        self.assertEqual((result.source_rows, result.source_matched, result.skipped), (2, 1, 1))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertEqual(book.active["AA2"].value, 1)
        self.assertEqual(book.active["S2"].value, "可见")
        book.close()
        stats = load_workbook(result.destination / "源文件匹配统计.xlsx")
        self.assertEqual(stats["文件统计"]["H2"].value, .5)
        stats.close()

    def test_deprecated_rows_are_excluded_from_matching_context_and_details(self):
        path = self.write_source([['前1','译文'], ['废弃独有','旧译文'], ['前2','译文'],
                                  ['前3','译文'], ['目标','正确'], ['非废弃说明','仍有效']])
        book = load_workbook(path)
        book.active['AZ1'] = ' 特殊说明 '
        book.active['AZ3'] = ' \n废弃\t '
        book.active['AZ7'] = '非废弃'
        book.save(path)
        book.close()
        obsolete = self.write_source([['目标','错误']], name='旧.xlsx')
        book = load_workbook(obsolete)
        book.active['C1'] = '特殊说明'
        book.active['C2'] = '废弃'
        book.save(obsolete)
        book.close()
        self.write_baseline(['前1','前2','前3','目标','非废弃说明'])
        before = {p: hashlib.sha256(p.read_bytes()).digest() for p in (path, obsolete)}
        for include, expected in ((False, (5,5,0)), (True, (7,6,1))):
            with self.subTest(include=include):
                report = scan_folder(self.source, 1, include_deprecated=include)
                self.assertEqual(report.results[0].notes_column, 52)
                result = match_folder(report, self.baseline, self.output)
                self.assertEqual(result.check_report.include_deprecated, include)
                self.assertEqual((result.source_rows,result.source_matched,result.conflict_texts), expected)
                book = load_workbook(result.destination / '匹配明细.xlsx')
                matched = list(book['匹配明细'].values)[1:]
                target = [r for r in matched if r[8] == '目标']
                if not include:
                    self.assertEqual({r[9] for r in target}, {'正确'})
                    self.assertEqual(target[0][12:14], ('前1句 [行4] 前3\n前2句 [行3] 前2\n前3句 [行2] 前1', '后1句 [行6] 非废弃说明'))
                    self.assertNotIn('废弃独有', str(list(book['未推荐明细'].values)))
                    self.assertNotIn('旧.xlsx', str(matched))
                else:
                    self.assertEqual({r[9] for r in target}, {'正确','错误'})
                book.close()
        self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).digest() for p in before})

    def test_deprecated_marker_validation_and_formula_rows(self):
        path = self.write_source([['=1+1','=2+2'], ['有效','번역']])
        book = load_workbook(path)
        book.active['C1'] = '特殊说明'
        book.active['C2'] = '废弃'
        book.save(path)
        book.close()
        self.write_baseline(['有效'])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual(result.source_rows, 1)
        book = load_workbook(path)
        book.active['D1'] = '特殊说明'
        book.save(path)
        book.close()
        self.assertIn('多个特殊说明列', scan_folder(self.source, 1).exceptions[0].issue)
        self.assertEqual(scan_folder(self.source, 1, include_deprecated=True).passed_sheets, 1)

    def test_plain_tag_fallback_preserves_values_exact_priority_and_statistics(self):
        first = self.write_source([['这是基准文件的原文','<color=#3266b4>번역</color>'],
                                   ['{playername}你好\\n世界','이름{playername}\\n원문'],
                                   ['{playername}<color=red></color>','标签空句'],
                                   ['价格100，OK','숫자'], ['精确优先','原推荐']])
        second = self.write_source([['<b>精确优先</b>','其他译文'], ['<b>废弃候选</b>','废弃译文']], name='B.xlsx')
        book = load_workbook(second)
        book.active['C1'] = '特殊说明'
        book.active['C3'] = '废弃'
        book.save(second)
        book.close()
        baseline = self.write_baseline(['<color=#3266b4>这是基准文件的原文</color>', '你好世界',
                                        '{other}<size=12></size>', '价格101，OK', '精确优先', '废弃候选'])
        before = {p: hashlib.sha256(p.read_bytes()).digest() for p in (first,second,baseline)}
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_rows,result.matched_rows,result.recommended_rows,result.source_matched,result.match_pairs), (6,3,3,3,3))
        book = load_workbook(result.destination / '基准匹配结果/story.xlsx')
        sheet = book.active
        for row, ko in ((2,'<color=#3266b4>번역</color>'), (3,'이름{playername}\\n원문')):
            self.assertEqual(sheet.cell(row,30).value, ko)
            self.assertIn('仅纯文本匹配，仍需复核标签', sheet.cell(row,32).value)
            self.assertEqual(sheet.cell(row,27).value, 1)
        self.assertEqual(sheet['AD6'].value, '原推荐')
        self.assertNotIn('仅纯文本匹配', sheet['AF6'].value)
        self.assertEqual((sheet['AA4'].value,sheet['AA5'].value), (0,0))
        self.assertIsNone(sheet['AD7'].value)
        book.close()
        book = load_workbook(result.destination / '匹配明细.xlsx')
        rows = list(book['匹配明细'].values)
        self.assertEqual(rows[0][23], '匹配方式')
        self.assertEqual(rows[0][-1], '上下文查询统计')
        self.assertEqual({r[8]:r[23] for r in rows[1:]},
                         {'这是基准文件的原文':'纯文本补匹配','{playername}你好\\n世界':'纯文本补匹配','精确优先':'全文精确匹配'})
        unused = list(book['未推荐明细'].values)[1:]
        self.assertFalse(any(r[4] in ('这是基准文件的原文','{playername}你好\\n世界') for r in unused))
        self.assertFalse(any('废弃候选' in str(r) for r in unused))
        book.close()
        book = load_workbook(result.destination / '基准文件匹配统计.xlsx')
        self.assertEqual(list(book['文件统计'].values)[1][2:11], (6,3,3,3,3,0,0,0,0))
        book.close()
        book = load_workbook(result.destination / '源文件匹配统计.xlsx')
        self.assertEqual(sum(r[5] for r in list(book['文件统计'].values)[1:]), 3)
        book.close()
        self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).digest() for p in before})
        self.assertEqual(plain_key('{playername}<color=#3266b4>你好</color>\\n\r\n世界'), '你好世界')
        self.assertEqual(plain_key('价格100，OK < 200'), '价格100，OK < 200')
        self.assertEqual(plain_key('{name}<color=red></color>\\n'), '')
        included = match_folder(scan_folder(self.source, 1, include_deprecated=True), self.baseline, self.output)
        self.assertEqual((included.matched_rows,included.recommended_rows), (4,4))

    def test_plain_fallback_can_reuse_source_recommended_by_another_baseline(self):
        self.write_source([['共同句','번역']])
        self.write_baseline(['共同句'], name='A.xlsx')
        self.write_baseline(['<color=red>共同句</color>'], name='B.xlsx')
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_rows,result.matched_rows,result.recommended_rows,result.source_matched,result.match_pairs), (2,2,2,1,2))
        book = load_workbook(result.destination / '匹配明细.xlsx')
        rows = list(book['匹配明细'].values)[1:]
        self.assertTrue(all(r[20]=='是' for r in rows))
        self.assertEqual({r[0]:r[23] for r in rows}, {'A.xlsx':'全文精确匹配','B.xlsx':'纯文本补匹配'})
        self.assertEqual(list(book['未推荐明细'].values)[1:], [])
        book.close()

    def test_trailing_notes_only_retry_unused_sources_preserve_recommendations_and_conflicts(self):
        source = self.write_source([
            ['保持', '기존'], ['保持（新说明）', '다른'],
            ['<color=red>你好</color>（小声）(动作)', '안녕 (작게)'],
            ['中（注）间', '중'], ['完整（备注）', 'full'], ['缺译（说明）', None],
            ['冲突（甲）', '번역1'], ['冲突（乙）', '번역2'],
        ])
        first = self.write_baseline(['保持', '你好', '中间', '完整（备注）', '完整', '缺译', '冲突'])
        second = self.write_baseline(['你好'], 'B.xlsx')
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, first, second)}
        result = match_folder(scan_folder(self.source), self.baseline, self.output)
        self.assertEqual((result.baseline_rows, result.matched_rows, result.recommended_rows,
                          result.source_matched, result.match_pairs), (8, 6, 4, 7, 8))
        book = load_workbook(result.destination / '基准匹配结果/story.xlsx')
        self.assertEqual(book.active['AD2'].value, '기존')
        self.assertIn('保留先前匹配阶段推荐', book.active['AF2'].value)
        self.assertEqual(book.active['AD3'].value, '안녕 (작게)')
        self.assertIn('忽略源中文句尾括号', book.active['AF3'].value)
        self.assertIn('仍需复核标签', book.active['AF3'].value)
        self.assertEqual(book.active['AG7'].value, '全部缺译')
        self.assertEqual(book.active['AG8'].value, '待确认')
        book.close()
        book = load_workbook(result.destination / '匹配明细.xlsx')
        rows = list(book['匹配明细'].values)[1:]
        self.assertEqual(sum(r[23] == TAIL_MATCH for r in rows), 6)
        unused = {r[4]: r[-1] for r in list(book['未推荐明细'].values)[1:]}
        self.assertIn('找到了但不符合推荐规则', unused['保持（新说明）'])
        self.assertIn('找不到', unused['中（注）间'])
        book.close()
        self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).digest() for p in hashes})
        for original, expected in [('甲（乙(丙)）', '甲'), ('甲（乙）丁', '甲（乙）丁'),
                                   ('甲（乙)', '甲（乙)'), ('甲（未闭合(注)', '甲（未闭合(注)'),
                                   ('甲）乙（注）', '甲）乙（注）'), ('（说明）', '')]:
            self.assertEqual(trailing_note_key(original), expected)

    def test_plain_context_resolves_exact_conflict_and_skips_tag_only_neighbors(self):
        a = ['<b>前1</b>','{name}','<b>前2</b>','<b>前3</b>','争议',
             '<i>后1</i>','{empty}','<i>后2</i>','<i>后3</i>']
        self.write_source([[text,'正确' if text=='争议' else '译文'] for text in a], name='A.xlsx')
        self.write_source([[text,'错误' if text=='争议' else '译文'] for text in
                           ['别前1','别前2','别前3','争议','别后1','别后2','别后3']], name='B.xlsx')
        self.write_baseline(['前1','前2','前3','争议','后1','后2','后3'])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.matched_rows,result.recommended_rows,result.match_pairs), (7,7,8))
        book = load_workbook(result.destination / '基准匹配结果/story.xlsx')
        self.assertEqual(book.active['AD5'].value, '正确')
        self.assertEqual(book.active['AG5'].value, '上下文推荐')
        self.assertIn('仅纯文本匹配，仍需复核标签', book.active['AF5'].value)
        self.assertIn('前3/3句一致，后3/3句一致', book.active['AF5'].value)
        book.close()
        book = load_workbook(result.destination / '匹配明细.xlsx')
        chosen = next(r for r in list(book['匹配明细'].values)[1:] if r[8]=='争议' and r[20]=='是')
        self.assertEqual(chosen[14], '前1句 [行5] 前3\n前2句 [行4] 前2\n前3句 [行2] 前1')
        self.assertEqual(chosen[15], '后1句 [行7] 后1\n后2句 [行9] 后2\n后3句 [行10] 后3')
        self.assertEqual(chosen[16:20], (3,3,2,6))
        self.assertEqual(chosen[23], '纯文本补匹配')
        self.assertTrue(any(r[0]=='B.xlsx' and r[4]=='争议' and r[6]=='是' for r in list(book['未推荐明细'].values)[1:]))
        book.close()

    def test_plain_conflicts_and_empty_translations_still_require_recommendation_rules(self):
        self.write_source([['<b>争议</b>','갑'], ['<b>缺译</b>',None], ['<b>唯一</b>','译文']])
        self.write_source([['<i>争议</i>','을']], name='B.xlsx')
        self.write_baseline(['争议','缺译','唯一'])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.matched_rows,result.recommended_rows,result.match_pairs), (3,1,4))
        book = load_workbook(result.destination / '基准匹配结果/story.xlsx')
        self.assertIsNone(book.active['AD2'].value)
        self.assertEqual(book.active['AG2'].value, '待确认')
        self.assertEqual(book.active['AA2'].value, 2)
        self.assertEqual(book.active['AC2'].value, '是')
        self.assertEqual(book.active['AG3'].value, '全部缺译')
        self.assertIsNone(book.active['AD3'].value)
        self.assertEqual(book.active['AD4'].value, '译文')
        book.close()
        book = load_workbook(result.destination / '匹配明细.xlsx')
        unused = list(book['未推荐明细'].values)[1:]
        self.assertEqual(len(unused), 3)
        self.assertTrue(all(r[6]=='是' and r[8]=='基准文件中找到了但不符合推荐规则' for r in unused))
        self.assertEqual(list(book['未匹配明细'].values)[1:], [])
        book.close()

    def test_hidden_header_row_is_not_used_when_disabled(self):
        path = self.write_source([["中文", "번역"]])
        book = load_workbook(path)
        book.active.row_dimensions[1].hidden = True
        book.save(path)
        book.close()
        self.assertEqual(scan_folder(self.source, 1).passed_sheets, 1)
        report = scan_folder(self.source, 1, include_hidden_rows=False)
        self.assertEqual(report.passed_sheets, 0)
        self.assertEqual(len(report.exceptions), 1)
        self.assertIn("排除隐藏行：1", report.exceptions[0].details)

    def test_validation_existing_output_columns_and_repeat_outputs(self):
        self.write_source([["中文", "한국어"]])
        baseline = self.write_baseline(["中文"])
        report = scan_folder(self.source, 1)
        for source, output in ((self.source, self.output), (self.baseline, self.source / "输出")):
            with self.assertRaises(ValueError):
                match_folder(report, source, output)
        first = match_folder(report, self.baseline, self.output)
        second = match_folder(report, self.baseline, self.output)
        self.assertNotEqual(first.destination, second.destination)
        book = load_workbook(baseline)
        book.active["R2"] = "原有内容"
        book.save(baseline)
        book.close()
        before = sorted(self.output.iterdir())
        with self.assertRaisesRegex(ValueError, "已有内容"):
            match_folder(report, self.baseline, self.output)
        self.assertEqual(sorted(self.output.iterdir()), before)

    def test_streaming_preserves_package_parts_prefixes_and_bad_dimensions(self):
        self.write_source([["中文", "한국어"]])
        baseline = self.write_baseline(["中文"])
        book = load_workbook(baseline)
        book.active["R2"].font = Font(italic=True)  # Empty styled cell must not duplicate the new R2.
        book.save(baseline)
        book.close()

        with ZipFile(baseline) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        xml = parts['xl/worksheets/sheet1.xml'].decode()
        xml = re.sub(r'<dimension\b[^>]*/>', '<dimension ref="A1"/>', xml)
        xml = xml.replace('xmlns=', 'xmlns:m=', 1)
        xml = re.sub(r'<(/?)([A-Za-z][A-Za-z0-9]*)(?=[\s/>])', r'<\1m:\2', xml)
        xml = xml.replace('<m:worksheet ', '<m:worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac" ', 1)
        parts['xl/worksheets/sheet1.xml'] = xml.encode()
        with ZipFile(baseline, 'w') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        target = result.destination / '基准匹配结果/story.xlsx'
        with ZipFile(target) as archive:
            for name, data in parts.items():
                if name != 'xl/worksheets/sheet1.xml':
                    self.assertEqual(archive.read(name), data)
            xml = archive.read('xl/worksheets/sheet1.xml')
            self.assertIn(b'mc:Ignorable="x14ac"', xml)
            ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            self.assertEqual(len(ET.fromstring(xml).findall('.//m:c[@r="R2"]', ns)), 1)
        book = load_workbook(target, read_only=True)
        self.assertEqual(book.active['R2'].value, '中文')
        self.assertEqual(book.active['AA2'].value, 1)
        self.assertEqual(book.active['A2'].value, '=1+1')
        book.close()

    def test_complete_details_missing_and_conflicts_keep_all_candidates(self):
        path = self.write_source([
            ["重复","韩语甲"], ["重复",None], ["首尾一致"," 번역 "], ["首尾一致","번역"],
            ["单一缺译"," \t "], ["四条","韩1"], ["四条","韩1"], ["四条","韩2"], ["四条",None],
            ["不在基准",None], ["不在基准","例1"], ["不在基准","例2"],
        ], name="A.xlsx")
        self.write_source([["重复","韩语乙"]], name="B.xlsx")
        book = load_workbook(path)
        ignored = book.create_sheet("忽略")
        ignored.append(["台词","韩文"])
        ignored.append(["重复","不应算冲突"])
        ignored.append(["单一缺译",None])
        invalid = book.create_sheet("异常")
        invalid.append(["台词"])
        invalid.append(["重复","不应算冲突"])
        book.save(path)
        book.close()
        exclusions = {sheet_key(path, "忽略")}
        self.write_baseline(["重复","首尾一致","单一缺译","四条","未找到"])
        self.write_baseline(["重复"], name="char.xlsx")
        report = scan_folder(self.source, 1, ignored_sheets=exclusions)
        result = match_folder(report, self.baseline, self.output, ignored_sheets=exclusions)
        self.assertEqual((result.source_rows,result.source_matched,result.missing_translations,result.conflict_texts,result.match_pairs), (13,10,4,3,13))
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        sheet = book.active
        self.assertEqual([(sheet.cell(r,27).value,sheet.cell(r,28).value,sheet.cell(r,29).value) for r in range(2,7)],
                         [(3,1,"是"),(2,0,"否"),(1,1,"否"),(4,1,"是"),(0,0,"未匹配")])
        self.assertTrue(all(sheet.cell(5,c).value is None for c in range(18,27)))
        book.close()
        book = load_workbook(result.destination / "匹配明细.xlsx")
        matched = list(book["匹配明细"].values)[1:]
        self.assertEqual(len(matched),13)
        repeat = [r for r in matched if r[8]=="重复"]
        self.assertEqual(len(repeat),6)
        self.assertEqual({r[0] for r in repeat},{"story.xlsx","char.xlsx"})
        self.assertTrue(all(r[11]=="是" for r in repeat))
        self.assertFalse(any(r[6] in ("忽略","异常") for r in matched))
        unmatched = list(book["未匹配明细"].values)[1:]
        self.assertEqual(len(unmatched),3)
        self.assertEqual({r[4] for r in unmatched},{"不在基准"})
        self.assertEqual({r[3] for r in unmatched},{11,12,13})
        missing = list(book["缺译明细"].values)[1:]
        self.assertEqual(len(missing),4)
        self.assertEqual(sum(r[6]=="否" for r in missing),1)
        conflicts = list(book["译文冲突"].values)[1:]
        self.assertEqual({r[4] for r in conflicts},{"重复","四条","不在基准"})
        self.assertTrue(all(r[8]==2 for r in conflicts))
        self.assertEqual(len(conflicts),10)
        self.assertEqual(book["匹配明细"].freeze_panes,"A2")
        book.close()
        book = load_workbook(result.destination / "源文件匹配统计.xlsx")
        rows = list(book["文件统计"].values)[1:]
        self.assertEqual([(r[0],r[9],r[10]) for r in rows],[("A.xlsx",4,3),("B.xlsx",0,1)])
        book.close()

    def test_details_split_at_sheet_limit_without_losing_rows(self):
        book = Workbook(write_only=True)
        count = write_detail_sheet(book,"明细",["中文"],[("=原文",),("第二行",),("第三行",),("第四行",),("第五行",)],[30],row_limit=3)
        path = self.root / "分页.xlsx"
        book.save(path)
        book.close()
        self.assertEqual(count,5)
        book = load_workbook(path)
        self.assertEqual(book.sheetnames,["明细","明细_2","明细_3"])
        self.assertEqual([row[0] for sheet in book for row in list(sheet.values)[1:]], ["=原文","第二行","第三行","第四行","第五行"])
        self.assertEqual(book["明细"]["A2"].data_type,"s")
        self.assertTrue(all(s.max_row<=3 for s in book))
        book.close()

    def test_unrecommended_details_classify_each_source_row_once(self):
        path = self.write_source([["唯一", "번역"], [" 唯一 ", "번역"], ["缺译", None],
                                  ["冲突", "甲"], ["空白", None], ["只在源里", " 원문\n"],
                                  ["=未匹配", "=번역"], ["排除隐藏行", "不参与"]], name="A.xlsx")
        self.write_source([["冲突", "乙"], ["空白", " \t "]], name="B.xlsx")
        book = load_workbook(path)
        book.active['A8'].data_type = book.active['B8'].data_type = 's'
        book.active.row_dimensions[9].hidden = True
        for name in ('忽略', '异常', '读取失败', '隐藏表'):
            sheet = book.create_sheet(name)
            sheet.append(['台词', '未知' if name == '异常' else '韩文'])
            sheet.append(['不应列出', '不参与'])
            if name == '读取失败':
                sheet.append(['失败公式', '=1+1'])
            if name == '隐藏表':
                sheet.sheet_state = 'hidden'
        book.save(path)
        book.close()
        ignored = {sheet_key(path, '忽略')}
        # The same chosen source is reused by multiple baselines, never duplicated here.
        self.write_baseline(['唯一'], name='A.xlsx')
        self.write_baseline(['唯一'], name='B.xlsx')
        self.write_baseline(['缺译', '冲突', '空白'], name='C.xlsx')
        report = scan_folder(self.source, 1, ignored_sheets=ignored,
                             include_hidden_sheets=False, include_hidden_rows=False)
        result = match_folder(report, self.baseline, self.output, ignored_sheets=ignored)
        self.assertEqual((result.source_rows, result.source_matched, result.recommended_rows), (9, 7, 2))
        book = load_workbook(result.destination / '匹配明细.xlsx')
        sheet = book['未推荐明细']
        records = list(sheet.values)
        self.assertEqual(records[0], ('源文件名','源完整路径','源 Sheet','源行号','源中文','韩语','命中基准','缺译','未推荐原因'))
        rows = {(r[0], r[3]): r for r in records[1:]}
        self.assertEqual(set(rows), {('A.xlsx', i) for i in range(4, 9)} | {('B.xlsx', 2), ('B.xlsx', 3)})
        self.assertEqual(len(records) - 1, len(rows))
        for key, row in rows.items():
            unmatched = key in {('A.xlsx', 7), ('A.xlsx', 8)}
            self.assertEqual(row[6], '否' if unmatched else '是')
            self.assertEqual(row[8], '基准文件中找不到' if unmatched else '基准文件中找到了但不符合推荐规则')
            self.assertEqual(row[2], '正文')
        self.assertNotIn(('A.xlsx', 3), rows)
        self.assertEqual(rows['A.xlsx', 7][5], ' 원문\n')
        self.assertEqual(rows['A.xlsx', 4][7], '是')
        self.assertEqual(rows['B.xlsx', 3][5:8], (' \t ', '是', '是'))
        formula_row = next(cell.row for row in sheet for cell in row if cell.value == '=未匹配')
        self.assertEqual(sheet.cell(formula_row, 5).data_type, 's')
        self.assertEqual(sheet.cell(formula_row, 6).data_type, 's')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:I8')
        book.close()

    def test_unrecommended_details_exclude_recommended_pairs_across_files_and_rounds(self):
        first = self.write_source([
            ['已推荐', '번역'], [' 已推荐 \n', '\n번역 '],
            ['<b>补匹配</b>', '보충'], ['尾注（说明）', '주석'],
            ['未匹配', '미사용'], ['未匹配', '미사용'],
        ], name='A.xlsx')
        second = self.write_source([
            ['已推荐', '번역'], ['<b>补匹配</b>', '보충'], ['尾注（说明）', '주석'],
            ['别的原文', '번역'], ['<b>已推荐</b>', '번역'],
        ], name='B.xlsx', title='重复来源')
        baseline = self.write_baseline(['已推荐', '补匹配', '尾注'])
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in (first, second, baseline)}
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.source_rows, result.source_matched, result.recommended_rows, result.match_pairs),
                         (11, 7, 3, 7))
        book = load_workbook(result.destination / '匹配明细.xlsx')
        unused = list(book['未推荐明细'].values)[1:]
        self.assertEqual([(r[0], r[3], r[4]) for r in unused], [
            ('A.xlsx', 6, '未匹配'), ('A.xlsx', 7, '未匹配'),
            ('B.xlsx', 5, '别的原文'), ('B.xlsx', 6, '<b>已推荐</b>'),
        ])
        # All candidates and their actual selected positions remain traceable.
        matches = list(book['匹配明细'].values)[1:]
        self.assertEqual(len(matches), 7)
        self.assertEqual(sum(r[20] == '是' for r in matches), 3)
        self.assertEqual({r[23] for r in matches}, {'全文精确匹配', '纯文本补匹配', TAIL_MATCH})
        book.close()
        self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).digest() for p in hashes})

    def test_unrecommended_details_empty_still_has_headers(self):
        self.write_source([['中文', '번역']])
        self.write_baseline(['中文'])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        book = load_workbook(result.destination / '匹配明细.xlsx')
        self.assertEqual(list(book['未推荐明细'].values),
                         [('源文件名','源完整路径','源 Sheet','源行号','源中文','韩语','命中基准','缺译','未推荐原因')])
        self.assertEqual(book['未推荐明细'].auto_filter.ref, 'A1:I1')
        book.close()

    def test_context_recommends_per_baseline_and_keeps_full_evidence(self):
        a = ["甲前一", "甲前二", "甲前三", "共同句", "甲后一", "甲后二", "甲后三"]
        b = ["乙前一", "乙前二", "乙前三", "共同句", "乙后一", "乙后二", "乙后三"]
        self.write_source([[t, "韩甲" if t == "共同句" else "译文"] for t in a], name="A.xlsx")
        self.write_source([[t, "韩乙" if t == "共同句" else "译文"] for t in b], name="B.xlsx")
        # One-sided candidate must lose to the candidate matching both sides.
        self.write_source([[t, "干扰译文" if t == "共同句" else "译文"] for t in a[:4] + b[4:]], name="C.xlsx")
        # Empty Korean never wins, even with the best context.
        self.write_source([[t, None if t == "共同句" else "译文"] for t in a], name="D.xlsx")
        self.write_baseline(a, name="A.xlsx")
        (self.baseline / "子目录").mkdir()
        self.write_baseline(b, name="子目录/B.xlsx")
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_files, result.baseline_rows, result.matched_rows, result.recommended_rows), (2, 14, 14, 14))
        for name, expected in (("A.xlsx", "韩甲"), ("子目录/B.xlsx", "韩乙")):
            book = load_workbook(result.destination / "基准匹配结果" / name)
            self.assertEqual(book.active["AD5"].value, expected)
            self.assertEqual(book.active["AG5"].value, "上下文推荐")
            self.assertIn("前3/3句一致，后3/3句一致", book.active["AF5"].value)
            self.assertEqual(book.active["AC5"].value, "是")
            book.close()
        book = load_workbook(result.destination / "匹配明细.xlsx")
        evidence = [dict(zip(next(book["匹配明细"].values), row)) for row in list(book["匹配明细"].values)[1:] if row[8] == "共同句"]
        selected = [r for r in evidence if r["推荐来源"] == "是"]
        self.assertEqual(len(selected), 2)
        self.assertEqual({r["韩语"] for r in selected}, {"韩甲", "韩乙"})
        self.assertTrue(all(r["完整一致侧数"] == 2 and r["一致位置总数"] == 6 for r in selected))
        self.assertIn("前1句 [行4] 甲前三", selected[0]["基准前3句"])
        unused = [r for r in list(book['未推荐明细'].values)[1:] if r[4] == '共同句']
        # A and B are chosen in different baselines; only C and D remain.
        self.assertEqual({r[0] for r in unused}, {'C.xlsx', 'D.xlsx'})
        self.assertTrue(all(r[8] == '基准文件中找到了但不符合推荐规则' for r in unused))
        book.close()
        book = load_workbook(result.destination / "基准文件匹配统计.xlsx")
        for row in list(book["文件统计"].values)[1:]:
            self.assertEqual(row[2:11], (7, 7, 7, 7, 0, 0, 1, 1, 0))
        self.assertEqual(list(book["总体统计"].values)[1][2:11], (14, 14, 14, 14, 0, 0, 2, 2, 0))
        self.assertEqual(book["文件统计"]["L2"].number_format, "0.00%")
        book.close()

    def test_context_ties_boundaries_single_side_and_hidden_rows(self):
        texts = ["前1", "前2", "前3", "争议", "后1", "后2", "后3"]
        self.write_source([[t, "韩甲" if t == "争议" else "同译"] for t in texts], name="A.xlsx")
        self.write_source([[t, "韩乙" if t == "争议" else "同译"] for t in texts], name="B.xlsx")
        self.write_baseline(texts)
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertIsNone(book.active["AD5"].value)
        self.assertEqual(book.active["AG5"].value, "待确认")
        self.assertIn("最高上下文得分", book.active["AF5"].value)
        book.close()
        book = load_workbook(result.destination / "匹配明细.xlsx")
        tied = [r for r in list(book["匹配明细"].values)[1:] if r[8] == "争议"]
        self.assertEqual({(r[20], r[21], r[24]) for r in tied}, {("否", "待确认", "是")})
        book.close()
        # At a boundary, even two identical neighbours are insufficient for a conflict.
        self.write_baseline(["争议", "后1", "后2"])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertEqual(book.active["AG2"].value, "待确认")
        self.assertIn("完整前3句或完整后3句", book.active["AF2"].value)
        book.close()
        # Skip a hidden distractor; the next three nonempty records then match.
        path = self.write_source([["争议", "韩甲"], ["隐藏干扰", "不可参与"], ["后1", "同译"],
                                  [None, "空中文"], ["后2", "同译"], ["后3", "同译"]], name="A.xlsx")
        book = load_workbook(path)
        book.active.row_dimensions[3].hidden = True
        book.save(path)
        book.close()
        self.write_source([["争议", "韩乙"], ["不同1", "同译"], ["不同2", "同译"], ["不同3", "同译"]], name="B.xlsx")
        self.write_baseline(["争议", "后1", "后2", "后3"])
        result = match_folder(scan_folder(self.source, 1, include_hidden_rows=False), self.baseline, self.output)
        book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
        self.assertEqual(book.active["AD2"].value, "韩甲")
        self.assertIn("前0/3句一致，后3/3句一致", book.active["AF2"].value)
        self.assertIn("非完美匹配", book.active["AF2"].value)
        book.close()

    def test_special_context_extends_repeatedly_cleans_information_and_counts_blank_stage(self):
        # Ten special rows require going beyond a fixed 3 + 3 window.
        for plain in (False, True):
            with self.subTest(plain=plain):
                path = self.write_source([])
                book = load_workbook(path)
                sheet = book.active
                sheet.append(["前1", "译文", "", "角色"])
                sheet.append(["前2", "译文", "", "角色"])
                sheet.append(["前3", "译文", "奇波", "角色"])
                sheet.append(["场景说明", "설명", "", ""])
                sheet.append(["嗯？（惊讶）", "응? (놀람)", " 气息 ", "角色"])
                for _ in range(9):
                    sheet.append([None, None, "演出", None])
                sheet.append(["目标", "正确", "", "角色"])  # row 16
                sheet.append(["（半疑惑的气）", "(의문)", "气息", "角色"])
                sheet.append(["镜头移动", None, "演出", ""])
                sheet.append(["后1", "译文", "", "角色"])
                sheet.append(["哈(喘息)", "하", "气息", "角色"])
                sheet.append(["后2", "译文", "演出说明", "角色"])
                sheet.append(["后3", "译文", "", "角色"])
                sheet["C1"], sheet["D1"] = "特殊说明", "说话角色名"
                if plain:
                    for row in (2, 3, 4, 19, 21, 22):
                        sheet.cell(row, 1, f"<b>{sheet.cell(row, 1).value}</b>")
                book.save(path)
                book.close()
                self.write_source([["其他1", "译文"], ["其他2", "译文"], ["其他3", "译文"],
                                   ["目标", "错误"], ["另后1", "译文"]], name="错误.xlsx")
                # One breath also exists in baseline; other special rows are source-only.
                baseline = self.write_baseline(["前1", "前2", "前3", "嗯？（不同说明）",
                                                "目标", "后1", "哈", "后2", "后3"])
                before = {p: hashlib.sha256(p.read_bytes()).digest() for p in (path, baseline)}
                result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
                book = load_workbook(result.destination / "匹配明细.xlsx")
                values = list(book["匹配明细"].values)
                selected = next(dict(zip(values[0], r)) for r in values[1:] if r[8] == "目标" and r[20] == "是")
                self.assertEqual(selected["韩语"], "正确")
                self.assertEqual(selected["完美匹配"], "是")
                self.assertEqual(selected["匹配方式"], "纯文本补匹配" if plain else "全文精确匹配")
                self.assertEqual(selected["上下文查询统计"],
                    "向上：查询13行，演出9行，气息1行，匹配3句\n"
                    "向下：查询6行，演出1行，气息2行，匹配3句\n"
                    "合计（不含当前句）：查询19行，演出10行，气息3行，匹配6句")
                self.assertEqual(selected["来源前3句"], "前1句 [行4] 前3\n前2句 [行3] 前2\n前3句 [行2] 前1")
                self.assertEqual(selected["来源后3句"], "后1句 [行19] 后1\n后2句 [行21] 后2\n后3句 [行22] 后3")
                self.assertIn("同时越过基准行5", selected["上下文跳过位置"])
                self.assertIn("同时越过基准行8", selected["上下文跳过位置"])
                self.assertNotIn("场景说明", str(values))
                excluded = list(book["上下文排除明细"].values)[1:]
                self.assertEqual(len(excluded), 14)
                self.assertTrue(any(r[2] == 5 and "两列均为空" in r[-1] for r in excluded))
                book.close()
                book = load_workbook(result.destination / "基准匹配结果/story.xlsx")
                self.assertEqual(book.active["AD6"].value, "正确")
                self.assertIn("完美匹配：上下各3句", book.active["AF6"].value)
                book.close()
                self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).digest() for p in before})

    def test_breath_brackets_only_affect_breath_rows_and_preserve_translation(self):
        path = self.write_source([["嗯？（惊讶(小声)）", "응? (놀람)", "气息"],
                                  ["正常(说明)", "일반", ""],
                                  ["哈(喘息)", "하", "气息"],
                                  ["（只有说明）", "(설명)", "气息"],
                                  ["哎（未闭合", "아", "气息"]])
        book = load_workbook(path)
        book.active["C1"] = "特殊说明"  # Missing speaker header must not drop normal rows.
        book.save(path)
        book.close()
        self.write_baseline(["嗯？", "嗯？（惊讶(小声)）", "嗯？(别的说明)", "<b>哈</b>",
                             "正常", "正常(说明)", "哎"])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.matched_rows, result.source_matched, result.match_pairs), (5, 3, 5))
        book = load_workbook(result.destination / "匹配明细.xlsx")
        rows = list(book["匹配明细"].values)[1:]
        breath_rows = [r for r in rows if r[8].startswith("嗯")]
        self.assertEqual(len(breath_rows), 3)
        self.assertEqual({r[9] for r in breath_rows}, {"응? (놀람)"})
        self.assertEqual({r[23] for r in breath_rows}, {"全文精确匹配", "气息去括号匹配"})
        self.assertTrue(all(r[24] == "否" for r in breath_rows))
        self.assertEqual(next(r[23] for r in rows if r[8] == "哈(喘息)"), "纯文本补匹配")
        book.close()
        self.assertEqual(breath_key("啊（动作）哈(小声)"), "啊哈")
        self.assertEqual(breath_key(" （只有说明） "), "")
        self.assertEqual(breath_key("嗯（不配对)"), "嗯（不配对)")
        self.assertEqual(breath_key("哎（未闭合"), "哎（未闭合")

    def test_special_context_stops_at_sheet_boundary_and_does_not_skip_normal_mismatch(self):
        path = self.write_source([["目标", "译文"], ["错误邻句", "译文"],
                                  [None, None, "演出"], ["（动作）", None, "气息"],
                                  ["后2", "译文"], ["后3", "译文"], ["后1", "译文"]])
        book = load_workbook(path)
        book.active["C1"] = "特殊说明"
        other = book.create_sheet("不能跨Sheet补句")
        other.append(["台词", "한국어"])
        other.append(["前1", "译文"])
        book.save(path)
        book.close()
        self.write_baseline(["前1", "前2", "前3", "目标", "后1", "后2", "后3"])
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        book = load_workbook(result.destination / "匹配明细.xlsx")
        row = next(r for r in list(book["匹配明细"].values)[1:] if r[8] == "目标")
        self.assertEqual(row[16:20], (0, 2, 0, 2))
        self.assertEqual(row[24], "否")
        self.assertEqual(row[-1], "向上：查询0行，演出0行，气息0行，匹配0句\n"
                         "向下：查询5行，演出1行，气息1行，匹配2句\n"
                         "合计（不含当前句）：查询5行，演出1行，气息1行，匹配2句")
        self.assertIsNone(row[14])
        book.close()

    def test_context_metadata_validation_rolls_back_all_evidence(self):
        self.write_source([["安全", "번역"]], name="安全.xlsx")
        self.write_baseline(["安全", "失败"])
        for metadata in (("C3", "=1+1", "特殊说明含公式"),
                         ("D3", "=1+1", "说话角色名含公式"),
                         ("E1", "特殊说明", "多个特殊说明列"),
                         ("E1", "说话角色名", "多个说话角色名列")):
            with self.subTest(metadata=metadata):
                path = self.write_source([["先读的气息(动作)", "译文", "气息", "角色"],
                                          ["失败", "错误", "", "角色"]])
                book = load_workbook(path)
                book.active["C1"], book.active["D1"] = "特殊说明", "说话角色名"
                book.active[metadata[0]] = metadata[1]
                book.save(path)
                book.close()
                result = match_folder(scan_folder(self.source, 1, include_deprecated=True), self.baseline, self.output)
                self.assertEqual(result.source_rows, 1)
                self.assertTrue(any(metadata[2] in r.issue for r in result.check_report.results if r.issue))
                book = load_workbook(result.destination / "匹配明细.xlsx")
                self.assertEqual(book["上下文排除明细"].max_row, 1)
                book.close()

    def test_many_baselines_summary_empty_missing_and_source_deduplication(self):
        self.write_source([["唯一", " 번역 "], ["唯一", "번역"], ["缺译", None]])
        for i in range(100):
            self.write_baseline(["唯一", "缺译", "无匹配"], name=f"{i:03}.xlsx")
        empty = self.write_baseline([], name="空.xlsx")
        result = match_folder(scan_folder(self.source, 1), self.baseline, self.output)
        self.assertEqual((result.baseline_files, result.baseline_rows, result.matched_rows, result.recommended_rows,
                          result.source_rows, result.source_matched, result.match_pairs), (101, 300, 200, 100, 3, 3, 300))
        book = load_workbook(result.destination / "基准文件匹配统计.xlsx")
        rows = {r[0]: r for r in list(book["文件统计"].values)[1:]}
        self.assertEqual(rows["000.xlsx"][2:11], (3, 2, 1, 1, 1, 1, 0, 0, 0))
        self.assertEqual(rows["空.xlsx"][2:11], (0,) * 9)
        self.assertEqual(rows["空.xlsx"][11:], (None, None, None))
        total = list(book["总体统计"].values)[1]
        self.assertEqual(total[:11], (101, 101, 300, 200, 100, 100, 100, 100, 0, 0, 0))
        self.assertAlmostEqual(total[11], 2/3)
        book.close()
        book = load_workbook(result.destination / "基准匹配结果/000.xlsx")
        self.assertEqual(book.active["AG2"].value, "唯一译文")
        self.assertEqual(book.active["AG3"].value, "全部缺译")
        self.assertEqual(book.active["AG4"].value, "未匹配")
        book.close()
        # New recommendation columns must also be protected from overwrite.
        book = load_workbook(empty)
        book.active["AG2"] = "已有判断"
        book.save(empty)
        book.close()
        with self.assertRaisesRegex(ValueError, "R～AG 列已有内容"):
            match_folder(scan_folder(self.source, 1), self.baseline, self.output)



if __name__ == "__main__":
    unittest.main()
