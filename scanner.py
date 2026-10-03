"""Read-only Excel header checks. Source workbooks are never saved."""

import csv
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree.ElementTree import iterparse

from openpyxl import load_workbook

from ignored_sheets import load_ignored, sheet_key

CHINESE_HEADERS = ("对白文本", "台词", "文本内容（示例）", "语音", "最新(2.18)", "中文原文")
KOREAN_HEADERS = ("韩文", "한국어", "韩语译文")
SUPPORTED = {".xlsx", ".xlsm"}
UNSUPPORTED = {".xls", ".xlsb", ".ods"}


def normalize(value):
    # NFKC handles full-width parentheses; NFC-compatible normalization also
    # handles decomposed Korean text. Internal whitespace is not discarded.
    return unicodedata.normalize("NFKC", value).strip() if isinstance(value, str) else ""


@dataclass
class CheckResult:
    path: Path
    sheet: str = ""
    visibility: str = ""
    issue: str = ""
    chinese: str = ""
    korean: str = ""
    header_row: int | None = None
    details: str = ""
    chinese_column: int | None = None
    korean_column: int | None = None

    @property
    def passed(self):
        return not self.issue


@dataclass
class ScanReport:
    root: Path
    header_rows: int
    files: int = 0
    results: list[CheckResult] = field(default_factory=list)
    ignored_sheets: int = 0
    chinese_headers: tuple = CHINESE_HEADERS
    korean_headers: tuple = KOREAN_HEADERS
    recursive: bool = True
    include_hidden_sheets: bool = True
    include_hidden_rows: bool = True
    skipped_sheets: list[CheckResult] = field(default_factory=list)
    file_paths: list[Path] = field(default_factory=list)

    @property
    def exceptions(self):
        return [result for result in self.results if not result.passed]

    @property
    def abnormal_files(self):
        return len({result.path for result in self.exceptions})

    @property
    def passed_sheets(self):
        return sum(result.passed for result in self.results)


def hidden_row_numbers(workbook, sheet, max_row=None):
    """Read row visibility metadata omitted by openpyxl's read-only sheets."""
    hidden = set()
    # The ZIP/path handles belong to openpyxl's read-only workbook; do not save it.
    with workbook._archive.open(sheet._worksheet_path) as stream:
        sheet_data = None
        for event, element in iterparse(stream, events=("start", "end")):
            if event == "start" and element.tag.endswith("}sheetData"):
                sheet_data = element
            if event == "end" and element.tag == "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row":
                number = int(element.get("r", "0"))
                if max_row is not None and number > max_row:
                    break
                if element.get("hidden", "0").lower() in ("1", "true"):
                    hidden.add(number)
                if sheet_data is not None:
                    sheet_data.remove(element)
                element.clear()
    return hidden


def check_sheet(path, sheet, header_rows, chinese_headers, korean_headers, hidden_rows=()):
    chinese, korean = [], []
    top_rows, candidate_rows = {}, {}
    # Ignore unreliable Excel dimension metadata; the explicit row limit still
    # bounds iteration. Visibility exclusions are determined separately.
    sheet.reset_dimensions()
    for number, row in enumerate(sheet.iter_rows(max_row=header_rows), 1):
        if number in hidden_rows:
            continue
        preview = []
        is_candidate = False
        row_number = None
        for cell in row:
            text = normalize(cell.value)
            if cell.value is not None and str(cell.value).strip():
                row_number = cell.row
                if len(preview) < 32:
                    value = str(cell.value).replace("\n", " ").replace("\r", " ")
                    preview.append(f"{cell.coordinate}={value[:80]}{'…' if len(value) > 80 else ''}")
            if text in chinese_headers:
                chinese.append((cell.row, cell.coordinate, str(cell.value)))
                is_candidate = True
            if text in korean_headers:
                korean.append((cell.row, cell.coordinate, str(cell.value)))
                is_candidate = True
        if preview and len(top_rows) < 3:
            top_rows[row_number] = "；".join(preview)
        if is_candidate and len(candidate_rows) < 3:
            candidate_rows[row_number] = "；".join(preview)

    describe = lambda hits: "；".join(f"{coordinate}={text}" for _, coordinate, text in hits)
    row_numbers = {hit[0] for hit in chinese} & {hit[0] for hit in korean}
    problems = []
    selected_row = None
    if not chinese and not korean:
        if not top_rows:
            problems.append("检查范围内没有非空单元格，未找到中文或韩语表头；需确认是否为翻译表")
        else:
            problems.append("未找到符合名称规则的中文表头，也未找到韩语表头；需确认是否为翻译表")
    elif not chinese:
        problems.append(f"缺少中文表头：已找到韩语表头 {describe(korean)}，未找到符合名称规则的中文表头")
    elif not korean:
        problems.append(f"缺少韩语表头：已找到中文表头 {describe(chinese)}，未找到符合名称规则的韩语表头")
    elif not row_numbers:
        problems.append(f"中文和韩语表头不在同一行：中文 {describe(chinese)}；韩语 {describe(korean)}")
    elif len(row_numbers) > 1:
        problems.append(f"存在多个中韩表头行（第 {'、'.join(map(str, sorted(row_numbers)))} 行），需确认使用哪一行")
    else:
        selected_row = next(iter(row_numbers))
        if sum(hit[0] == selected_row for hit in chinese) > 1:
            problems.append(f"第 {selected_row} 行存在多个中文候选列：{describe([hit for hit in chinese if hit[0] == selected_row])}")
        if sum(hit[0] == selected_row for hit in korean) > 1:
            problems.append(f"第 {selected_row} 行存在多个韩语候选列：{describe([hit for hit in korean if hit[0] == selected_row])}")
        if {hit[1] for hit in chinese} & {hit[1] for hit in korean}:
            problems.append("中文列与韩语列重合")

    sample_rows = {**top_rows, **candidate_rows}
    details = "\n".join([
        f"检查范围：实际行号 1～{header_rows}。" + (f"排除隐藏行：{'、'.join(map(str, sorted(hidden_rows))) or '无'}。" if hidden_rows else "本范围没有排除行。"),
        "匹配方式：完整名称匹配，忽略首尾空白并规范化 Unicode；不做包含或模糊匹配。",
        f"允许的中文表头：{'、'.join(sorted(chinese_headers))}",
        f"允许的韩语表头：{'、'.join(sorted(korean_headers))}",
        f"已找到中文表头：{describe(chinese) or '无'}",
        f"已找到韩语表头：{describe(korean) or '无'}",
        "顶部及候选表头行文本样例（每行最多 32 项，每项最多 80 字，供人工判断）：",
        *(f"第 {number} 行：{sample_rows[number]}" for number in sorted(sample_rows)),
        *([] if sample_rows else ["检查范围内没有非空单元格。"]),
    ])
    return CheckResult(
        path, sheet.title,
        {"visible": "可见", "hidden": "隐藏", "veryHidden": "深度隐藏"}.get(sheet.sheet_state, sheet.sheet_state),
        "；".join(problems), describe(chinese), describe(korean), selected_row, details,
        next((cell_column(coordinate) for number, coordinate, _ in chinese if number == selected_row), None),
        next((cell_column(coordinate) for number, coordinate, _ in korean if number == selected_row), None),
    )


def cell_column(coordinate):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    return column_index_from_string(coordinate_from_string(coordinate)[0])


def scan_folder(root, header_rows=10, chinese_headers=CHINESE_HEADERS,
                korean_headers=KOREAN_HEADERS, progress=None, ignored_sheets=(),
                recursive=True, include_hidden_sheets=True, include_hidden_rows=True):
    root = Path(root).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("请选择存在的源文件夹")
    if not isinstance(header_rows, int) or header_rows < 1:
        raise ValueError("检查行数必须是大于等于 1 的整数")
    chinese_headers = {normalize(value) for value in chinese_headers if normalize(value)}
    korean_headers = {normalize(value) for value in korean_headers if normalize(value)}
    if not chinese_headers or not korean_headers:
        raise ValueError("中文和韩语表头列表均不能为空")
    if chinese_headers & korean_headers:
        raise ValueError("中文和韩语表头列表不能使用同一个名称")
    files = sorted(
        (path for path in (root.rglob("*") if recursive else root.glob("*")) if path.is_file()
         and not path.name.startswith("~$")
         and path.suffix.lower() in SUPPORTED | UNSUPPORTED),
        key=lambda path: str(path).casefold(),
    )
    if not files:
        raise ValueError("所选扫描范围内没有 Excel 文件，请检查文件夹及子文件夹扫描开关")
    report = ScanReport(root, header_rows, len(files))
    report.chinese_headers = tuple(sorted(chinese_headers))
    report.korean_headers = tuple(sorted(korean_headers))
    report.recursive = recursive
    report.include_hidden_sheets = include_hidden_sheets
    report.include_hidden_rows = include_hidden_rows
    report.file_paths = files
    ignored_sheets = set(ignored_sheets)
    for index, path in enumerate(files, 1):
        if progress:
            progress(index, len(files), path)
        if path.suffix.lower() not in SUPPORTED:
            report.results.append(CheckResult(path, issue="暂不支持此格式，请转存为 .xlsx 后重新检查"))
            continue
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            if not workbook.worksheets:
                report.results.append(CheckResult(path, issue="没有可检查的工作表"))
            for sheet in workbook.worksheets:
                if sheet_key(path, sheet.title) in ignored_sheets:
                    report.ignored_sheets += 1
                    continue
                if not include_hidden_sheets and sheet.sheet_state != "visible":
                    report.skipped_sheets.append(CheckResult(path, sheet.title, issue="扫描隐藏工作表开关已关闭"))
                    continue
                try:
                    hidden_rows = hidden_row_numbers(workbook, sheet, header_rows) if not include_hidden_rows else set()
                    report.results.append(check_sheet(path, sheet, header_rows, chinese_headers, korean_headers, hidden_rows))
                except Exception as error:
                    report.results.append(CheckResult(path, sheet.title, issue=f"工作表读取失败：{type(error).__name__}: {error}"))
        except Exception as error:
            report.results.append(CheckResult(path, issue=f"文件读取失败：{type(error).__name__}: {error}"))
        finally:
            if workbook is not None:
                workbook.close()
    return report


def export_exceptions(report, destination):
    """UTF-8 BOM CSV opens in Windows Excel without garbled Chinese/Korean."""
    destination = Path(destination)
    if destination.suffix.lower() != ".csv":
        raise ValueError("异常清单必须保存为 .csv 文件")
    columns = ("文件名", "相对路径", "完整路径", "工作表", "显示状态", "异常原因",
               "中文表头位置", "韩语表头位置", "检查前几行", "检查依据", "该文件已通过的工作表")
    passed_sheets = {}
    for result in report.results:
        if result.passed:
            passed_sheets.setdefault(result.path, []).append(result.sheet)

    def safe_text(value):
        text = str(value)
        # Source-controlled filenames/sheet names must remain text in Excel.
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text

    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(columns)
        for result in report.exceptions:
            writer.writerow([
                safe_text(result.path.name), safe_text(result.path.relative_to(report.root)),
                safe_text(result.path), safe_text(result.sheet), result.visibility, result.issue,
                safe_text(result.chinese), safe_text(result.korean), report.header_rows,
                safe_text(result.details), safe_text("；".join(passed_sheets.get(result.path, [])) or "无"),
            ])


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="递归检查 Excel 中韩表头，导出异常清单")
    parser.add_argument("folder")
    parser.add_argument("--rows", type=int, default=10)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--no-recursive", action="store_true")
    parser.add_argument("--no-hidden-sheets", action="store_true")
    parser.add_argument("--no-hidden-rows", action="store_true")
    args = parser.parse_args()
    try:
        report = scan_folder(args.folder, args.rows, ignored_sheets=load_ignored(),
                             recursive=not args.no_recursive, include_hidden_sheets=not args.no_hidden_sheets,
                             include_hidden_rows=not args.no_hidden_rows)
        export_exceptions(report, args.output)
    except (ValueError, OSError) as error:
        parser.exit(1, f"检查失败：{error}\n")
    print(f"检查 {report.files} 个文件，{report.passed_sheets} 张表通过，"
          f"{report.abnormal_files} 个文件需确认，{len(report.exceptions)} 项异常，"
          f"跳过 {report.ignored_sheets} 张已忽略表，按开关跳过 {len(report.skipped_sheets)} 张隐藏表。")
    print(f"异常清单：{args.output.resolve()}")
