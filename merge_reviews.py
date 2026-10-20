"""Persist explicit merge reviews; never infer approval from workbook colors."""

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


DECISIONS = {"可适用", "不适用", "待确认"}


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def candidate_identity(sources, baseline, neighbors, evidence):
    def position(row):
        return [row["kind"], row["path"], row["sheet"], row["row"]]

    def content(row):
        return [position(row), row["original"], row["korean"], row["scene"], row["metadata"]]

    identity = digest([position(row) for row in [*sources, baseline]])
    # Version the evidence contract. Record IDs change between scans and are excluded.
    fingerprint = digest([1, [content(row) for row in [*sources, baseline, *neighbors]], evidence])
    return identity, fingerprint


def load_reviews(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}
    except (ValueError, UnicodeError) as error:
        raise ValueError(f"复核记录无法读取，请检查：{path}") from error
    if not isinstance(data, dict) or data.get("version") != 1 or not isinstance(data.get("reviews"), dict):
        raise ValueError(f"复核记录格式无效：{path}")
    for key, value in data["reviews"].items():
        if (not re.fullmatch(r"[0-9a-f]{64}", key) or not isinstance(value, dict)
                or value.get("decision") not in DECISIONS
                or not isinstance(value.get("comment"), str)
                or not isinstance(value.get("fingerprint"), str)
                or not re.fullmatch(r"[0-9a-f]{64}", value["fingerprint"])):
            raise ValueError(f"复核记录内容无效：{path}")
    return data["reviews"]


def apply_review_file(path, reviews, candidates):
    """Validate the entire import before changing the in-memory ledger."""
    if not path:
        return reviews
    current = {c["identity"]: c["fingerprint"] for c in candidates.values()}
    changes = {}
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if "复核填写" not in book.sheetnames:
            raise ValueError("请选择试合并结果中含“复核填写”工作表的文件；只读取明确填写的结论，不读取颜色")
        sheet = book["复核填写"]
        sheet.reset_dimensions()
        rows = iter(sheet.rows)
        headers = [c.value for c in next(rows, [])]
        required = ("对应标识", "内容指纹", "复核结论", "复核意见")
        if any(headers.count(h) != 1 for h in required):
            raise ValueError("复核填写表缺少唯一的对应标识、内容指纹、复核结论或复核意见列")
        columns = [headers.index(h) for h in required]
        for number, row in enumerate(rows, 2):
            if max(columns) >= len(row):
                continue
            cells = [row[i] for i in columns]
            identity, fingerprint, decision, comment = [c.value for c in cells]
            if decision is None or decision == "":
                continue
            if any(c.data_type == "f" for c in cells):
                raise ValueError(f"复核第 {number} 行含公式，请填写普通文本")
            if not isinstance(decision, str) or decision.strip() not in DECISIONS:
                raise ValueError(f"复核第 {number} 行结论只能填写：可适用、不适用、待确认")
            if not isinstance(identity, str) or identity not in current:
                raise ValueError(f"复核第 {number} 行对应关系已不存在或不在当前合并范围，请使用最新扫描生成的复核表")
            if fingerprint != current[identity]:
                raise ValueError(f"复核第 {number} 行内容或定位证据已变化，请重新扫描后复核")
            if identity in changes:
                raise ValueError(f"复核第 {number} 行对应关系重复，请每组只保留一条结论")
            changes[identity] = dict(fingerprint=fingerprint, decision=decision.strip(),
                                     comment="" if comment is None else str(comment))
    finally:
        book.close()
    return {**reviews, **changes}
