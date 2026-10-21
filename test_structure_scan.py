"""Evidence-based checks for structural scans and read-only report generation."""

import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from ignored_sheets import sheet_key
from matcher import match_folder
from scanner import scan_folder
from structure_scan import bilingual_evidence, comparison_key, scan_structure, space_merged_korean, split_patterns, split_translation, trial_merge, trial_numbered_split


class StructureScanTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "源"
        self.baseline = self.root / "基准"
        self.output = self.root / "结果"
        self.source.mkdir()
        self.baseline.mkdir()

    def workbook(self, root, rows, name=None, title="正文", hidden=()):
        name = name or ("char.xlsx" if root == self.baseline else "样本.xlsx")
        book = Workbook()
        sheet = book.active
        sheet.title = title
        for row in rows:
            sheet.append(row)
        for number in hidden:
            sheet.row_dimensions[number].hidden = True
        path = root / name
        book.save(path)
        book.close()
        return path

    def scan(self, max_rows=8, ignored=(), review_report=None, **options):
        report = scan_folder(self.source, ignored_sheets=ignored, **options)
        result = scan_structure(report, self.baseline, self.output, ignored_sheets=ignored,
                                max_rows=max_rows, review_report=review_report)
        book = load_workbook(result.destination / "结构差异扫描.xlsx", read_only=True)
        self.addCleanup(book.close)
        return result, {sheet.title: list(sheet.values) for sheet in book}

    def test_merges_splits_evidence_and_originals(self):
        source = self.workbook(self.source, [
            ["台词", "한국어"],
            ["前锚", "앞"], ["甲", "가"], ["乙", "나"], ["丙", "다"], ["后锚", "뒤"],
            ["A.请扶我一下，谢谢 B.什么…爆炸魔法？", "A.도와주세요 B.폭발 마법?"],
            ["无分隔甲无分隔乙", "분리할 수 없는 번역"],
            ["未知内容", "알 수 없음"],
            ["=literal", "=not_a_formula"],
        ])
        # Excel uses leading '=' as a formula unless explicitly written as text.
        book = load_workbook(source)
        book.active["A10"].data_type = book.active["B10"].data_type = "s"
        book.save(source)
        book.close()
        baseline = self.workbook(self.baseline, [["Text"], ["前锚"], ["甲乙丙"], ["后锚"],
                                                   ["请扶我一下，谢谢"], ["什么…爆炸魔法？"],
                                                   ["无分隔甲"], ["无分隔乙"]])
        hashes = [hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)]
        result, sheets = self.scan()
        self.assertEqual(result.merge_candidates, 1)
        self.assertEqual(result.merge_trial_rows, 1)
        merged = sheets["合并候选"][1]
        self.assertEqual(merged[4], "3、4、5")
        self.assertIn("前1句一致（双方全文唯一）", merged[11])
        self.assertIn("后1句一致（双方全文唯一）", merged[11])
        split = sheets["拆分候选"][1:]
        self.assertEqual(len(split), 2)
        self.assertTrue(any("编号、顺序与段数一致" in row[12] for row in split))
        self.assertTrue(any("韩文无可识别分隔符" in row[12] for row in split))
        unknown = {row[5] for row in sheets["未解释"][1:]}
        self.assertIn("未知内容", unknown)
        self.assertIn("=literal", unknown)
        self.assertEqual(len(sheets["规则样本"][0]), 15)
        self.assertTrue(all(len(row) == 15 for row in sheets["规则样本"]))
        self.assertEqual(sum(row[2] for row in sheets["规则汇总"][1:]), result.pattern_samples)
        book = load_workbook(result.destination / "结构差异扫描.xlsx")
        self.addCleanup(book.close)
        raw = list(book["原始内容"].rows)
        cell = next(row[6] for row in raw[1:] if row[6].value == "=literal")
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(hashes, [hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)])

    def test_patterns_preserve_raw_boundaries_and_do_not_claim_translation_alignment(self):
        text = "序言 Ａ：甲 Ｂ：乙"
        numbered = split_patterns(text)[0]
        self.assertEqual(numbered["labels"], ["A", "B"])
        self.assertEqual(numbered["parts"], [text[a:b] for a, b in numbered["spans"]])
        self.assertIn("首个编号前有未归属文字", numbered["issues"])
        repeated = split_patterns("A.甲 B.乙 A.丙 B.丁")[0]
        self.assertIn("编号重复，可能包含多组选项", repeated["issues"])
        mixed = split_patterns("1:甲 B:乙")[0]
        self.assertEqual(mixed["rule"], "混合编号")
        self.assertFalse(split_patterns("价格1.5，价格2.5"))
        pattern = split_patterns("A.甲 B.乙")[0]
        self.assertIn("歧义", bilingual_evidence(pattern, "B.나 A.가")[0])
        self.assertIn("仅段数相同", bilingual_evidence(pattern, "가\n나")[0])
        self.assertIn("数量不一致", bilingual_evidence(pattern, "가\n나\n다")[0])
        self.assertEqual(split_patterns("甲\\n乙\n丙")[0]["rule"], "混合换行")
        self.assertEqual(comparison_key('<color=#fff>甲</color><ruby="a">乙</ruby>\\n{playername}%s%d'),
                         "甲乙{playername}%s%d")
        self.assertNotEqual(comparison_key("甲{playername}"), comparison_key("甲"))

    def test_trial_split_preserves_text_and_rejects_unsafe_boundaries(self):
        cn = " １、<color=#fff>甲{playername}%s</color> \n２、乙 %d"
        ko = " 1. <color=#fff>가{playername}%s</color>\n2. 나 %d "
        patterns, reason = trial_numbered_split(cn, ko, ["甲{playername}%s", "乙 %d"])
        self.assertFalse(reason)
        for text, pattern in zip((cn, ko), patterns):
            self.assertEqual(pattern["parts"], [text[a:b] for a, b in pattern["spans"]])
        self.assertEqual(patterns[1]["parts"][1].strip(), "나 %d")
        for cn, ko, baseline, expected in [
            ("1、甲\n2、乙", "가\n나", ["甲", "乙"], "数字编号"),
            ("1、甲\n2、乙", "2. 나\n1. 가", ["甲", "乙"], "连续排列"),
            ("1、甲\n2、乙", "1. 가\n2. 나\n3. 다", ["甲", "乙"], "段数不一致"),
            ("1、甲乙\n2、丙", "1. 가\n2. 나", ["甲", "乙丙"], "逐条一致"),
            ("1、甲{playername}\n2、乙", "1. 가\n2. 나{playername}", ["甲{playername}", "乙"], "占位符"),
            ("1、<color=red>甲\n2、乙</color>", "1. 가\n2. 나", ["甲", "乙"], "标签"),
            ("前言1、甲\n2、乙", "1. 가\n2. 나", ["甲", "乙"], "未归属"),
            ("1、甲\n2、乙\n1、丙\n2、丁", "1. 가\n2. 나", ["甲", "乙", "丙", "丁"], "重复"),
            ("1、甲\n内部句\n2、乙", "1. 가\n2. 나", ["甲内部句", "乙"], "内部含换行"),
            ("A.甲 B.乙", "A.가 B.나", ["甲", "乙"], "数字编号"),
        ]:
            with self.subTest(cn=cn, ko=ko):
                patterns, reason = trial_numbered_split(cn, ko, baseline)
                self.assertIsNone(patterns)
                self.assertIn(expected, reason)

    def test_whole_matching_precedes_reverse_lookup_including_missing_and_conflicts(self):
        self.workbook(self.source, [
            ["台词", "한국어"],
            ["已有", "기존"], ["整句", "문장"], ["已有整句", "정규"], [None],
            ["标签", "표"], ["匹配", "문"], ["<color=red>标签匹配</color>", "태그"], [None],
            ["缺译", "없"], ["匹配", "음"], ["缺译匹配", None], [None],
            ["冲突", "충"], ["匹配", "돌"], ["冲突匹配", "가"], ["冲突匹配", "나"], [None],
            ["新增一", " =새1"], ["新增二", "새2"],
        ])
        self.workbook(self.baseline, [["Text"], ["已有整句"], ["标签匹配"], ["缺译匹配"],
                                        ["冲突匹配"], ["新增一新增二"]])
        result, sheets = self.scan()
        self.assertEqual(result.whole_matched_rows, 4)
        self.assertEqual((result.merge_candidates, result.merge_trial_rows, result.merge_pending_rows), (1, 0, 1))
        self.assertEqual(sheets["合并候选"][1][9], "6")
        book = load_workbook(result.destination / "试合并结果.xlsx")
        self.addCleanup(book.close)
        self.assertEqual(book["全部反查候选"]["D2"].value, "=새1 새2")
        self.assertEqual(book["全部反查候选"]["D2"].data_type, "s")
        self.assertEqual(book["全部反查候选"]["M2"].value, 6)
        whole = {row[3]: row for row in list(book["整句命中检查"].values)[1:]}
        self.assertEqual(whole["标签匹配"][4], "纯文本补匹配")
        self.assertEqual(whole["缺译匹配"][6], 0)
        self.assertIn("全部缺译", whole["缺译匹配"][7])
        self.assertEqual(whole["冲突匹配"][6], 2)
        self.assertNotIn("标签匹配", {row[5] for row in sheets["未解释"][1:] if row[1] == "基准"})

    def test_whole_matching_shares_source_filters_breath_and_source_coverage(self):
        self.workbook(self.source, [
            ["台词", "한국어", "特殊说明", "说话角色名"],
            ["甲乙", "无效说明的译文", None, None],
            ["甲", "가", None, "角色"], ["乙", "나", None, "角色"],
            ["哈（喘息）", "하", "气息", None],
            ["<COLOR=#3266b4>你好</COLOR>{playername}", "<b>안녕</b>", None, "角色"],
            ["哈（喘息）", None, "气息", None],  # One physical row counts once despite two matching keys.
        ])
        self.workbook(self.baseline, [["Text"], ["甲乙"], ["哈"], ["你好"], ["哈（喘息）"]])
        report = scan_folder(self.source)
        matched = match_folder(report, self.baseline, self.output)
        result, sheets = self.scan()
        self.assertEqual(result.source_rows, matched.source_rows)
        self.assertEqual((result.whole_matched_rows, matched.matched_rows), (3, 3))
        self.assertEqual(result.merge_candidates, 1)
        self.assertEqual(sheets["合并候选"][1][4], "3、4")
        self.assertFalse(sheets["未解释"][1:], "Whole matches must cover both baseline and source rows")
        self.assertTrue(any("无效说明" in row[3] for row in sheets["跳过清单"][1:]))
        book = load_workbook(result.destination / "试合并结果.xlsx", read_only=True)
        try:
            whole = {r[3]: r for r in list(book["整句命中检查"].values)[1:]}
            self.assertEqual(whole["哈"][4:7], ("气息去括号匹配", 2, 1))
            self.assertEqual(whole["哈（喘息）"][4:7], ("全文精确匹配", 2, 1))
            self.assertEqual(whole["你好"][4:7], ("纯文本补匹配", 1, 1))
        finally:
            book.close()

    def test_tags_are_ignored_for_structure_comparison_and_retained_in_preview(self):
        source = self.workbook(self.source, [
            ["台词", "한국어"], ["<color=#3266b4>甲</color>", "<color=red>가</color>"],
            ["<b>乙</b>", "<b>나</b>"], [None],
            ["1、<size=20>丙</size>\n2、<i>丁</i>", "1. <size=18>다</size>\n2. <i>라</i>"],
        ])
        baseline = self.workbook(self.baseline, [["Text"], ["<i>甲乙</i>"], [None], ["丙"], ["丁"]])
        before = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)}
        result, sheets = self.scan()
        self.assertEqual(result.merge_candidates, 1)
        self.assertEqual((result.trial_groups, result.trial_segments), (1, 2))
        self.assertIn("<color=#3266b4>", sheets["合并候选"][1][5])
        book = load_workbook(result.destination / "试拆分结果.xlsx", read_only=True)
        try:
            rows = list(book["拆分明细"].values)[1:]
            self.assertEqual([r[3] for r in rows], ["<size=18>다</size>", "<i>라</i>"])
        finally:
            book.close()
        self.assertEqual(comparison_key("<b>甲</b>{playername}%s%d"), "甲{playername}%s%d")
        self.assertEqual(comparison_key("1 < 2；中文<说明>"), "1 < 2；中文<说明>")
        self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).digest() for p in before})

    def test_failed_review_save_removes_published_reports_and_preserves_saved_decisions(self):
        self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]])
        self.workbook(self.baseline, [["Text"], ["甲乙"]])
        initial, _ = self.scan()
        review = self.review_file(initial, [(2, "可适用", "确认位置")])
        accepted, _ = self.scan(review_report=review)
        self.assertEqual(accepted.merge_trial_rows, 1)
        ledger = self.output / "合并复核记录.json"
        before = ledger.read_bytes()
        outputs = set(self.output.iterdir())
        revoked = self.review_file(accepted, [(2, "不适用", "撤回")])
        with patch("structure_scan.write_json", side_effect=PermissionError("复核保存失败")):
            with self.assertRaisesRegex(PermissionError, "复核保存失败"):
                self.scan(review_report=revoked)
        self.assertEqual(ledger.read_bytes(), before)
        self.assertEqual(set(self.output.iterdir()), outputs)

    def test_merge_preview_keeps_all_equal_sources_and_defers_position_conflicts(self):
        self.workbook(self.source, [["台词", "한국어"], ["共", "공"], ["同", "동"], [None],
                                   ["甲", "가"], ["乙", "나"], [None], ["缺", None], ["译", "문"]])
        self.workbook(self.source, [["台词", "한국어"], ["共", "공"], ["同", "동"]], name="重复.xlsx")
        self.workbook(self.baseline, [["Text"], ["共同"], [None], ["甲乙"], [None], ["甲乙"], [None], ["缺译"]])
        result, _ = self.scan()
        self.assertEqual((result.merge_candidates, result.merge_trial_rows, result.merge_pending_rows), (5, 0, 4))
        book = load_workbook(result.destination / "试合并结果.xlsx", read_only=True)
        self.addCleanup(book.close)
        accepted = list(book["合并建议"].values)[1:]
        self.assertEqual(accepted, [])
        candidates = list(book["全部反查候选"].values)[1:]
        self.assertEqual(len(candidates), 5)
        self.assertTrue(any("归属不唯一" in row[4] for row in candidates))
        self.assertTrue(any("缺译" in row[4] for row in candidates))
        self.assertTrue(any("即使译文相同也不自动采用" in row[4] for row in candidates))

    def review_file(self, result, decisions):
        path = self.root / "填写后的复核.xlsx"
        book = load_workbook(result.destination / "试合并结果.xlsx")
        for row, decision, comment in decisions:
            book["复核填写"].cell(row, 2, decision)
            book["复核填写"].cell(row, 3, comment)
        book.save(path)
        book.close()
        return path

    def test_only_char_baselines_merge_while_story_still_matches_and_splits(self):
        self.workbook(self.source, [["台词", "한국어"], ["已有整句", "문장"],
                                   ["甲", "가"], ["乙", "나"], [None],
                                   ["1、选甲\n2、选乙", "1. 가\n2. 나"]], name="story.xlsx")
        self.workbook(self.baseline, [["Text"], ["已有整句"], ["甲乙"]], name="CHAR.XLSX")
        story = self.workbook(self.baseline, [["Text"], ["已有整句"], ["甲乙"], ["选甲"], ["选乙"]],
                              name="story.xlsx")
        for name in ("char_extra.xlsx", "character.xlsx", "char.xlsm"):
            self.workbook(self.baseline, [["Text"], ["甲乙"]], name=name)
        result, sheets = self.scan()
        self.assertEqual((result.baseline_rows, result.whole_matched_rows, result.merge_candidates), (9, 2, 1))
        self.assertTrue(all(Path(row[7]).name == "CHAR.XLSX" for row in sheets["合并候选"][1:]))
        self.assertTrue(all(row[7] == str(story.resolve()) for row in sheets["拆分候选"][1:]))
        self.assertEqual((result.trial_groups, result.trial_segments), (1, 2))
        self.assertTrue(any(r[2] == str(story.resolve()) and r[5] == "甲乙" and "未执行合并反查" in r[7]
                            for r in sheets["未解释"][1:]))
        book = load_workbook(result.destination / "试合并结果.xlsx", read_only=True)
        try:
            stats = dict(book["流程统计与说明"].values)
            self.assertEqual(stats["进入合并反查的基准行"], 1)
            self.assertEqual(stats["未命中整句但不在合并范围的基准行"], 6)
            self.assertEqual(stats["反查仍无候选的基准行"], 0)
            scope = {Path(r[0]).name: r for r in list(book["合并范围"].values)[1:]}
            self.assertEqual(scope["story.xlsx"][1:4], (4, 1, 0))
            self.assertEqual(scope["CHAR.XLSX"][1:4], (2, 1, 1))
        finally:
            book.close()
        # The identical story text must not block char approval as a competing merge.
        review = self.review_file(result, [(2, "可适用", "已确认角色技能台词")])
        result, _ = self.scan(review_report=review)
        self.assertEqual(result.merge_trial_rows, 1)

    def test_no_char_file_and_old_story_approval_cannot_enable_merging(self):
        self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]])
        self.workbook(self.baseline, [["Text"], ["甲乙"]], name="story.xlsx")
        # Reproduce a saved approval from the previous unrestricted version.
        with patch("structure_scan.is_merge_baseline", return_value=True):
            old, _ = self.scan()
            review = self.review_file(old, [(2, "可适用", "旧版确认")])
            approved, _ = self.scan(review_report=review)
            self.assertEqual(approved.merge_trial_rows, 1)
        ledger = (self.output / "合并复核记录.json").read_bytes()
        result, sheets = self.scan()
        self.assertEqual((result.merge_candidates, result.merge_trial_rows, result.merge_pending_rows), (0, 0, 0))
        self.assertEqual(len(sheets["合并候选"]), 1)
        self.assertIn("未执行合并反查", sheets["未解释"][1][7])
        book = load_workbook(result.destination / "试合并结果.xlsx", read_only=True)
        try:
            stats = dict(book["流程统计与说明"].values)
            self.assertEqual(stats["进入合并反查的基准行"], 0)
            self.assertEqual(stats["未命中整句但不在合并范围的基准行"], 1)
            self.assertEqual(stats["反查仍无候选的基准行"], 0)
            self.assertEqual(list(book["复核填写"].values)[1:], [])
        finally:
            book.close()
        with self.assertRaisesRegex(ValueError, "不在当前合并范围"):
            self.scan(review_report=review)
        self.assertEqual((self.output / "合并复核记录.json").read_bytes(), ledger)

    def test_roles_segments_and_optional_formulas_are_evidence(self):
        self.workbook(self.source, [
            ["台词", "한국어", "说话角色名", "序号", "语音类型"],
            ["甲", "가", "甲角色", "A", "VO"], ["乙", "나", "乙角色", "A", "VO"], [None],
            ["丙", "다", "主角", "B", "VO"], ["丁", "라", "主角", "C", "VO"], [None],
            ["戊", "마", "主角", '=1+1', "VO"], ["己", "바", "主角", '=1+1', "VO"], [None],
            ["庚", "사", "主角"], ["辛", "아", "主角"],
        ])
        self.workbook(self.baseline, [["Text", "Speaker"], ["甲乙", "甲角色"], ["丙丁", "主角"],
                                      ["戊己", "主角"], ["庚辛", "配角"]])
        result, sheets = self.scan()
        self.assertEqual(result.source_rows, 8)
        self.assertEqual((result.merge_candidates, result.merge_trial_rows, result.merge_pending_rows,
                          result.merge_rejected_rows), (4, 0, 1, 3))
        decisions = {r[9]: r for r in sheets["合并候选"][1:]}
        self.assertIn("不同角色", decisions["2"][12])
        self.assertIn("跨剧情段", decisions["3"][12])
        self.assertIn("未知：公式", decisions["4"][14])
        self.assertIn("源角色与基准角色不一致", decisions["5"][12])
        review = self.review_file(result, [(2, "可适用", "仍不能覆盖硬性角色冲突")])
        result, _ = self.scan(review_report=review)
        self.assertEqual(result.merge_trial_rows, 0)

    def test_review_persistence_changes_and_stale_evidence(self):
        source = self.workbook(self.source, [["台词", "한국어", "说话角色名"],
                                             ["甲", "가", "主角"], ["乙", "나", "主角"],
                                             ["后邻句", "이웃", "主角"]])
        self.workbook(self.baseline, [["Text", "Speaker"], ["甲乙", "主角"], ["另一个后邻句", "主角"]])
        initial, _ = self.scan()
        self.assertEqual(initial.merge_trial_rows, 0)
        review = self.review_file(initial, [(2, "可适用", "剧情位置人工确认")])
        accepted, _ = self.scan(review_report=review)
        self.assertEqual(accepted.merge_trial_rows, 1)
        again, _ = self.scan()
        self.assertEqual(again.merge_trial_rows, 1)
        for column, row, value in [(2, 2, "新的韩文"), (3, 2, "主角改名"), (1, 4, "邻句变更")]:
            original = source.read_bytes()
            book = load_workbook(source)
            book.active.cell(row, column, value)
            book.save(source)
            book.close()
            stale, sheets = self.scan()
            self.assertEqual(stale.merge_trial_rows, 0)
            # A role conflict remains a hard rejection even with an old approval.
            self.assertTrue(any("复核已过期" in r[12] or "不同角色" in r[12]
                                for r in sheets["合并候选"][1:]))
            before = (self.output / "合并复核记录.json").read_bytes()
            with self.assertRaisesRegex(ValueError, "证据已变化"):
                self.scan(review_report=review)
            self.assertEqual((self.output / "合并复核记录.json").read_bytes(), before)
            source.write_bytes(original)
        pending = self.review_file(again, [(2, "待确认", "理论上可用，还需核实")])
        result, _ = self.scan(review_report=pending)
        self.assertEqual((result.merge_trial_rows, result.merge_pending_rows), (0, 1))
        rejected = self.review_file(result, [(2, "不适用", "实际对应另一处")])
        result, _ = self.scan(review_report=rejected)
        self.assertEqual((result.merge_pending_rows, result.merge_rejected_rows), (0, 1))

    def test_review_does_not_use_color_or_accept_invalid_import(self):
        from openpyxl.styles import PatternFill
        self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]])
        self.workbook(self.baseline, [["Text"], ["甲乙"]])
        result, _ = self.scan()
        path = self.review_file(result, [])
        book = load_workbook(path)
        book["复核填写"]["B2"].fill = PatternFill("solid", fgColor="92D050")
        book.save(path)
        book.close()
        result, _ = self.scan(review_report=path)
        self.assertEqual(result.merge_trial_rows, 0)
        ledger = (self.output / "合并复核记录.json").read_bytes()
        for value in ("理论上可用", "=1+1"):
            path = self.review_file(result, [(2, value, "invalid")])
            outputs = set(self.output.iterdir())
            with self.assertRaises(ValueError):
                self.scan(review_report=path)
            self.assertEqual(set(self.output.iterdir()), outputs)
            self.assertEqual((self.output / "合并复核记录.json").read_bytes(), ledger)

    def test_review_selects_one_source_but_not_multiple_approvals(self):
        for name in ("一.xlsx", "二.xlsx"):
            self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]], name=name)
        self.workbook(self.baseline, [["Text"], ["甲乙"]])
        result, _ = self.scan()
        path = self.review_file(result, [(2, "可适用", "已定位"), (3, "待确认", "未定位")])
        result, _ = self.scan(review_report=path)
        self.assertEqual(result.merge_trial_rows, 1)
        path = self.review_file(result, [(2, "可适用", "已定位"), (3, "可适用", "重复确认")])
        result, sheets = self.scan(review_report=path)
        self.assertEqual((result.merge_trial_rows, result.merge_pending_rows), (0, 1))
        self.assertTrue(all("多组通过复核" in r[12] for r in sheets["合并候选"][1:]))

    def test_merge_text_validation(self):
        sources = [dict(row=2, original="甲{playername}", korean="=가{playername} "),
                   dict(row=3, original="乙%s", korean="<color=red>나%s</color>")]
        self.assertEqual(trial_merge(sources, dict(original="甲{playername}乙%s")),
                         ("=가{playername} <color=red>나%s</color>", ""))
        for sources, baseline, expected in [
            ([dict(row=2, original="……", korean="..."), dict(row=3, original="……", korean="...")],
             "…………", "缺少可确认"),
            ([dict(row=2, original="甲", korean="")], "甲", "缺译"),
            ([dict(row=2, original="甲%s", korean="가")], "甲%s", "占位符"),
            ([dict(row=2, original="甲", korean="<color=red>가")], "甲", "标签"),
            ([dict(row=2, original="甲", korean="가")], "乙", "未完整覆盖"),
        ]:
            with self.subTest(expected=expected):
                translation, reason = trial_merge(sources, dict(original=baseline))
                self.assertIsNone(translation)
                self.assertIn(expected, reason)

    def test_merge_full_row_prefix_suffix_reuse_keeps_ambiguous_overlap_pending(self):
        self.workbook(self.source, [['台词', '한국어'], ['前锚', '앞'],
                                   ['玉枝倚宝扇…', '옥가지를 부채에 기대어...'],
                                   ['花间千丝系尘缘——', '천 갈래 실이 인연을 묶나니—'],
                                   ['因果已结！', '인과여, 맺어져라!'], ['后锚', '뒤']])
        parts = ['玉枝倚宝扇…', '花间千丝系尘缘——', '因果已结！']
        for shortened, expected, side in [
            (parts[1:], '천 갈래 실이 인연을 묶나니— 인과여, 맺어져라!', '开头'),
            (parts[:2], '옥가지를 부채에 기대어... 천 갈래 실이 인연을 묶나니—', '结尾'),
        ]:
            with self.subTest(side=side):
                self.workbook(self.baseline, [['Text'], ['前锚'], [''.join(parts)], ['后锚'], [None], [''.join(shortened)]])
                result, sheets = self.scan()
                self.assertEqual((result.merge_candidates, result.merge_trial_rows), (2, 2))
                short = next(r for r in sheets['合并候选'][1:] if r[9] == '6')
                self.assertIn(f'省略{side}1条完整源台词', short[12])
                book = load_workbook(result.destination / '试合并结果.xlsx')
                try:
                    accepted = {r[7]: r for r in list(book['合并建议'].values)[1:]}
                    self.assertEqual(accepted[6][2], expected)
                finally:
                    book.close()
        # Identical windows at different baseline positions remain ambiguous.
        self.workbook(self.baseline, [['Text'], [''.join(parts)], [None], [''.join(parts)]])
        result, _ = self.scan()
        self.assertEqual((result.merge_trial_rows, result.merge_pending_rows), (0, 2))

    def test_merge_korean_sentence_spacing_preserves_protected_content(self):
        for original, expected in [
            ('가...나!다.라…마！바。끝...', '가... 나! 다. 라… 마！ 바。 끝...'),
            ('가...?!나!!다', '가...?! 나!! 다'),
            ('가. 나!  다…\n라', '가. 나!  다…\n라'),
            ('가.\\n나!\\r\\n다', '가.\\n나!\\r\\n다'),
            ('3.14 v1.2 A.B 가.나', '3.14 v1.2 A.B 가. 나'),
            ('.5 -.5 (.5) 가.5개', '.5 -.5 (.5) 가. 5개'),
            ('「가!」나', '「가!」 나'),
            ('<color=#1.2>가!</color><b>나</b>', '<color=#1.2>가!</color><b> 나</b>'),
            ('가!{player.name}나', '가! {player.name}나'),
            ('<b>가...</b>', '<b>가...</b>'),
        ]:
            with self.subTest(original=original):
                self.assertEqual(space_merged_korean(original), expected)
                self.assertEqual(space_merged_korean(expected), expected)
        sources = [dict(row=2, original='甲…', korean=' 가... '),
                   dict(row=3, original='乙！', korean='나!다.'),
                   dict(row=4, original='丙。', korean='라.')]
        self.assertEqual(trial_merge(sources, dict(original='甲…乙！丙。')), ('가... 나! 다. 라.', ''))

    def test_structure_recommendations_fill_copies_and_keep_whole_recommendations_and_partial_sources(self):
        source = self.workbook(self.source, [
            ['台词', '한국어'], ['前锚', '앞'], ['甲…', '가...'], ['乙！', '나!다.'], ['后锚', '뒤'],
            ['1、已命中\n2、新拆', '1. 다른 번역\n2. =새'], ['已命中', '整句译文'],
            ['留一\n缺一', '부분1\n부분2'], ['独甲\n独乙', '단독1\n단독2'], ['留一', None],
        ])
        baseline = self.workbook(self.baseline, [
            ['说明'], [None], ['Key', 'Text', 'TextTrans'],
            *[[str(i), text, '已有译文'] for i, text in enumerate(
                ['前锚', '\n甲…乙！\n', '后锚', '已命中', '新拆', '留一', '独甲', '独乙'])],
        ])
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)}
        result, _ = self.scan()
        matched = result.match_report
        self.assertEqual((matched.baseline_rows, matched.matched_rows, matched.recommended_rows), (8, 8, 8))
        self.assertTrue(all(p.is_file() for p in matched.outputs))
        book = load_workbook(matched.destination / '基准匹配结果/char.xlsx')
        try:
            sheet = book.active
            headers = {c.value: c.column for c in sheet[3] if c.value}
            for row, translation, status in [
                (5, '\n가... 나! 다.\n', '合并推荐·待复核'),
                (7, '整句译文', '唯一译文'), (8, '=새', '拆分推荐·待复核'),
                (9, '부분1', '拆分推荐·待复核'), (10, '단독1', '拆分推荐·待复核'),
                (11, '단독2', '拆分推荐·待复核'),
            ]:
                self.assertEqual(sheet.cell(row, headers['推荐韩语']).value, translation)
                self.assertEqual(sheet.cell(row, headers['推荐状态']).value, status)
                self.assertEqual(sheet.cell(row, headers['TextTrans']).value, '已有译文')
            self.assertEqual(sheet.cell(8, headers['推荐韩语']).data_type, 's')
            self.assertEqual(sheet.cell(9, headers['匹配条数']).value, 2)
            self.assertEqual(sheet.cell(9, headers['缺译条数']).value, 1)
            self.assertIn('句中省略号', sheet.cell(5, headers['推荐依据']).value)
            self.assertIn('韩文字符', sheet.cell(8, headers['推荐依据']).value)
        finally:
            book.close()
        book = load_workbook(matched.destination / '匹配明细.xlsx')
        try:
            structural = list(book['合并拆分推荐'].values)[1:]
            self.assertEqual(len(structural), 5)
            self.assertEqual({r[7] for r in structural}, {'合并推荐·待复核', '拆分推荐·待复核'})
            pending = list(book['结构推荐未写入'].values)[1:]
            self.assertEqual([(r[3], r[-1]) for r in pending], [('已命中', '保留已有整句推荐')])
            unused = list(book['未推荐明细'].values)[1:]
            self.assertEqual({r[4] for r in unused}, {'1、已命中\n2、新拆', '留一\n缺一', '留一'})
            self.assertTrue(all(r[-1] == '拆分后仅部分分段获推荐，剩余分段待复核' for r in unused if '\n' in r[4]))
            self.assertEqual(next(r[7] for r in unused if r[4] == '留一'), '是')
            self.assertEqual(list(book['未匹配明细'].values)[1:], [])
        finally:
            book.close()
        book = load_workbook(matched.destination / '基准文件匹配统计.xlsx')
        try:
            self.assertEqual(list(book['总体统计'].values)[1][2:6], (8, 8, 8, 8))
        finally:
            book.close()
        self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).digest() for p in hashes})

    def test_fuzzy_pass_follows_structure_and_excludes_used_or_partially_used_sources(self):
        self.workbook(self.source, [
            ['台词', '한국어'], ['前锚', '앞'], ['春风吹过山岗', '봄바람'], ['花香飘进小院', '꽃향기'], ['后锚', '뒤'],
            ['春风吹过山岗花香飘进小院啊', '低优先级模糊译文'],
            ['单独需要拆分\n剩余未命中', '부분1\n부분2'], [None, None], ['春风吹过山岗', '봄바람'],
        ])
        self.workbook(self.baseline, [['Text'], ['前锚'], ['春风吹过山岗花香飘进小院'], ['后锚'],
                                     ['春风吹过山冈'], ['单独需要拆分'], ['单独需要拆分剩余未命中啊']])
        result, _ = self.scan()
        book = load_workbook(result.match_report.destination / '基准匹配结果/char.xlsx')
        try:
            self.assertEqual((book.active['AD3'].value, book.active['AG3'].value), ('봄바람 꽃향기', '合并推荐·待复核'))
            self.assertEqual(book.active['AG5'].value, '未匹配')
            self.assertEqual(book.active['AG6'].value, '拆分推荐·待复核')
            self.assertEqual(book.active['AG7'].value, '未匹配')
        finally:
            book.close()
        book = load_workbook(result.match_report.destination / '匹配明细.xlsx')
        try:
            self.assertFalse(any(r[23] == '模糊匹配' for r in list(book['匹配明细'].values)[1:]))
            remaining = list(book['未推荐明细'].values)[1:]
            self.assertTrue(any(r[4] == '单独需要拆分\n剩余未命中' and '部分分段' in r[-1] for r in remaining))
            self.assertFalse(any(r[4] == '春风吹过山岗' for r in remaining))
        finally:
            book.close()

    def test_merge_split_competition_does_not_pick_an_arbitrary_translation(self):
        self.workbook(self.source, [['台词', '한국어'], ['前锚', '앞'], ['甲', '가'], ['乙', '나'],
                                   ['后锚', '뒤'], ['甲乙\n独立', '다른\n독립'], ['甲乙啊', '不可绕过结构冲突']])
        self.workbook(self.baseline, [['Text'], ['前锚'], ['甲乙'], ['后锚'], ['独立']])
        result, _ = self.scan()
        self.assertEqual(result.match_report.recommended_rows, 3)
        book = load_workbook(result.match_report.destination / '匹配明细.xlsx')
        try:
            pending = list(book['结构推荐未写入'].values)[1:]
            self.assertEqual(len(pending), 2)
            self.assertTrue(all('多个结构推荐' in r[-1] for r in pending))
            unused = list(book['未推荐明细'].values)[1:]
            self.assertTrue(all('结构候选未写入' in r[-1] for r in unused if r[4] in ('甲', '乙')))
        finally:
            book.close()
        book = load_workbook(result.match_report.destination / '基准匹配结果/char.xlsx')
        try:
            self.assertIsNone(book.active['AD3'].value)
        finally:
            book.close()

    def test_changed_source_between_structure_and_recommendations_aborts_output(self):
        source = self.workbook(self.source, [['台词', '한국어'], ['甲\n乙', '가\n나']])
        self.workbook(self.baseline, [['Text'], ['甲'], ['乙']])
        def change_and_match(*args, **kwargs):
            book = load_workbook(source)
            book.active['B2'] = '새가\n새나'
            book.save(source)
            book.close()
            return match_folder(*args, **kwargs)
        with patch('structure_scan.match_folder', side_effect=change_and_match):
            with self.assertRaisesRegex(ValueError, '源内容或参与范围已变化'):
                self.scan()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_trial_report_unique_mapping_conflicts_and_formula_like_translation(self):
        source = self.workbook(self.source, [
            ["台词", "한국어"],
            ["1、甲\n2、乙", "1. =literal\n2. 나"],  # After removing the label, still plain text.
            ["1、丙\n2、丁", "1. 다\n2. 라"],
            ["1、戊\n2、己", "1. 마\n2. 바"],
            ["1、戊\n2、己", "1. 다른 번역\n2. 바"],
            ["庚\n辛", "사\n아"],
            ["1、壬\n2、癸", "1. 자\n2. 차"],
        ])
        baseline = self.workbook(self.baseline, [
            ["Text"], ["甲"], ["乙"], [None], ["丙"], ["丁"], [None],
            ["丙"], ["丁"], [None], ["戊"], ["己"], [None], ["庚"], ["辛"],
        ])
        hashes = [hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)]
        result, _ = self.scan()
        self.assertEqual((result.trial_groups, result.trial_segments, result.trial_pending), (2, 4, 4))
        book = load_workbook(result.destination / "试拆分结果.xlsx")
        self.addCleanup(book.close)
        rows = list(book["拆分明细"].values)
        self.assertEqual([row[2:5] for row in rows[1:]],
                         [("甲", "=literal", "甲"), ("乙", "나", "乙"), ("庚", "사", "庚"), ("辛", "아", "辛")])
        self.assertEqual([row[8] for row in rows[1:]], [2, 3, 14, 15])
        self.assertEqual([row[11] for row in rows[1:]], [2, 2, 6, 6])
        self.assertEqual(book["拆分明细"]["D2"].data_type, "s")
        self.assertEqual(book["整组原文"]["B2"].value, "1、甲\n2、乙")
        reasons = [row[3] for row in list(book["暂不拆分"].values)[1:]]
        for expected in ("归属不唯一", "来源冲突", "没有对应"):
            self.assertTrue(any(expected in reason for reason in reasons), expected)
        self.assertEqual(hashes, [hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)])

    def test_boundaries_limits_and_placeholder_mismatch(self):
        self.workbook(self.source, [
            ["台词", "한국어", "特殊说明", "场景"],
            ["甲", "가", None, "场景1"], ["乙", "나", None, "场景2"],
            [None, None], ["丙", "다"], ["废", "폐", "废弃"], ["丁", "라"],
            [None, None], ["隐前", "앞"], ["隐藏", "숨김"], ["隐后", "뒤"],
            [None, None], ["长一", "일"], ["长二", "이"], ["长三", "삼"],
            [None, None], ["前{playername}", "앞"], ["后", "뒤"],
        ], hidden=(10,))
        self.workbook(self.baseline, [["Text"], ["甲乙"], ["丙丁"], ["隐前隐后"], ["长一长二长三"], ["前后"]])
        result, sheets = self.scan(max_rows=2, include_hidden_rows=False)
        self.assertEqual(result.merge_candidates, 0)
        self.assertTrue(any("2 行" in row[7] for row in sheets["未解释"][1:]))
        self.assertTrue(any("排除隐藏行" in row[3] for row in sheets["跳过清单"][1:]))
        result, sheets = self.scan(max_rows=3, include_hidden_rows=False)
        self.assertEqual(result.merge_candidates, 1)
        self.assertIn("长一", sheets["合并候选"][1][5])

    def test_added_number_separators_and_line_split_validation(self):
        for separator in (";", "；", ",", "，", "、"):
            with self.subTest(separator=separator):
                patterns, reason = trial_numbered_split(f"1{separator}甲\n2{separator}乙", "1. 가\n2. 나", ['甲', '乙'])
                self.assertEqual(reason, '')
                self.assertEqual([p.strip() for p in patterns[0]['parts']], ['甲', '乙'])
        self.assertFalse(split_patterns('价格1,234与2,345；小数1.23与2.34'))
        for chinese, korean, expected in [
            ('甲\n乙', '가', '韩文没有'), ('甲\n乙', '가\n나\n다', '段数不一致'),
            ('甲\n\n乙', '가\n\n나', '空分段'), ('甲\n乙', '2.가\n3.나', '歧义'),
            ('甲{name}\n乙', '가\n나', '占位符'),
            ('<color=red>甲\n乙</color>', '가\n나', '标签'),
        ]:
            with self.subTest(chinese=chinese, korean=korean):
                patterns, reason = split_translation(chinese, korean)
                self.assertIsNone(patterns)
                self.assertIn(expected, reason)

    def test_partial_line_preview_keeps_numbered_korean_positions_and_reports_remaining(self):
        source = self.workbook(self.source, [['台词', '한국어'], ['\n甲\\n未找到\n丙\n', '\n1. 가\n2. 나\\n3. 다\n']])
        baseline = self.workbook(self.baseline, [['Text', 'TextTrans'], ['甲', '原译文'], ['丙', '原译文']])
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in (source, baseline)}
        result, sheets = self.scan()
        self.assertEqual((result.trial_groups, result.trial_segments, result.trial_pending), (1, 2, 1))
        book = load_workbook(result.destination / '试拆分结果.xlsx')
        rows = list(book['拆分明细'].values)[1:]
        self.assertEqual([(r[1], r[2], r[3], r[8], r[-1]) for r in rows],
                         [('1', '甲', '1. 가', 2, '2/3段'), ('3', '丙', '3. 다', 3, '2/3段')])
        missing = list(book['分段未匹配'].values)[1:]
        self.assertEqual([r[1:4] for r in missing], [(2, '未找到', '2. 나')])
        self.assertIn('没有对应', missing[0][4])
        self.assertEqual(book['整组原文']['D2'].value, 3)
        book.close()
        unresolved = sheets['未解释'][1:]
        self.assertEqual(len(unresolved), 1)
        self.assertIn('已定位2/3段', unresolved[0][-1])
        self.assertEqual(hashes, {p: hashlib.sha256(p.read_bytes()).digest() for p in hashes})

    def test_trailing_whole_matches_prevent_false_merge_and_unresolved(self):
        self.workbook(self.source, [['台词', '한국어'], ['你好（小声）', '안녕 (작게)'], ['你', '너'], ['好', '좋아'],
                                   ['完整（注）', 'full'], ['中（注）间', '중']])
        self.workbook(self.baseline, [['Text'], ['你好'], ['完整（注）'], ['完整'], ['中间']])
        result, sheets = self.scan()
        self.assertEqual(result.whole_matched_rows, 2)
        self.assertEqual(result.merge_candidates, 0)
        self.assertNotIn('你好（小声）', [r[5] for r in sheets['未解释'][1:]])
        self.assertIn('完整', [r[5] for r in sheets['未解释'][1:]])
        matched = match_folder(scan_folder(self.source), self.baseline, self.output)
        self.assertEqual(matched.matched_rows, result.whole_matched_rows)

    def test_duplicate_candidates_and_misaligned_segment_boundaries_remain_visible(self):
        self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]])
        self.workbook(self.source, [["台词", "한국어"], ["甲", "다"], ["乙", "라"]], name="另一个.xlsx")
        self.workbook(self.source, [["台词", "한국어"], ["A.一二 B.三", "A.하나 B.둘"]], name="选项.xlsx")
        self.workbook(self.baseline, [["Text"], ["甲乙"], ["一"], ["二三"]])
        result, sheets = self.scan()
        self.assertEqual(result.merge_candidates, 2)
        self.assertEqual((result.merge_trial_rows, result.merge_pending_rows), (0, 1))
        self.assertTrue(all("2 个合并候选" in row[12] for row in sheets["合并候选"][1:]))
        self.assertTrue(all("译文序列不同" in row[12] for row in sheets["合并候选"][1:]))
        self.assertTrue(any("分段边界或前缀仍需复核" in row[12] for row in sheets["拆分候选"][1:]))

    def test_formula_sheet_rollback_ignored_sheet_and_atomic_failure(self):
        bad = self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "=1+1"]], name="公式.xlsx")
        ignored = self.workbook(self.source, [["台词", "한국어"], ["甲", "가"], ["乙", "나"]], name="忽略.xlsx")
        self.workbook(self.source, [["台词", "한국어"], ["有效", "유효"]])
        self.workbook(self.baseline, [["Text"], ["甲乙"], ["有效"]])
        result, sheets = self.scan(ignored={sheet_key(ignored, "正文")})
        self.assertEqual(result.source_rows, 1)
        self.assertEqual(result.merge_candidates, 0)
        self.assertTrue(any("正文读取失败" in row[3] for row in sheets["跳过清单"][1:]))
        self.assertTrue(any("本地已忽略" in row[3] for row in sheets["跳过清单"][1:]))
        self.assertTrue(any(r.issue and r.path == bad.resolve() for r in result.check_report.results))
        before = set(self.output.iterdir())
        with self.assertRaisesRegex(ValueError, "2～50"):
            scan_structure(scan_folder(self.source), self.baseline, self.output, max_rows=1)
        self.assertEqual(before, set(self.output.iterdir()))
        with self.assertRaisesRegex(ValueError, "输出文件夹不能"):
            scan_structure(scan_folder(self.source), self.baseline, self.source / "out")


if __name__ == "__main__":
    unittest.main()
