import csv
import hashlib
import tempfile
import unittest
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from unittest.mock import patch

from openpyxl import Workbook

from scanner import export_exceptions, scan_folder
from ignored_sheets import load_ignored, save_ignored, sheet_key


class HeaderChecks(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write_book(self, name, sheets):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, rows in sheets:
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(row)
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()
        return path

    def test_nested_unicode_hidden_filtered_and_source_unchanged(self):
        path = self.root / "子文件夹" / "客户한국어.XLSX"
        path.parent.mkdir()
        workbook = Workbook()
        visible = workbook.active
        visible.title = "正文"
        visible.append(["说明"])
        visible.append([" 文本内容(示例) ", "한국어"])
        visible.row_dimensions[2].hidden = True
        visible.auto_filter.ref = "A1:B3"
        hidden = workbook.create_sheet("隐藏表")
        hidden.append(["语音", "韩文"])
        hidden.sheet_state = "hidden"
        workbook.save(path)
        workbook.close()
        before = hashlib.sha256(path.read_bytes()).digest()
        report = scan_folder(self.root)
        self.assertEqual(report.files, 1)
        self.assertEqual(report.passed_sheets, 2)
        self.assertEqual(report.results[0].header_row, 2)
        self.assertEqual(report.results[1].visibility, "隐藏")
        self.assertFalse(report.exceptions)
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_missing_duplicate_and_separate_rows(self):
        self.write_book("客户.xlsx", [
            ("无表头", [["参考资料"]]),
            ("缺中文", [["한국어"]]),
            ("缺韩语", [["台词"]]),
            ("重复中文", [["语音", "台词", "韩文"]]),
            ("重复韩语", [["语音", "韩文", "한국어"]]),
            ("不同行", [["语音"], [None, "韩文"]]),
            ("重复表头行", [["语音", "韩文"], ["台词", "한국어"]]),
            ("正文关键词", [["语音", "韩文"], ["语音", "正文"]]),
        ])
        report = scan_folder(self.root)
        self.assertEqual(report.abnormal_files, 1)
        self.assertEqual(len(report.exceptions), 7)
        self.assertEqual(report.passed_sheets, 1)
        issues = [result.issue for result in report.exceptions]
        for fragment in ("需确认是否为翻译表", "缺少中文", "缺少韩语", "多个中文", "多个韩语", "不在同一行", "多个中韩表头行"):
            self.assertTrue(any(fragment in issue for issue in issues), fragment)

    def test_limit_uses_actual_row_numbers_and_is_adjustable(self):
        self.write_book("深表头.xlsx", [("Sheet", [[None]] * 50 + [["台词", "韩文"]])])
        self.assertEqual(len(scan_folder(self.root, 50).exceptions), 1)
        self.assertEqual(scan_folder(self.root, 51).passed_sheets, 1)

    def test_default_ten_rows_and_custom_minimum_one(self):
        self.write_book("表头范围.xlsx", [
            ("第一行", [["台词", "韩文"]]),
            ("第十一行", [[None]] * 10 + [["台词", "韩文"]]),
        ])
        default = scan_folder(self.root)
        self.assertEqual(default.header_rows, 10)
        self.assertEqual(default.passed_sheets, 1)
        self.assertEqual([result.sheet for result in default.exceptions], ["第十一行"])
        self.assertEqual(scan_folder(self.root, 1).passed_sheets, 1)
        self.assertEqual(scan_folder(self.root, 11).passed_sheets, 2)
        for rows in (0, -1, 1.5, "10"):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                scan_folder(self.root, rows)

    def test_custom_headers(self):
        self.write_book("定制.xlsx", [("Sheet", [["新中文", "新韩语"]])])
        self.assertEqual(scan_folder(self.root, chinese_headers=["新中文"], korean_headers=["新韩语"]).passed_sheets, 1)

    def test_corrupt_unsupported_and_temporary_files(self):
        self.write_book("正常.xlsx", [("Sheet", [["台词", "韩文"]])])
        (self.root / "损坏.xlsx").write_text("broken")
        (self.root / "旧格式.xls").write_text("old")
        (self.root / "~$忽略.xlsx").write_text("lock")
        (self.root / "说明.txt").write_text("text")
        report = scan_folder(self.root)
        self.assertEqual(report.files, 3)
        self.assertEqual(report.passed_sheets, 1)
        self.assertEqual(report.abnormal_files, 2)
        self.assertTrue(any("文件读取失败" in result.issue for result in report.exceptions))

    def test_wrong_dimension_metadata_does_not_hide_headers(self):
        path = self.write_book("范围错误.xlsx", [("Sheet", [["说明"], ["台词", "韩文"]])])
        with ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b'ref="A1:B2"', b'ref="A1:A1"')
        with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
            for name, value in entries.items():
                archive.writestr(name, value)
        self.assertEqual(scan_folder(self.root).passed_sheets, 1)

    def test_csv_unicode_bom_positions_and_formula_text(self):
        self.write_book("=客户한국어.xlsx", [("=表名", [["语音"]])])
        report = scan_folder(self.root)
        destination = self.root / "异常.csv"
        export_exceptions(report, destination)
        self.assertTrue(destination.read_bytes().startswith(b"\xef\xbb\xbf"))
        with destination.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["文件名"], "'=客户한국어.xlsx")
        self.assertEqual(rows[0]["工作表"], "'=表名")
        self.assertEqual(rows[0]["中文表头位置"], "A1=语音")
        self.assertEqual(rows[0]["检查前几行"], "10")
        self.assertIn("缺少韩语表头", rows[0]["异常原因"])
        self.assertIn("A1=语音", rows[0]["异常原因"])
        self.assertIn("允许的韩语表头", rows[0]["检查依据"])

    def test_diagnostics_show_actual_names_and_other_passed_sheet(self):
        self.write_book("客户.xlsx", [
            ("中文", [["对白文本", "한국어"]]),
            ("日文", [["对白文本", "日文"]]),
            ("修改记录", [["日期", "修改人", "修改内容"]]),
            ("空白表", []),
        ])
        report = scan_folder(self.root)
        by_sheet = {result.sheet: result for result in report.results}
        self.assertTrue(by_sheet["中文"].passed)
        self.assertIn("A1=对白文本", by_sheet["日文"].issue)
        self.assertIn("B1=日文", by_sheet["日文"].details)
        self.assertIn("已找到韩语表头：无", by_sheet["日文"].details)
        self.assertIn("允许的韩语表头：", by_sheet["日文"].details)
        self.assertIn("한국어", by_sheet["日文"].details)
        self.assertIn("A1=日期", by_sheet["修改记录"].details)
        self.assertIn("没有非空单元格", by_sheet["空白表"].issue)
        destination = self.root / "异常.csv"
        export_exceptions(report, destination)
        with destination.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(row["该文件已通过的工作表"] == "中文" for row in rows))

    def test_ambiguous_diagnostics_include_candidate_coordinates(self):
        self.write_book("候选.xlsx", [
            ("不同表头行", [["语音"], [None, "한국어"]]),
            ("多个韩语列", [["语音", "韩文", "한국어"]]),
        ])
        results = scan_folder(self.root).results
        self.assertIn("A1=语音", results[0].issue)
        self.assertIn("B2=한국어", results[0].issue)
        self.assertIn("B1=韩文", results[1].issue)
        self.assertIn("C1=한국어", results[1].issue)

    def test_empty_report_exports_column_names(self):
        self.write_book("正常.xlsx", [("Sheet", [["台词", "韩文"]])])
        destination = self.root / "异常.csv"
        export_exceptions(scan_folder(self.root), destination)
        with destination.open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(len(list(csv.reader(stream))), 1)

    def test_invalid_inputs(self):
        for root, kwargs in ((self.root / "missing", {}), (self.root, {}), (self.root, {"header_rows": 0}), (self.root, {"chinese_headers": []}), (self.root, {"chinese_headers": ["韩文"]})):
            with self.subTest(root=root, kwargs=kwargs), self.assertRaises(ValueError):
                scan_folder(root, **kwargs)

    def test_persisted_ignores_skip_inspection_and_export_then_restore(self):
        path = self.write_book("客户.xlsx", [
            ("中文", [["对白文本", "한국어"]]),
            ("修改记录", [["日期", "修改人"]]),
            ("日文", [["对白文本", "日文"]]),
        ])
        config = self.root / "settings" / "ignored_sheets.json"
        records = {sheet_key(path, "修改记录"), sheet_key(path, "日文")}
        save_ignored(records, config)
        self.assertEqual(load_ignored(config), records)
        from scanner import check_sheet

        with patch("scanner.check_sheet", wraps=check_sheet) as inspect:
            report = scan_folder(self.root, ignored_sheets=load_ignored(config))
        self.assertEqual([call.args[1].title for call in inspect.call_args_list], ["中文"])
        self.assertEqual(report.ignored_sheets, 2)
        self.assertEqual(report.passed_sheets, 1)
        self.assertFalse(report.exceptions)
        destination = self.root / "异常.csv"
        export_exceptions(report, destination)
        with destination.open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(len(list(csv.reader(stream))), 1)
        save_ignored(records - {sheet_key(path, "日文")}, config)
        restored = scan_folder(self.root, ignored_sheets=load_ignored(config))
        self.assertEqual(restored.ignored_sheets, 1)
        self.assertEqual([result.sheet for result in restored.exceptions], ["日文"])

    def test_ignore_is_specific_to_file_even_for_same_sheet_name(self):
        first = self.write_book("first.xlsx", [("修改记录", [["日期"]])])
        self.write_book("second.xlsx", [("修改记录", [["日期"]])])
        report = scan_folder(self.root, ignored_sheets={sheet_key(first, "修改记录")})
        self.assertEqual(report.ignored_sheets, 1)
        self.assertEqual(len(report.exceptions), 1)
        self.assertEqual(report.exceptions[0].path.name, "second.xlsx")

    def test_ignoring_all_sheets_is_not_an_error(self):
        path = self.write_book("客户.xlsx", [("参考表", [["标题"]])])
        report = scan_folder(self.root, ignored_sheets={sheet_key(path, "参考表")})
        self.assertEqual(report.files, 1)
        self.assertEqual(report.ignored_sheets, 1)
        self.assertFalse(report.results)
        self.assertEqual(report.abnormal_files, 0)

    def test_ignore_records_validate_and_do_not_overwrite_on_write_failure(self):
        config = self.root / "ignored_sheets.json"
        self.assertEqual(load_ignored(config), set())
        for text in ('{broken', '{"version":1,"sheets":[{"file":"relative.xlsx","sheet":"参考"}]}', '{"version":2,"sheets":[]}'):
            config.write_text(text, encoding="utf-8")
            with self.assertRaises(ValueError):
                load_ignored(config)
            self.assertEqual(config.read_text(encoding="utf-8"), text)
        records = {sheet_key(self.root / "客户한국어.xlsx", "修改记录")}
        save_ignored(records, config)
        before = config.read_bytes()
        with patch("ignored_sheets.os.replace", side_effect=OSError("disk failure")):
            with self.assertRaises(OSError):
                save_ignored(set(), config)
        self.assertEqual(config.read_bytes(), before)
        self.assertEqual(load_ignored(config), records)
        self.assertFalse(list(self.root.glob("ignored_*.tmp")))


if __name__ == "__main__":
    unittest.main()
